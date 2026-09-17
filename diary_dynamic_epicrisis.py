"""Sick-leave dynamic epicrisis layer for the text diary route.

The regular diary generator is deliberately left untouched.  When sick leave is
selected, this module inserts the historical Dokkomplekt dynamic epicrises into
the already generated text diary in chronological order.
"""
from __future__ import annotations

import os
import re
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from diary_dates import parse_full_date, parse_optional_discharge_date

FIXED_HOLIDAY_RANGES: tuple[tuple[int, int, int], ...] = ((1, 1, 9), (5, 1, 9))
_LEADING_DATE_RE = re.compile(r"^\s*([0-3]?\d)[./-]([01]?\d)[./-](\d{2}|20\d{2})(?=\s|$)")


@dataclass(frozen=True)
class DynamicEpicrisisInput:
    patient_name: str = ""
    birth_date: str = ""
    sick_leave_from: str = ""
    complaints: str = ""
    treatment: str = ""
    profile_status: str = ""
    treatment_correction: str = ""
    treating_physician: str = ""
    department_head: str = ""


def is_fixed_holiday(day: date) -> bool:
    return any(month == day.month and start <= day.day <= end for month, start, end in FIXED_HOLIDAY_RANGES)


def is_non_working_day(day: date) -> bool:
    return day.weekday() >= 5 or is_fixed_holiday(day)


def next_working_day(day: date, *, used=()) -> date:
    used_set = set(used)
    current = day
    for _ in range(370):
        if not is_non_working_day(current) and current not in used_set:
            return current
        current += timedelta(days=1)
    raise RuntimeError("Cannot find an available calendar day within one year.")


def _optional_full_date(value: str) -> date | None:
    text = str(value or "").strip()
    if not text:
        return None
    try:
        return parse_full_date(text)
    except ValueError:
        return None


def dynamic_epicrisis_base_date(admission: date, sick_leave_from: str) -> date:
    """Use the later of admission and sick-leave start, matching Dokkomplekt."""
    sick_leave_date = _optional_full_date(sick_leave_from)
    return max(admission, sick_leave_date) if sick_leave_date is not None else admission


def dynamic_epicrisis_dates(
    admission: date,
    *,
    discharge_date: date | None = None,
    limit: int = 12,
) -> tuple[date, ...]:
    """Plan each 10-day epicrisis, shifting non-working dates forward.

    The discharge date is a hard exclusive boundary: no dynamic epicrisis is
    created on or after discharge, including after a weekend/holiday shift.
    """
    result: list[date] = []
    current = admission + timedelta(days=10)
    while len(result) < limit:
        if discharge_date is not None and current >= discharge_date:
            break
        adjusted = next_working_day(current, used=result)
        if discharge_date is not None and adjusted >= discharge_date:
            break
        result.append(adjusted)
        current += timedelta(days=10)
    return tuple(result)


def _signature_person_name(value: object) -> str:
    text = " ".join(str(value or "").strip().split())
    if not text:
        return ""
    return re.sub(
        r"(?i)^\s*(?:лечащий\s+врач|врач(?:-психиатр)?|заведующ(?:ий|ая)\s+отделением|зав\.?\s*отделением|зав\.?\s*отд\.?)\s*[:—–-]?\s*",
        "",
        text,
    ).strip()


def dynamic_epicrisis_signature_lines(
    treating_physician: object = "",
    department_head: object = "",
) -> tuple[str, str]:
    doctor = _signature_person_name(treating_physician) or "____________________"
    head = _signature_person_name(department_head) or "____________________"
    return (f"Лечащий врач {doctor}", f"Зав.отделением {head}")


def build_dynamic_epicrisis_text(data: DynamicEpicrisisInput) -> str:
    correction = str(data.treatment_correction or "").strip() or "Лекарства принимает согласно назначениям."
    return "\n".join(
        [
            "Динамический эпикриз.",
            f"ФИО: {data.patient_name or 'не указано'}.",
            f"Дата рождения: {data.birth_date or 'не указана'}.",
            f"Лечится с: {data.sick_leave_from or 'не указано'}.",
            f"Жалобы: {data.complaints or 'без существенной динамики'}.",
            f"Принимает: {data.treatment or 'согласно листу назначений'}.",
            f"Профильный статус: {data.profile_status or 'без существенной динамики'}.",
            correction,
            "Продолжение лечения по листу нетрудоспособности.",
            *dynamic_epicrisis_signature_lines(data.treating_physician, data.department_head),
        ]
    )


def _leading_date(text: str) -> date | None:
    match = _LEADING_DATE_RE.match(str(text or ""))
    if not match:
        return None
    year = int(match.group(3))
    if year < 100:
        year += 2000 if year < 70 else 1900
    try:
        return date(year, int(match.group(2)), int(match.group(1)))
    except ValueError:
        return None


def _insert_dynamic_block(doc: Document, item_date: date, lines: tuple[str, ...]) -> None:
    """Insert after all regular blocks on the same date and before later dates."""
    anchor = None
    for paragraph in doc.paragraphs:
        paragraph_date = _leading_date(paragraph.text)
        if paragraph_date is not None and paragraph_date > item_date:
            anchor = paragraph
            break

    rendered = (f"{item_date:%d.%m.%y} {lines[0]}".rstrip(), *lines[1:])
    if anchor is None:
        if doc.paragraphs:
            doc.add_paragraph("")
        for line in rendered:
            paragraph = doc.add_paragraph(line)
            if line.startswith("Лечащий врач ") or line.startswith("Зав.отделением "):
                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        return

    for line in rendered:
        paragraph = anchor.insert_paragraph_before(line)
        if line.startswith("Лечащий врач ") or line.startswith("Зав.отделением "):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    anchor.insert_paragraph_before("")


def apply_sick_leave_dynamic_epicrises(
    path: str | Path,
    *,
    admission_value: str,
    discharge_value: str = "",
    sick_leave_from: str = "",
    patient_name: str = "",
    birth_date: str = "",
    complaints: str = "",
    treatment: str = "",
    profile_status: str = "",
    treatment_correction: str = "",
    treating_physician: str = "",
    department_head: str = "",
) -> int:
    """Atomically add the historical sick-leave epicrisis blocks to one DOCX."""
    target = Path(path)
    admission = parse_full_date(admission_value)
    discharge = parse_optional_discharge_date(discharge_value)
    base_date = dynamic_epicrisis_base_date(admission, sick_leave_from)
    dates = dynamic_epicrisis_dates(base_date, discharge_date=discharge, limit=12)
    if not dates:
        return 0

    data = DynamicEpicrisisInput(
        patient_name=patient_name,
        birth_date=birth_date,
        sick_leave_from=f"{base_date:%d.%m.%Y}",
        complaints=complaints,
        treatment=treatment,
        profile_status=profile_status,
        treatment_correction=treatment_correction,
        treating_physician=treating_physician,
        department_head=department_head,
    )
    lines = tuple(build_dynamic_epicrisis_text(data).splitlines())
    doc = Document(str(target))
    for item_date in dates:
        _insert_dynamic_block(doc, item_date, lines)

    staged = target.with_name(f".{target.name}.dynamic-epicrisis.tmp.docx")
    try:
        doc.save(str(staged))
        os.replace(staged, target)
    finally:
        try:
            staged.unlink()
        except FileNotFoundError:
            pass
    return len(dates)
