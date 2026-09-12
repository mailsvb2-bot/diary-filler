from docx.shared import RGBColor
created, data = service.create_documents(
    navigation_path=nav,
    output_dir=OUT / "medical_with_epi",
    discharge_date="11.06.2026",
    epi_path=epi,
    selected_docs=DOCUMENT_ORDER,
    override_data=manual_data,
)
assert len(created) == len(DOCUMENT_ORDER), created
assert all(path.exists() for path in created)
for created_path in created:
    assert "_" not in created_path.name, created_path.name
assert any(path.name == "Иванова Ирина Ивановна Выписной эпикриз.docx" for path in created), [p.name for p in created]
combined_text = "\n".join(extract_docx_text(path) for path in created)
assert "F99.9 Тестовый диагноз из UI" in combined_text
discharge_path = next(path for path in created if "Выписной" in path.name)
rvk_path = next(path for path in created if "РВК" in path.name)
primary_path = next(path for path in created if "Первичный" in path.name)
commission_path = next(path for path in created if "Совместный" in path.name)
admission_doctor_path = next(path for path in created if "приёмного покоя" in path.name)
vk_mse_path = next(path for path in created if "ВК на МСЭ" in path.name)
sick_leave_vk_path = next(path for path in created if "ВК больничный" in path.name)
discharge_text = extract_docx_text(discharge_path)
rvk_text = extract_docx_text(rvk_path)
primary_text = extract_docx_text(primary_path)
commission_text = extract_docx_text(commission_path)
admission_doctor_text = extract_docx_text(admission_doctor_path)
assert "На основании данных" in discharge_text and "F99.9 Тестовый диагноз из UI" in discharge_text, discharge_text
for occurrence_text in (primary_text, discharge_text, commission_text, admission_doctor_text, rvk_text):
    assert "В 3 отделение КДП поступает повторно добровольно" in occurrence_text, occurrence_text
assert "На учёте у психиатров" not in admission_doctor_text, admission_doctor_text
assert "На учете у психиатров" not in admission_doctor_text, admission_doctor_text
assert "Целесообразна госпитализация" not in commission_text, commission_text
# The bundled discharge placeholder is red, but final clinical text must be explicitly black.
discharge_doc = Document(discharge_path)
period_paragraph = next(
    paragraph for paragraph in discharge_doc.paragraphs
    if paragraph.text.startswith("Находилась на лечении в ГБУЗ НО")
)
period_runs = [run for run in period_paragraph.runs if run.text]
assert period_runs and all(run.font.color.rgb == RGBColor(0, 0, 0) for run in period_runs), [
    (run.text, run.font.color.rgb) for run in period_runs
]
assert "Экспертный анамнез: Работает в ООО Завод, в должности инженер." in primary_text, primary_text
primary_expert_pos = primary_text.index("Экспертный анамнез: Работает в ООО Завод, в должности инженер.")
assert primary_expert_pos > primary_text.index("Эпидемиологический анамнез"), primary_text
assert "Экспертный анамнез: Работает в ООО Завод, в должности инженер. Больничный лист" not in primary_text, primary_text
assert "О СОСТОЯНИИ ЗДОРОВЬЯ ГРАЖДАНИНА № 77-А" in combined_text
assert "Выписка из ПРОТОКОЛА № 42" in combined_text
assert "Выписка из ПРОТОКОЛА № 55" in combined_text
assert "Место работы: ГБУЗ НО Тест, санитар" in combined_text
assert "Место работы: ООО РВК, программист" not in extract_docx_text(next(path for path in created if "РВК" in path.name))
assert "военного комиссариата Ленинского района" in combined_text
assert "По направлению из Ленинского военкомата" in primary_text
assert "Место работы, должность: ООО Тест, инженер" in combined_text
assert "Экспертный анамнез: Работает в ООО Завод, в должности инженер. Больничный лист. Срок лечения с 10.06.2026 по 11.06.2026, 2 дня. К труду с 12.06.2026." in combined_text
assert "Экспертный анамнез: Работает в ООО Завод, в должности инженер. Больничный лист нужен с 15.06.2026." in combined_text
assert "К труду с 12.06.2026" in discharge_text
for path in created:
    if any(key in path.name for key in ("Первичный", "Выписной", "Совместный")):
        assert "Экспертный анамнез:" in extract_docx_text(path), path
assert "Находится на лечении с 10.06.2026 (9 дней)" in combined_text
assert "От 16.06.2026 г." in combined_text
assert "ЭПИ тестовая информация" in combined_text

# Dates entered in the dedicated popups must reach the actual document headers.
commission_doc = Document(commission_path)
assert commission_doc.paragraphs[0].text.startswith("18.06.2026 г. 10:00"), commission_doc.paragraphs[0].text

# A complaints sentence belongs only to the complaints block. It must never be
# duplicated as an unexplained trailing sentence at the end of Joint Examination.
import copy
complaint_data = copy.deepcopy(manual_data)
complaint_data.complaints = "Пациентка предъявляет жалобы на плохой сон"
complaint_created, _ = service.create_documents(
    navigation_path=nav,
    output_dir=OUT / "commission_complaint_position",
    selected_docs=["commission"],
    override_data=complaint_data,
)
complaint_text = extract_docx_text(complaint_created[0])
assert complaint_text.count("Пациентка предъявляет жалобы на плохой сон") == 1, complaint_text
complaint_lines = [line.strip() for line in complaint_text.splitlines() if line.strip()]
complaint_index = next(i for i, line in enumerate(complaint_lines) if "Пациентка предъявляет жалобы на плохой сон" in line)
assert complaint_lines[complaint_index].startswith("Жалобы при поступлении:"), complaint_lines[complaint_index]
assert all("Пациентка предъявляет жалобы на плохой сон" not in line for line in complaint_lines[-4:]), complaint_lines[-4:]
vk_mse_doc = Document(vk_mse_path)
assert any(p.text.strip() == "16.06.2026" for p in vk_mse_doc.paragraphs), [p.text for p in vk_mse_doc.paragraphs[:5]]

# Template yellow highlighting on the protocol date must not survive into either VK document.
for vk_path in (vk_mse_path, sick_leave_vk_path):
    vk_doc = Document(vk_path)
    protocol_date_paragraphs = [p for p in vk_doc.paragraphs if p.text.strip().lower().startswith("от ")]
    assert protocol_date_paragraphs, vk_path
    for paragraph in protocol_date_paragraphs:
        for run in paragraph.runs:
            assert run.font.highlight_color is None, (vk_path.name, paragraph.text, run.text, run.font.highlight_color)
            assert "w:highlight" not in run._r.xml and "w:shd" not in run._r.xml, (vk_path.name, run._r.xml)
        assert "w:shd" not in paragraph._p.xml, (vk_path.name, paragraph._p.xml)

# Clinical sections must be visually separated in generated medical documents.
for readable_path in (primary_path, discharge_path, commission_path, vk_mse_path, sick_leave_vk_path, admission_doctor_path, rvk_path):
    readable_doc = Document(readable_path)
    section_paragraphs = [p for p in readable_doc.paragraphs if p.text.strip().lower().startswith(("жалобы", "анамнез жизни", "анамнез заболевания", "психический статус"))]
    assert section_paragraphs, readable_path
    assert all(p.paragraph_format.space_before is not None and p.paragraph_format.space_before.pt >= 9.5 for p in section_paragraphs), [
        (readable_path.name, p.text, p.paragraph_format.space_before) for p in section_paragraphs
    ]

for treatment_path in (vk_mse_path, sick_leave_vk_path):
    treatment_doc = Document(treatment_path)
    treatment_paragraphs = [p for p in treatment_doc.paragraphs if p.text.strip().lower().startswith("получает лечение")]
    assert treatment_paragraphs, treatment_path
    assert all(p.paragraph_format.space_before is not None and p.paragraph_format.space_before.pt >= 9.5 for p in treatment_paragraphs), [
        (treatment_path.name, p.text, p.paragraph_format.space_before) for p in treatment_paragraphs
    ]

# ЭПИ must stay in its own block and never overwrite laboratory/analysis rows.
# This locks the exact user regression where unrelated text appeared around analyses.
lab_prefixes = (
    "ОАК", "ОАМ", "RW", "HCV", "HBsAg", "ВИЧ", "Биохимия крови",
    "Глюкоза крови", "Кал на яйца глист", "Флюорография", "ЭКГ", "ЭЭГ",
)
for generated in created:
    if not any(key in generated.name for key in ("Выписной", "Совместный", "РВК")):
        continue
    generated_doc = Document(generated)
    visible_paragraphs = [p.text.strip() for p in generated_doc.paragraphs if p.text.strip()]
    for table in generated_doc.tables:
        for row in table.rows:
            seen_cells = set()
            for cell in row.cells:
                cell_id = id(cell._tc)
                if cell_id in seen_cells:
                    continue
                seen_cells.add(cell_id)
                visible_paragraphs.extend(p.text.strip() for p in cell.paragraphs if p.text.strip())
    epi_sentinel_hits = [text for text in visible_paragraphs if "EPI_PLACEMENT_SENTINEL_7F31" in text]
    assert len(epi_sentinel_hits) == 1, (generated.name, epi_sentinel_hits)
    assert epi_sentinel_hits[0].startswith("ЭПИ"), (generated.name, epi_sentinel_hits[0])
    contaminated_labs = [
        text for text in visible_paragraphs
        if text.startswith(lab_prefixes) and "EPI_PLACEMENT_SENTINEL_7F31" in text
    ]
    assert not contaminated_labs, (generated.name, contaminated_labs)

# --- Legacy hospitalization recommendation must not leak into generated clinical blocks ---
phrase_data = service.parse_navigation(nav)
phrase_data.admission = "добровольно Целесообразна госпитализация пациентки в 3 отделение КДП"
phrase_data.admission_occurrence = "первично"
phrase_data.diagnosis = "F41.2 Тест"
phrase_data.commission_date = "18.06.2026"
phrase_data.commission_number = "10"
phrase_data.expert_sick_leave_needed = "нет"
phrase_data.disability_needed = "нет"
phrase_created, _ = service.create_documents(
    navigation_path=nav,
    output_dir=OUT / "medical_with_hospitalization_phrase",
    selected_docs=("primary", "commission"),
    override_data=phrase_data,
)
for path in phrase_created:
    phrase_text = extract_docx_text(path)
    assert "Целесообразна госпитализация" not in phrase_text, (path.name, phrase_text)
    assert "В 3 отделение КДП поступает первично добровольно" in phrase_text, (path.name, phrase_text)

# --- Representative medical document selection combinations must render without failure ---
# Полный перебор всех 2^N комбинаций заметно раздувает smoke-time при добавлении
# нового шаблона. Проверяем все одиночные документы, несколько смешанных наборов
# и полный пакет — этого достаточно для регрессии маршрутизации/шаблонов.
combo_failures = []
representative_combos = []
representative_combos.extend((kind,) for kind in DOCUMENT_ORDER)
representative_combos.extend([
    ("discharge", "admission_doctor_referral"),
    ("primary", "admission_doctor_referral", "rvk"),
    tuple(DOCUMENT_ORDER),
])
for combo in representative_combos:
    try:
        combo_out = OUT / "medical_combinations" / ("_".join(combo))
        combo_created, _ = service.create_documents(
            navigation_path=nav,
            output_dir=combo_out,
            selected_docs=combo,
            override_data=manual_data,
        )
        assert len(combo_created) == len(combo)
        assert all(path.exists() and path.suffix.lower() == ".docx" for path in combo_created)
    except Exception as exc:  # pragma: no cover - smoke report clarity
        combo_failures.append((combo, repr(exc)))
assert not combo_failures, combo_failures[:3]

# --- Generated documents must not keep template choice placeholders ---
choice_placeholder_fragments = (
    "состоит/не состоит",
    "нужен/ не нужен",
    "нужен / не нужен",
    "нужно/ не нужно",
    "нужно / не нужно",
    "да/нет",
)
for path in created:
    doc_text = extract_docx_text(path).lower()
    for fragment in choice_placeholder_fragments:
        assert fragment not in doc_text, (path.name, fragment)

# --- Medical documents without EPI: no ЭПИ mentions should remain ---
manual_no_epi = service.parse_navigation(nav)
manual_no_epi.discharge_date = "11.06.2026"
manual_no_epi.diagnosis = "F88 Диагноз без дополнительного исследования"
manual_no_epi.admission_occurrence = "первично"
manual_no_epi.rvk_act_number = "88-Б"
manual_no_epi.rvk_military_commissariat = "Советского района"
manual_no_epi.rvk_work_position = "не работает"
manual_no_epi.commission_date = "17.06.2026"
manual_no_epi.commission_number = "11"
manual_no_epi.vk_date = "17.06.2026"
manual_no_epi.vk_protocol_number = "43"
manual_no_epi.vk_protocol_date = "17.06.2026"
manual_no_epi.vk_mse_work_org = "не работает"
manual_no_epi.vk_mse_position = ""
manual_no_epi.sick_leave_vk_date = "19.06.2026"
manual_no_epi.sick_leave_vk_protocol_number = "56"
manual_no_epi.sick_leave_vk_protocol_date = "19.06.2026"
manual_no_epi.sick_leave_vk_commission_date = "19.06.2026"
manual_no_epi.sick_leave_vk_work_org = "не работает"
manual_no_epi.sick_leave_vk_position = ""
manual_no_epi.sick_leave_vk_work_position = ""
manual_no_epi.expert_work_status = "нет"
manual_no_epi.expert_sick_leave_needed = "нет"
manual_no_epi.disability_needed = "нет"
manual_no_epi.epi_present = "нет"
created_no_epi, _ = service.create_documents(
    navigation_path=nav,
    output_dir=OUT / "medical_without_epi",
    discharge_date="11.06.2026",
    epi_path=None,
    selected_docs=DOCUMENT_ORDER,
    override_data=manual_no_epi,
)
for path in created_no_epi:
    text = extract_docx_text(path)
    assert "сюда подставлять" not in text.lower(), path
    assert "выбирается в ui" not in text.lower(), path
    assert not __import__("re").search(r"(?<![A-Za-zА-ЯЁа-яё])ЭПИ(?![A-Za-zА-ЯЁа-яё])", text), path


# --- User contract: selection -> merged popup -> exact DOCX kit and content ---
# This is intentionally an application-level contract, not a renderer-only smoke:
# it simulates a doctor selecting output tiles in block 03, answers the popup
# fields, runs the same creation orchestrator, and then reads the produced DOCX.
class _ContractVar:
    def __init__(self, value=""):
        self.value = value

    def get(self):
        return self.value

    def set(self, value):
        self.value = value


class _ContractStatus:
    def __init__(self):
        self.text = ""

    def config(self, **kwargs):
        if "text" in kwargs:
            self.text = kwargs["text"]


class _ContractRoot:
    def __init__(self):
        self.cursor = ""

    def configure(self, **kwargs):
        if "cursor" in kwargs:
            self.cursor = kwargs["cursor"]

    def update_idletasks(self):
        return None


def _make_user_contract_primary(path: Path) -> None:
    doc = Document()
    doc.add_paragraph("10.06.2026 Первичный осмотр")
    doc.add_paragraph("Ф.И.О.: Петров Пётр Петрович")
    doc.add_paragraph("Год рождения: 1977")
    doc.add_paragraph("Зарегистрирован: Н. Новгород, ул. Проверочная, д. 1")
    doc.add_paragraph("Работает в организации: не работает")
    doc.add_paragraph("В 3 отделение КДП поступает добровольно")
    doc.add_paragraph("Жалобы на момент осмотра: тревога, сниженное настроение")
    doc.add_paragraph("Анамнез жизни: Со слов пациента, рос и развивался без особенностей.")
    doc.add_paragraph("Анамнез заболевания: Ухудшение состояния около двух недель.")
    doc.add_paragraph("Психический статус: Контактен, ориентирован, эмоционально напряжён.")
    doc.add_paragraph("Соматический статус: Без грубой соматической патологии.")
    # Важно: явного блока «План лечения» здесь нет, поэтому UI обязан спросить
    # лечение в popup и дальше протащить его в создаваемые DOCX.
    doc.add_paragraph("На основании данных анамнеза жизни и заболевания, психического статуса, данных клинических исследований был выставлен диагноз: F41.2 Смешанное тревожное и депрессивное расстройство")
    doc.save(path)


def _build_contract_app(*, primary_path: Path, output_dir: Path, selected: tuple[str, ...], popup_values: dict[str, str] | None):
    app = _main_module.CombinedMedicalDiaryApp.__new__(_main_module.CombinedMedicalDiaryApp)
    app.root = _ContractRoot()
    app.status_label = _ContractStatus()
    app.service = MedicalDocumentService()
    app._primary_parse_cache = {}
    app._diary_template_files_cache = {}
    app._diary_template_day_cache = {}
    app._diary_template_folder_contains_cache = {}
    app.data = app.service.parse_primary_document(primary_path)
    app._log_buffer = []
    app._last_preview_text = ""
    app._suspend_user_edit_tracking = False
    app._manual_output_dir = True
    app._manual_patient_name = False
    app._manual_admission_date = False
    app._manual_discharge_date = False
    app._manual_diagnosis = False
    app._popup_discharge_date_override = ""
    app._popup_diagnosis_override = ""
    app._work_details_manually_edited = False
    app._primary_work_org_default = ""
    app._primary_work_position_default = ""
    app.diary_files = []

    app.navigation_path_var = _ContractVar(str(primary_path))
    app.output_dir_var = _ContractVar(str(output_dir))
    app.primary_document_type_var = _ContractVar("primary_exam")
    app.patient_name_var = _ContractVar("")
    app.admission_date_var = _ContractVar("")
    app.discharge_date_var = _ContractVar("")
    app.diagnosis_var = _ContractVar("")
    app.case_number_var = _ContractVar("")
    app.admission_occurrence_var = _ContractVar("")
    app.assigned_treatment_var = _ContractVar("")
    app.epi_path_var = _ContractVar("")
    app.epi_present_var = _ContractVar("нет")
    app.strict_mode_var = _ContractVar(False)
    app.printer_var = _ContractVar("")
    app.open_result_folder_var = _ContractVar(True)
    app._opened_output_folders = []

    app.expert_work_status_var = _ContractVar("нет")
    app.expert_work_org_var = _ContractVar("")
    app.expert_position_var = _ContractVar("")
    app.expert_sick_leave_needed_var = _ContractVar("нет")
    app.expert_sick_leave_from_var = _ContractVar("")
    app.expert_sick_leave_number_var = _ContractVar("")
    app.disability_needed_var = _ContractVar("нет")

    app.commission_date_var = _ContractVar("")
    app.commission_number_var = _ContractVar("")
    app.rvk_act_number_var = _ContractVar("")
    app.rvk_military_commissariat_var = _ContractVar("")
    app.rvk_work_position_var = _ContractVar("")
    app.vk_date_var = _ContractVar("")
    app.vk_protocol_number_var = _ContractVar("")
    app.vk_protocol_date_var = _ContractVar("")
    app.vk_mse_work_org_var = _ContractVar("")
    app.vk_mse_position_var = _ContractVar("")
    app.sick_leave_vk_date_var = _ContractVar("")
    app.sick_leave_vk_protocol_number_var = _ContractVar("")
    app.sick_leave_vk_protocol_date_var = _ContractVar("")
    app.sick_leave_vk_commission_date_var = _ContractVar("")
    app.sick_leave_vk_work_org_var = _ContractVar("")
    app.sick_leave_vk_position_var = _ContractVar("")
    app.sick_leave_vk_work_position_var = _ContractVar("")

    app.output_vars = {kind: _ContractVar(kind in selected) for kind in DOCUMENT_ORDER}
    app.output_vars[DIARY_KIND] = _ContractVar(DIARY_KIND in selected)

    popup_calls: list[tuple[str, list[tuple[str, str]], dict[str, tuple[str, ...]] | None]] = []

    def _contract_prompt_fields(title, rows, width=72, linked_groups=None, choice_options=None):
        popup_calls.append((title, list(rows), choice_options))
        if popup_values is None:
            return None
        values: list[str] = []
        for label, _default in rows:
            assert label in popup_values, (title, label, rows)
            values.append(popup_values[label])
        return values

    app._prompt_fields = _contract_prompt_fields
    app._write_creation_report = lambda **kwargs: None
    app._redraw_selection_controls = lambda: None
    app._update_expert_sick_leave_display = lambda: None
    app._open_result_folder_silent = lambda folder: app._opened_output_folders.append(Path(folder))
    return app, popup_calls


from tkinter import messagebox as _contract_messagebox
_original_messagebox_functions = {
    "showinfo": _contract_messagebox.showinfo,
    "showwarning": _contract_messagebox.showwarning,
    "showerror": _contract_messagebox.showerror,
    "askyesno": _contract_messagebox.askyesno,
}
_contract_messagebox_events: list[tuple[str, str]] = []
_contract_messagebox.showinfo = lambda title, message, **kwargs: _contract_messagebox_events.append(("info", title))
_contract_messagebox.showwarning = lambda title, message, **kwargs: _contract_messagebox_events.append(("warning", title))
_contract_messagebox.showerror = lambda title, message, **kwargs: _contract_messagebox_events.append(("error", title))
_contract_messagebox.askyesno = lambda title, message, **kwargs: (_contract_messagebox_events.append(("askyesno", title)) or True)
try:
    contract_dir = OUT / "user_contract_selection_popup_docx"
    if contract_dir.exists():
        shutil.rmtree(contract_dir)
    contract_dir.mkdir(parents=True, exist_ok=True)
    contract_primary = contract_dir / "Первичный_без_лечения.docx"
    _make_user_contract_primary(contract_primary)

    contract_app, contract_popup_calls = _build_contract_app(
        primary_path=contract_primary,
        output_dir=contract_dir / "created",
        selected=("primary", "discharge"),
        popup_values={
            "Номер истории болезни": "К-900",
            "Нужен ли больничный лист": "нет",
            "Нужно ли оформление инвалидности": "нет",
            "Есть ли ЭПИ": "нет",
            "Лечение": "терапия из пользовательского popup",
            "Поступает в 3 отделение КДП": "первично",
            "Дата выписки": "11062026",
        },
    )
    assert contract_app.selected_medical_docs() == ["primary", "discharge"]
    contract_app.create_selected_outputs(print_after=False)
    assert len(contract_popup_calls) == 2, contract_popup_calls
    assert contract_popup_calls[0][0] == "Дополнительные данные"
    assert [label for label, _default in contract_popup_calls[0][1]] == [
        "Нужен ли больничный лист",
        "Нужно ли оформление инвалидности",
        "Есть ли ЭПИ",
    ]
    assert contract_popup_calls[0][2] == {
        "Нужен ли больничный лист": ("нет", "да"),
        "Нужно ли оформление инвалидности": ("нет", "да"),
        "Есть ли ЭПИ": ("нет", "да"),
    }
    assert contract_popup_calls[1][0] == "Данные для выписного эпикриза"
    assert [label for label, _default in contract_popup_calls[1][1]] == [
        "Номер истории болезни",
        "Лечение",
        "Поступает в 3 отделение КДП",
        "Дата выписки",
    ]
    assert contract_popup_calls[1][2] == {"Поступает в 3 отделение КДП": ("первично", "повторно")}

    contract_created = sorted((contract_dir / "created").glob("*.docx"))
    assert [path.name for path in contract_created] == [
        "Петров Пётр Петрович Выписной эпикриз.docx",
        "Петров Пётр Петрович Первичный осмотр.docx",
    ], [path.name for path in contract_created]
    contract_text_by_name = {path.name: extract_docx_text(path) for path in contract_created}
    contract_primary_text = contract_text_by_name["Петров Пётр Петрович Первичный осмотр.docx"]
    contract_discharge_text = contract_text_by_name["Петров Пётр Петрович Выписной эпикриз.docx"]
    assert "История болезни № К-900" in contract_primary_text, contract_primary_text
    assert "План лечения: терапия из пользовательского popup" in contract_primary_text, contract_primary_text
    assert "Выписной эпикриз № К-900" in contract_discharge_text, contract_discharge_text
    assert "по 11.06.2026" in contract_discharge_text, contract_discharge_text
    assert "В 3 отделение КДП поступает первично добровольно" in contract_discharge_text, contract_discharge_text
    assert "Лечение: терапия из пользовательского popup" in contract_discharge_text, contract_discharge_text
    assert "F41.2 Смешанное тревожное и депрессивное расстройство" in contract_discharge_text, contract_discharge_text
    assert contract_app._opened_output_folders == [contract_dir / "created"], contract_app._opened_output_folders
    assert not any(event[0] in {"info", "warning", "error", "askyesno"} for event in _contract_messagebox_events), _contract_messagebox_events

    # Whole-set transaction: medical generation may succeed internally, but if
    # diaries then fail the final user folder must receive none of that staged
    # medical output. This locks the top-level all-or-nothing user contract.
    _contract_messagebox_events.clear()
    rollback_dir = OUT / "user_contract_atomic_rollback"
    if rollback_dir.exists():
        shutil.rmtree(rollback_dir)
    rollback_dir.mkdir(parents=True, exist_ok=True)
    rollback_primary = rollback_dir / "Первичный_без_лечения.docx"
    _make_user_contract_primary(rollback_primary)
    rollback_output = rollback_dir / "created"
    rollback_app, rollback_popup_calls = _build_contract_app(
        primary_path=rollback_primary,
        output_dir=rollback_output,
        selected=("primary", "discharge", DIARY_KIND),
        popup_values={
            "Номер истории болезни": "К-901",
            "Нужен ли больничный лист": "нет",
            "Нужно ли оформление инвалидности": "нет",
            "Есть ли ЭПИ": "нет",
            "Лечение": "терапия для rollback",
            "Поступает в 3 отделение КДП": "повторно",
            "Дата выписки": "11062026",
        },
    )
    rollback_snapshot_ids = {}
    _rollback_original_medical_impl = rollback_app._create_medical_documents_impl

    def _rollback_recording_medical(*args, **kwargs):
        snapshot = kwargs.get("patient_data_snapshot")
        assert snapshot is not None
        rollback_snapshot_ids["medical"] = id(snapshot)
        result = _rollback_original_medical_impl(*args, **kwargs)
        # Simulate a live UI mutation between the two generators. The diary
        # route must still receive the immutable-per-run patient snapshot.
        rollback_app.patient_name_var.set("Чужой Пациент Из UI")
        rollback_app.discharge_date_var.set("01012099")
        return result

    def _rollback_failing_diary(**kwargs):
        snapshot = kwargs.get("patient_data_snapshot")
        assert snapshot is not None
        rollback_snapshot_ids["diary"] = id(snapshot)
        assert snapshot.output_fio == "Петров Пётр Петрович", snapshot
        assert snapshot.discharge_date == "11.06.2026", snapshot
        assert snapshot.output_fio != rollback_app.patient_name_var.get()
        assert snapshot.discharge_date != rollback_app.discharge_date_var.get()
        raise RuntimeError("synthetic diary failure")

    rollback_app._create_medical_documents_impl = _rollback_recording_medical
    rollback_app._create_diaries_impl = _rollback_failing_diary
    rollback_app.create_selected_outputs(print_after=False)
    assert rollback_snapshot_ids["medical"] == rollback_snapshot_ids["diary"], rollback_snapshot_ids
    assert len(rollback_popup_calls) == 2, rollback_popup_calls
    assert rollback_popup_calls[0][0] == "Дополнительные данные"
    assert rollback_popup_calls[1][0] == "Данные для выписного эпикриза"
    assert rollback_output.exists(), rollback_output
    assert not list(rollback_output.glob("*.docx")), list(rollback_output.glob("*.docx"))
    assert not list(rollback_output.glob(".medical-autofill-set-*")), list(rollback_output.iterdir())
    assert ("error", "Комплект не создан") in _contract_messagebox_events, _contract_messagebox_events
    assert rollback_app._opened_output_folders == [], rollback_app._opened_output_folders

    _contract_messagebox_events.clear()
    cancel_dir = OUT / "user_contract_cancelled_popup"
    if cancel_dir.exists():
        shutil.rmtree(cancel_dir)
    cancel_dir.mkdir(parents=True, exist_ok=True)
    cancel_primary = cancel_dir / "Первичный_без_лечения.docx"
    _make_user_contract_primary(cancel_primary)
    cancel_app, cancel_popup_calls = _build_contract_app(
        primary_path=cancel_primary,
        output_dir=cancel_dir / "created",
        selected=("primary", "discharge"),
        popup_values=None,
    )
    cancel_app.create_selected_outputs(print_after=False)
    assert len(cancel_popup_calls) == 1, cancel_popup_calls
    assert not (cancel_dir / "created").exists() or not list((cancel_dir / "created").glob("*.docx"))
finally:
    for _name, _func in _original_messagebox_functions.items():
        setattr(_contract_messagebox, _name, _func)
