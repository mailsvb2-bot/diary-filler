"""Windows GUI wiring smoke for the real Tk application.

This check intentionally instantiates the production CombinedMedicalDiaryApp,
uses real Tk widgets/TkDND, and drives a couple of user-facing controls through
mouse events. Heavy document correctness remains covered by the DOCX smoke
suite; this file closes the UI-wiring gap between headless method tests and the
packaged EXE startup probe.
"""

from __future__ import annotations

import os
from pathlib import Path
from tempfile import TemporaryDirectory
import time

from docx import Document

import files_mixin
from app import CombinedMedicalDiaryApp
from startup import _create_root


def _click(root, widget) -> None:
    # Keep press/release contiguous. Pumping Tk's idle queue between them can
    # synthesize a <Leave> from the runner's real mouse position, which resets
    # the production button's pressed state before ButtonRelease arrives.
    root.update_idletasks()
    widget.event_generate("<Enter>")
    widget.event_generate("<ButtonPress-1>", x=8, y=8)
    widget.event_generate("<ButtonRelease-1>", x=8, y=8)
    root.update_idletasks()


def main() -> None:
    if os.name != "nt":
        print("GUI RUNTIME CHECK SKIPPED: Windows-only")
        return

    with TemporaryDirectory(prefix="medical-autofill-gui-") as temp_dir:
        temp_root = Path(temp_dir)
        os.environ["APPDATA"] = str(temp_root / "appdata")

        dates_dir = temp_root / "Даты"
        dates_dir.mkdir()
        Document().save(dates_dir / "01.docx")
        texts_dir = temp_root / "Тексты"
        texts_dir.mkdir()

        # Real production contract: filenames in the physician's large text
        # folder are WORDS, not ICD codes. The first file is deliberately a
        # descriptive phrase; the second uses the informal «шизофреника» form.
        initial_text_name = "дневники на тревожно депрессивное расстройство.docx"
        schizophrenia_text_name = "дневники на шизофреника.docx"
        text_doc = Document()
        text_doc.add_paragraph("Пациент спокоен, жалоб не предъявляет.")
        text_doc.save(texts_dir / initial_text_name)
        text_doc = Document()
        text_doc.add_paragraph("Пациент доступен продуктивному контакту.")
        text_doc.save(texts_dir / schizophrenia_text_name)

        root = _create_root(require_dnd=True)
        original_askdirectory = files_mixin.filedialog.askdirectory
        original_askopenfilename = files_mixin.filedialog.askopenfilename
        try:
            app = CombinedMedicalDiaryApp(root)
            # Mouse events on custom Canvas buttons are meaningful only after the
            # real top-level and its descendants are mapped. A withdrawn window
            # can still accept direct bindings (the drop-zone) while discarding
            # stateful press/release button semantics.
            root.update_idletasks()
            root.deiconify()
            root.update()

            assert app._register_tkinterdnd_drop_targets(), "TkDND targets did not register"
            assert hasattr(app, "drop_zone"), "Primary drop zone missing"
            assert hasattr(app, "diary_dates_button"), "Dates button missing"
            assert hasattr(app, "status_files_button"), "Texts button missing"
            assert hasattr(app, "status_file_button"), "Diary text folder button missing"
            assert app.drop_zone.winfo_ismapped(), "Primary drop-zone is not mapped"
            assert app.diary_dates_button.winfo_ismapped(), "Dates button is not mapped"
            assert app.status_files_button.winfo_ismapped(), "Texts button is not mapped"
            assert app.status_file_button.winfo_ismapped(), "Diary text folder button is not mapped"

            # Drive the actual visible primary drop-zone. Its click binding must
            # call the production navigation chooser; dormant compatibility
            # controls are not treated as part of the current user path.
            navigation_clicks: list[str] = []
            app.choose_navigation = lambda: navigation_clicks.append("choose_navigation")
            _click(root, app.drop_zone)
            assert navigation_clicks == ["choose_navigation"], "Primary drop-zone click is not wired"

            # User clicks «Даты»: exactly one folder dialog is used and the
            # numbered template is selected from that folder.
            app.admission_date_var.set("01.01.2026")
            calls: list[str] = []

            def choose_dates_folder(*_args, **_kwargs):
                calls.append("askdirectory")
                return str(dates_dir)

            files_mixin.filedialog.askdirectory = choose_dates_folder
            _click(root, app.diary_dates_button)
            assert calls == ["askdirectory"], f"Unexpected Dates-dialog flow: {calls!r}"
            assert Path(app.diary_template_dir) == dates_dir
            assert app.diary_files and Path(app.diary_files[0]).name == "01.docx"

            # User clicks «Тексты»: the primary action is a concrete Word-file
            # picker. This is the installed user path that must visibly expose
            # .doc/.docx/.docm instead of opening a folder-only chooser.
            initial_diagnosis = "F41.2 Смешанное тревожно-депрессивное расстройство"
            app.diagnosis_var.set(initial_diagnosis)
            file_calls: list[str] = []

            def choose_text_file(*_args, **kwargs):
                file_calls.append(str(kwargs.get("filetypes", "")))
                return str(texts_dir / initial_text_name)

            files_mixin.filedialog.askopenfilename = choose_text_file
            _click(root, app.status_files_button)
            assert len(file_calls) == 1, f"Unexpected Texts file-dialog flow: {file_calls!r}"
            assert "*.doc" in file_calls[0] and "*.docx" in file_calls[0] and "*.docm" in file_calls[0], file_calls
            assert app.status_files == [str(texts_dir / initial_text_name)]
            assert app._diary_text_files_auto_selected is False

            # User clicks «Папка»: this is the separate directory chooser used
            # only to enable automatic VERBAL diagnosis matching. The folder has
            # no ICD-coded filenames; matching must use words only.
            folder_calls: list[str] = []

            def choose_texts_folder(*_args, **_kwargs):
                folder_calls.append("askdirectory")
                return str(texts_dir)

            files_mixin.filedialog.askdirectory = choose_texts_folder
            _click(root, app.status_file_button)
            assert folder_calls == ["askdirectory"], f"Unexpected Folder-dialog flow: {folder_calls!r}"
            assert Path(app.diary_texts_dir) == texts_dir
            assert app.status_files and Path(app.status_files[0]).name == initial_text_name, app.status_files
            assert app._diary_text_files_auto_selected is True

            # Regression: changing Diagnosis in the visible patient card is a
            # clinical state change. The new value must immediately replace the
            # parsed/popup diagnosis and invalidate the automatically selected
            # text from the old diagnosis. The replacement filename contains NO
            # F20/20/etc; «Параноидная шизофрения» must find «шизофреника» by
            # words, without reloading the primary document.
            new_diagnosis = "F20.0 Параноидная шизофрения"
            app.diagnosis_entry.delete(0, "end")
            app.diagnosis_entry.insert(0, new_diagnosis)
            app.diagnosis_entry.event_generate("<KeyRelease>")
            root.update_idletasks()
            assert app.data.diagnosis == new_diagnosis, app.data.diagnosis
            assert app._popup_diagnosis_override == new_diagnosis, app._popup_diagnosis_override
            assert app.status_files == [], app.status_files
            time.sleep(0.25)
            root.update()
            assert app.status_files and Path(app.status_files[0]).name == schizophrenia_text_name, app.status_files
            assert app._diary_text_files_auto_selected is True

            print("GUI RUNTIME CHECK OK")
        finally:
            files_mixin.filedialog.askdirectory = original_askdirectory
            files_mixin.filedialog.askopenfilename = original_askopenfilename
            try:
                root.destroy()
            except Exception:
                pass


if __name__ == "__main__":
    main()
