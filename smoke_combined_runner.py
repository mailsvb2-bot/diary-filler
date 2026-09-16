"""Runner for the split combined smoke regression suite."""

from __future__ import annotations

from pathlib import Path

PARTS = (
    "smoke_combined_part01_setup_contracts.py",
    "smoke_combined_part02_ui_parser_regressions.py",
    "smoke_combined_part03_medical_parser_manual.py",
    "smoke_combined_part04_medical_generation.py",
    "smoke_combined_part05_diary_basic_templates.py",
    "smoke_combined_part06_diary_columns_settings.py",
)


def _run_docx_cache_regression() -> None:
    """Prove stable cache hits, invalidation, and fail-closed unstable reads."""
    import os
    from tempfile import TemporaryDirectory

    from docx import Document

    import medical_docx_blocks as docx_blocks

    with TemporaryDirectory(prefix="medical-autofill-docx-cache-") as temp_dir:
        root = Path(temp_dir)

        stable_path = root / "stable.docx"
        stable_doc = Document()
        stable_doc.add_paragraph("Первичный осмотр")
        stable_doc.add_paragraph("CACHE_SENTINEL_ALPHA")
        stable_doc.save(stable_path)

        original_document_loader = docx_blocks.Document
        stable_open_calls: list[str] = []

        def counting_document_loader(path):
            stable_open_calls.append(str(path))
            return original_document_loader(path)

        docx_blocks.Document = counting_document_loader
        try:
            stable_first = docx_blocks.extract_docx_text(stable_path)
            stable_second = docx_blocks.extract_docx_text(stable_path)
            assert stable_first == stable_second
            assert "CACHE_SENTINEL_ALPHA" in stable_first
            assert len(stable_open_calls) == 1, "unchanged DOCX must be opened only once"

            before_change = stable_path.stat()
            changed_doc = Document(str(stable_path))
            changed_doc.add_paragraph("CACHE_SENTINEL_BETA")
            changed_doc.save(stable_path)
            after_change = stable_path.stat()
            if (
                after_change.st_size == before_change.st_size
                and after_change.st_mtime_ns == before_change.st_mtime_ns
            ):
                os.utime(
                    stable_path,
                    ns=(after_change.st_atime_ns, after_change.st_mtime_ns + 1_000_000),
                )

            stable_third = docx_blocks.extract_docx_text(stable_path)
            assert "CACHE_SENTINEL_BETA" in stable_third
            assert len(stable_open_calls) == 2, "changed DOCX must force a real re-read"
        finally:
            docx_blocks.Document = original_document_loader

        racy_path = root / "racy.docx"
        racy_doc = Document()
        racy_doc.add_paragraph("CACHE_SENTINEL_RACE")
        racy_doc.save(racy_path)

        racy_open_calls: list[str] = []
        racy_mutated = [False]

        def mutating_document_loader(path):
            loaded = original_document_loader(path)
            racy_open_calls.append(str(path))
            candidate = Path(path)
            if candidate.resolve() == racy_path.resolve() and not racy_mutated[0]:
                stat = candidate.stat()
                os.utime(candidate, ns=(stat.st_atime_ns, stat.st_mtime_ns + 1_000_000))
                racy_mutated[0] = True
            return loaded

        docx_blocks.Document = mutating_document_loader
        try:
            racy_first = docx_blocks.extract_docx_text(racy_path)
            racy_second = docx_blocks.extract_docx_text(racy_path)
            racy_third = docx_blocks.extract_docx_text(racy_path)
            assert racy_first == racy_second == racy_third
            assert "CACHE_SENTINEL_RACE" in racy_first
            assert len(racy_open_calls) == 2, "unstable read must not be cached; stable retry must be"
        finally:
            docx_blocks.Document = original_document_loader


def _run_facility_reference_regression() -> None:
    """Prove facility prefilters preserve every established replacement route."""
    from docx import Document

    from medical_constants import TARGET_MEDICAL_FACILITY
    from medical_gender import normalize_facility_references_in_document

    sources_and_expected = (
        ("ГБУЗ НО ПБ №2", TARGET_MEDICAL_FACILITY),
        ("гбуз но пб № 2", TARGET_MEDICAL_FACILITY),
        (
            "ГБУЗНО «Психиатрическая больница № 2» г. Н. Новгорода",
            TARGET_MEDICAL_FACILITY,
        ),
        (
            "ГБУЗ НО «Психиатрическая больница №2» г. Н. Новгорода",
            TARGET_MEDICAL_FACILITY,
        ),
        (
            "Переведен в отделение №3 для продолжения лечения",
            f"Переведен в {TARGET_MEDICAL_FACILITY} для продолжения лечения",
        ),
        (
            "Из отделения № 3 переведен по профилю",
            f"Из {TARGET_MEDICAL_FACILITY} переведен по профилю",
        ),
        ("Направляется на лечение в старое учреждение", f"Направляется в {TARGET_MEDICAL_FACILITY}"),
        ("   Направляется   в ГБУЗ НО ПБ №2", f"Направляется в {TARGET_MEDICAL_FACILITY}"),
        ("Психиатрическая помощь амбулаторно", "Психиатрическая помощь амбулаторно"),
        ("отделение №4", "отделение №4"),
        ("Без специальных слов", "Без специальных слов"),
    )

    doc = Document()
    for source, _expected in sources_and_expected:
        doc.add_paragraph(source)

    normalize_facility_references_in_document(doc)

    actual = tuple(paragraph.text for paragraph in doc.paragraphs)
    expected = tuple(expected for _source, expected in sources_and_expected)
    assert actual == expected, f"facility reference regression mismatch: {actual!r}"



def _run_gender_regex_cache_regression() -> None:
    """Prove regex caching preserves gender conversion while reusing patterns."""
    import shared_gender

    pattern_cache = shared_gender._gender_pair_pattern
    pattern_cache.cache_clear()

    source = "Находился спокоен, ориентирован и был выписан."
    first = shared_gender.adapt_text_to_patient_gender(source, "female")
    first_info = pattern_cache.cache_info()
    second = shared_gender.adapt_text_to_patient_gender(source, "female")
    second_info = pattern_cache.cache_info()

    assert first == second
    assert first[0] == "Находилась спокойна, ориентирована и была выписана."
    assert first[1] == 5
    assert first_info.misses > 0
    assert second_info.misses == first_info.misses
    assert second_info.hits > first_info.hits
    assert second_info.maxsize == 512

def _run_normalize_match_fast_path_regression() -> None:
    """Prove match-only normalization keeps established Unicode semantics."""
    from medical_text_utils import normalize_match

    sources_and_expected = (
        ("", ""),
        ("  Анамнез\tжизни  ", "анамнез жизни"),
        ("Ёжик — тест", "ежик - тест"),
        ("ГБУЗ\xa0НО\vПБ №2", "гбуз но пб №2"),
        ("A\n\n\nB", "a b"),
        ("X\u2003Y\u202fZ", "x y z"),
        ("слово‑слово – слово − слово", "слово-слово - слово - слово"),
    )
    for source, expected in sources_and_expected:
        actual = normalize_match(source)
        assert actual == expected, f"normalize_match mismatch: {source!r} -> {actual!r}"


def run() -> None:
    root = Path(__file__).resolve().parent
    namespace = {"__name__": "__smoke_combined__", "__file__": str(root / "smoke_test_combined.py")}
    for part_name in PARTS:
        part_path = root / part_name
        code = compile(part_path.read_text(encoding="utf-8"), str(part_path), "exec")
        exec(code, namespace, namespace)
    _run_docx_cache_regression()
    _run_facility_reference_regression()
    _run_gender_regex_cache_regression()
    _run_normalize_match_fast_path_regression()
