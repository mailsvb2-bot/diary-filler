"""Focused regression for the bounded DOCX text cache."""
from __future__ import annotations

import os
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document

import medical_docx_blocks as _docx_blocks


with TemporaryDirectory(prefix="medical-autofill-docx-cache-") as _cache_temp_dir:
    _cache_root = Path(_cache_temp_dir)

    # Unchanged file: the second request must reuse the extracted normalized text
    # instead of opening the same Word package again.
    _stable_path = _cache_root / "stable.docx"
    _stable_doc = Document()
    _stable_doc.add_paragraph("Первичный осмотр")
    _stable_doc.add_paragraph("CACHE_SENTINEL_ALPHA")
    _stable_doc.save(_stable_path)

    _original_document_loader = _docx_blocks.Document
    _stable_open_calls: list[str] = []

    def _counting_document_loader(path):
        _stable_open_calls.append(str(path))
        return _original_document_loader(path)

    _docx_blocks.Document = _counting_document_loader
    try:
        _stable_first = _docx_blocks.extract_docx_text(_stable_path)
        _stable_second = _docx_blocks.extract_docx_text(_stable_path)
        assert _stable_first == _stable_second
        assert "CACHE_SENTINEL_ALPHA" in _stable_first
        assert len(_stable_open_calls) == 1, "unchanged DOCX must be opened only once"

        # Changing the file must invalidate the exact path+size+mtime_ns entry.
        _before_change = _stable_path.stat()
        _changed_doc = Document(str(_stable_path))
        _changed_doc.add_paragraph("CACHE_SENTINEL_BETA")
        _changed_doc.save(_stable_path)
        _after_change = _stable_path.stat()
        if (
            _after_change.st_size == _before_change.st_size
            and _after_change.st_mtime_ns == _before_change.st_mtime_ns
        ):
            os.utime(
                _stable_path,
                ns=(_after_change.st_atime_ns, _after_change.st_mtime_ns + 1_000_000),
            )

        _stable_third = _docx_blocks.extract_docx_text(_stable_path)
        assert "CACHE_SENTINEL_BETA" in _stable_third
        assert len(_stable_open_calls) == 2, "changed DOCX must force a real re-read"
    finally:
        _docx_blocks.Document = _original_document_loader

    # If the file signature changes while python-docx is opening it, that read
    # must not populate the cache. The next call therefore has to open it again.
    _racy_path = _cache_root / "racy.docx"
    _racy_doc = Document()
    _racy_doc.add_paragraph("CACHE_SENTINEL_RACE")
    _racy_doc.save(_racy_path)

    _racy_open_calls: list[str] = []
    _racy_mutated_during_first_read = False

    def _mutating_document_loader(path):
        global _racy_mutated_during_first_read
        loaded = _original_document_loader(path)
        _racy_open_calls.append(str(path))
        candidate = Path(path)
        if candidate.resolve() == _racy_path.resolve() and not _racy_mutated_during_first_read:
            stat = candidate.stat()
            os.utime(candidate, ns=(stat.st_atime_ns, stat.st_mtime_ns + 1_000_000))
            _racy_mutated_during_first_read = True
        return loaded

    _docx_blocks.Document = _mutating_document_loader
    try:
        _racy_first = _docx_blocks.extract_docx_text(_racy_path)
        _racy_second = _docx_blocks.extract_docx_text(_racy_path)
        _racy_third = _docx_blocks.extract_docx_text(_racy_path)
        assert _racy_first == _racy_second == _racy_third
        assert "CACHE_SENTINEL_RACE" in _racy_first
        assert len(_racy_open_calls) == 2, "unstable read must not be cached; stable retry must be"
    finally:
        _docx_blocks.Document = _original_document_loader
