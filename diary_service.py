"""Production service boundary for user-facing text diaries."""

from __future__ import annotations

import os
import re
from dataclasses import dataclass
from datetime import date, timedelta
from pathlib import Path
from tempfile import TemporaryDirectory
from typing import Mapping, Sequence

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from diary_batch import _clinical_diary_offsets, create_text_diaries
from diary_calendar import is_fixed_holiday, is_non_working_day
from diary_dates import parse_full_date, parse_optional_discharge_date
from diary_models import DiaryBatchResult


_LEADING_DATE_RE = re.compile(r"^\s*([0-3]?\d)[./-]([01]?\d)[./-](\d{2}|20\d{2})(?=\s|$)")


@dataclass(frozen=True)
class DynamicEpicrisisInput:
    patient_name: str = ""
    birth_date: str = ""
    sick_leave_from: str = ""
    # Clinical wording selected for this exact epicrisis date from the semantic
    # diary plan. Never substitute a neighboring/admission observation.
    clinical_state: str = ""
    complaints: str = ""
    treatment: str = ""
    profile_status: str = ""
    treatment_correction: str = ""
    treating_physician: str = ""
    department_head: str = ""


def next_working_day(day: date, *, used=()) -> date:
    used_set = set(used)
    current = day
    for _ in range(370):
        if not is_non_working_day(current) and current not in used_set:
            return current
        current += timedelta(days=1)
    raise RuntimeError("Cannot find an available calendar day within one year.")


def _optional_full_date(value: str) -> date | None:
    text = str(value or "").strip()
    if not text:
        return None
    try:
        return parse_full_date(text)
    except ValueError:
        return None


def dynamic_epicrisis_base_date(admission: date, sick_leave_from: str) -> date:
    """Use the later of admission and sick-leave start, matching Dokkomplekt."""
    sick_leave_date = _optional_full_date(sick_leave_from)
    return max(admission, sick_leave_date) if sick_leave_date is not None else admission


def dynamic_epicrisis_dates(
    admission: date,
    *,
    discharge_date: date | None = None,
    limit: int = 12,
) -> tuple[date, ...]:
    """Plan each 10-day epicrisis, shifting non-working dates forward.

    The discharge date is a hard exclusive boundary: no dynamic epicrisis is
    created on or after discharge, including after a weekend/holiday shift.
    """
    result: list[date] = []
    current = admission + timedelta(days=10)
    while len(result) < limit:
        if discharge_date is not None and current >= discharge_date:
            break
        adjusted = next_working_day(current, used=result)
        if discharge_date is not None and adjusted >= discharge_date:
            break
        result.append(adjusted)
        current += timedelta(days=10)
    return tuple(result)


def _signature_person_name(value: object) -> str:
    text = " ".join(str(value or "").strip().split())
    if not text:
        return ""
    return re.sub(
        r"(?i)^\s*(?:лечащий\s+врач|врач(?:-психиатр)?|заведующ(?:ий|ая)\s+отделением|зав\.?\s*отделением|зав\.?\s*отд\.?)\s*[:—–-]?\s*",
        "",
        text,
    ).strip()


def dynamic_epicrisis_signature_lines(
    treating_physician: object = "",
    department_head: object = "",
) -> tuple[str, str]:
    doctor = _signature_person_name(treating_physician) or "____________________"
    head = _signature_person_name(department_head) or "____________________"
    return (f"Лечащий врач {doctor}", f"Зав.отделением {head}")


def build_dynamic_epicrisis_text(data: DynamicEpicrisisInput) -> str:
    """Build only from explicit patient/source facts; never invent clinical state."""
    lines = ["Динамический эпикриз."]
    patient_name = str(data.patient_name or "").strip()
    birth_date = str(data.birth_date or "").strip()
    sick_leave_from = str(data.sick_leave_from or "").strip()
    clinical_state = str(data.clinical_state or "").strip()
    complaints = str(data.complaints or "").strip()
    treatment = str(data.treatment or "").strip()
    profile_status = str(data.profile_status or "").strip()
    correction = str(data.treatment_correction or "").strip()

    if patient_name:
        lines.append(f"ФИО: {patient_name}.")
    if birth_date:
        lines.append(f"Дата рождения: {birth_date}.")
    if sick_leave_from:
        lines.append(f"Лечится с: {sick_leave_from}.")
    if clinical_state:
        lines.append(f"Динамическое наблюдение: {clinical_state}")
    if complaints:
        lines.append(f"Жалобы: {complaints}.")
    if treatment:
        lines.append(f"Принимает: {treatment}.")
    if profile_status:
        lines.append(f"Психический статус: {profile_status}.")
    if correction:
        lines.append(correction)

    lines.append("Продолжение лечения по листу нетрудоспособности.")
    lines.extend(dynamic_epicrisis_signature_lines(data.treating_physician, data.department_head))
    return "\n".join(lines)


def _leading_date(text: str) -> date | None:
    match = _LEADING_DATE_RE.match(str(text or ""))
    if not match:
        return None
    year = int(match.group(3))
    if year < 100:
        year += 2000 if year < 70 else 1900
    try:
        return date(year, int(match.group(2)), int(match.group(1)))
    except ValueError:
        return None


def _remove_regular_diary_block_for_date(doc: Document, item_date: date) -> None:
    """Remove the ordinary diary block on a dynamic-epicrisis date.

    Historical diary behavior is mutually exclusive per date: when a dynamic
    epicrisis is due, that calendar date contains the epicrisis instead of the
    ordinary diary/joint-round block. A rendered ordinary entry is one dated
    paragraph followed by its undated signature paragraphs, so remove the whole
    span up to the first later dated entry.
    """
    paragraphs = list(doc.paragraphs)
    start = next(
        (
            index
            for index, paragraph in enumerate(paragraphs)
            if _leading_date(paragraph.text) == item_date
            and "Динамический эпикриз." not in paragraph.text
        ),
        None,
    )
    if start is None:
        return

    end = len(paragraphs)
    for index in range(start + 1, len(paragraphs)):
        paragraph_date = _leading_date(paragraphs[index].text)
        if paragraph_date is not None and paragraph_date > item_date:
            end = index
            break

    for paragraph in paragraphs[start:end]:
        element = paragraph._element
        parent = element.getparent()
        if parent is not None:
            parent.remove(element)


def _dynamic_epicrisis_rendered_paragraphs(
    item_date: date,
    lines: tuple[str, ...],
) -> tuple[str, ...]:
    """Render clinical epicrisis prose as one paragraph, like an ordinary diary.

    build_dynamic_epicrisis_text keeps semantic fields on separate lines so
    callers can still inspect them independently. Word output must not mirror
    those semantic line breaks: the clinical body is continuous prose and Word
    should wrap it naturally at the page margin. Signature lines remain
    separate, matching ordinary diary rendering.
    """
    signature_prefixes = ("Лечащий врач ", "Зав.отделением ")
    body_lines: list[str] = []
    signature_lines: list[str] = []
    for raw_line in lines:
        line = " ".join(str(raw_line or "").split())
        if not line:
            continue
        if line.startswith(signature_prefixes):
            signature_lines.append(line)
        else:
            body_lines.append(line)

    body = " ".join(body_lines).strip()
    dated_body = f"{item_date:%d.%m.%y} {body}".strip()
    return (dated_body, *signature_lines)


def _insert_dynamic_block(doc: Document, item_date: date, lines: tuple[str, ...]) -> None:
    """Insert a dynamic epicrisis in chronological order."""
    anchor = None
    for paragraph in doc.paragraphs:
        paragraph_date = _leading_date(paragraph.text)
        if paragraph_date is not None and paragraph_date > item_date:
            anchor = paragraph
            break

    rendered = _dynamic_epicrisis_rendered_paragraphs(item_date, lines)
    if anchor is None:
        if doc.paragraphs:
            doc.add_paragraph("")
        for line in rendered:
            paragraph = doc.add_paragraph(line)
            if line.startswith("Лечащий врач ") or line.startswith("Зав.отделением "):
                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        return

    for line in rendered:
        paragraph = anchor.insert_paragraph_before(line)
        if line.startswith("Лечащий врач ") or line.startswith("Зав.отделением "):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    anchor.insert_paragraph_before("")


def apply_sick_leave_dynamic_epicrises(
    path: str | Path,
    *,
    admission_value: str,
    discharge_value: str = "",
    sick_leave_from: str = "",
    patient_name: str = "",
    birth_date: str = "",
    complaints: str = "",
    treatment: str = "",
    profile_status: str = "",
    treatment_correction: str = "",
    treating_physician: str = "",
    department_head: str = "",
    clinical_state_by_date: Mapping[date, str] | None = None,
) -> int:
    """Atomically add the historical sick-leave epicrisis blocks to one DOCX."""
    target = Path(path)
    admission = parse_full_date(admission_value)
    sick_leave_text = str(sick_leave_from or "").strip()
    if not sick_leave_text:
        raise ValueError(
            "Для динамических эпикризов по листу нетрудоспособности укажите дату начала больничного."
        )
    try:
        sick_leave_date = parse_full_date(sick_leave_text)
    except ValueError as exc:
        raise ValueError(
            "Дата начала больничного для динамических эпикризов указана неверно."
        ) from exc
    discharge = parse_optional_discharge_date(discharge_value)
    base_date = max(admission, sick_leave_date)
    dates = dynamic_epicrisis_dates(base_date, discharge_date=discharge, limit=12)
    if not dates:
        return 0

    # A mapping (even an empty one) means the automatic production route is
    # enforcing date-specific evidence. Legacy undated clinical arguments are
    # retained only for explicit low-level callers that do not provide a map.
    use_legacy_undated = clinical_state_by_date is None
    state_by_date = clinical_state_by_date or {}

    doc = Document(str(target))
    for item_date in dates:
        data = DynamicEpicrisisInput(
            patient_name=patient_name,
            birth_date=birth_date,
            sick_leave_from=f"{base_date:%d.%m.%Y}",
            clinical_state=state_by_date.get(item_date, ""),
            complaints=complaints if use_legacy_undated else "",
            treatment=treatment if use_legacy_undated else "",
            profile_status=profile_status if use_legacy_undated else "",
            treatment_correction=treatment_correction if use_legacy_undated else "",
            treating_physician=treating_physician,
            department_head=department_head,
        )
        lines = tuple(build_dynamic_epicrisis_text(data).splitlines())
        _remove_regular_diary_block_for_date(doc, item_date)
        _insert_dynamic_block(doc, item_date, lines)

    staged = target.with_name(f".{target.name}.dynamic-epicrisis.tmp.docx")
    try:
        doc.save(str(staged))
        os.replace(staged, target)
    finally:
        try:
            staged.unlink()
        except FileNotFoundError:
            pass
    return len(dates)


class DiaryService:
    """Create the paragraph-based clinical diary used by the GUI.

    The legacy table writer intentionally is not exposed here. This keeps the
    production UI on one diary architecture while old integrations can continue
    using ``fill_diary_batch`` during the compatibility period.
    """

    @staticmethod
    def _make_fallback_date_source(
        directory: Path,
        *,
        admission_value: str,
        discharge_value: str,
    ) -> Path:
        """Create an ephemeral date source when no 01–31 set was supplied.

        A manually selected diary-text Word file is the doctor's explicit source
        of clinical wording. Numbered date files are optional calendar metadata;
        when absent (or when a remembered folder contains only text files), use
        the proven clinical cadence and never reject the selected text source.
        The temporary DOCX is deleted immediately after generation.
        """
        admission = parse_full_date(admission_value)
        discharge = parse_optional_discharge_date(discharge_value)
        max_offset = (discharge - admission).days if discharge is not None else 7
        path = directory / "generated-clinical-diary-dates.docx"
        doc = Document()
        for offset in _clinical_diary_offsets(max_offset):
            doc.add_paragraph((admission + timedelta(days=offset)).strftime("%d.%m.%Y"))
        doc.save(str(path))
        return path

    @staticmethod
    def _persistent_fallback_output_dir(
        output_dir: str | Path | None,
        status_files: Sequence[str | Path],
    ) -> str | Path:
        """Choose a persistent destination before creating the temporary Dates source."""
        if output_dir is not None and str(output_dir).strip():
            return output_dir
        for raw_path in status_files:
            if raw_path is not None and str(raw_path).strip():
                return Path(raw_path).expanduser().parent
        return Path.cwd()

    @staticmethod
    def _add_dynamic_epicrises_if_needed(
        result: DiaryBatchResult,
        *,
        sick_leave_dynamic_epicrisis: bool,
        admission_value: str,
        discharge_value: str,
        sick_leave_from: str,
        patient_name: str,
        birth_date: str,
        complaints: str,
        treatment: str,
        profile_status: str,
        treatment_correction: str,
        doctor_name: str,
        department_head_name: str,
    ) -> DiaryBatchResult:
        # The established diary generator is intentionally untouched when sick
        # leave is off. Dynamic epicrises are an additive historical layer.
        if not sick_leave_dynamic_epicrisis:
            return result

        total = 0
        for created in result.created_files:
            total += apply_sick_leave_dynamic_epicrises(
                created,
                admission_value=admission_value,
                discharge_value=discharge_value,
                sick_leave_from=sick_leave_from,
                patient_name=patient_name,
                birth_date=birth_date,
                complaints=complaints,
                treatment=treatment,
                profile_status=profile_status,
                # Only exact-date semantic diary observations may supply the
                # clinical body of an automatic dynamic epicrisis.
                clinical_state_by_date=result.dated_clinical_states,
                treating_physician=doctor_name,
                department_head=department_head_name,
            )
        # Kept outside the dataclass for compatibility with existing consumers.
        result.dynamic_epicrisis_count = total
        return result

    def create_text_diaries(
        self,
        *,
        status_files: Sequence[str | Path],
        diary_files: Sequence[str | Path],
        output_dir: str | Path | None,
        patient_name: str,
        admission_value: str,
        gender_source_name: str | None = None,
        discharge_value: str = "",
        repeat_statuses: bool = True,
        force_final_diary: bool = True,
        write_report: bool = False,
        doctor_name: str = "",
        department_head_name: str = "",
        sick_leave_dynamic_epicrisis: bool = False,
        sick_leave_from: str = "",
        birth_date: str = "",
        complaints: str = "",
        treatment: str = "",
        profile_status: str = "",
        treatment_correction: str = "",
    ) -> DiaryBatchResult:
        # Validate the dynamic-epicrisis contract before the underlying diary
        # generator commits any visible DOCX. A bad/missing sick-leave date must
        # never leave a partial kit behind.
        if sick_leave_dynamic_epicrisis:
            sick_leave_text = str(sick_leave_from or "").strip()
            if not sick_leave_text:
                raise ValueError(
                    "Для динамических эпикризов по листу нетрудоспособности укажите дату начала больничного."
                )
            try:
                parse_full_date(sick_leave_text)
            except ValueError as exc:
                raise ValueError(
                    "Дата начала больничного для динамических эпикризов указана неверно."
                ) from exc

        def finalize(result: DiaryBatchResult) -> DiaryBatchResult:
            return self._add_dynamic_epicrises_if_needed(
                result,
                sick_leave_dynamic_epicrisis=sick_leave_dynamic_epicrisis,
                admission_value=admission_value,
                discharge_value=discharge_value,
                sick_leave_from=sick_leave_from,
                patient_name=patient_name,
                birth_date=birth_date,
                complaints=complaints,
                treatment=treatment,
                profile_status=profile_status,
                treatment_correction=treatment_correction,
                doctor_name=doctor_name,
                department_head_name=department_head_name,
            )

        if diary_files:
            return finalize(
                create_text_diaries(
                    status_files=status_files,
                    diary_files=diary_files,
                    output_dir=output_dir,
                    patient_name=patient_name,
                    admission_value=admission_value,
                    gender_source_name=gender_source_name,
                    discharge_value=discharge_value,
                    repeat_statuses=repeat_statuses,
                    force_final_diary=force_final_diary,
                    write_report=write_report,
                    doctor_name=doctor_name,
                    department_head_name=department_head_name,
                )
            )

        # No numbered date template: the explicitly selected text file remains
        # authoritative. Resolve output against persistent inputs before entering
        # the temporary date-source directory, otherwise output_dir=None would
        # write the finished diary into that directory and delete it on return.
        persistent_output_dir = self._persistent_fallback_output_dir(output_dir, status_files)
        with TemporaryDirectory(prefix=".diary-date-source-") as tmp_dir:
            fallback = self._make_fallback_date_source(
                Path(tmp_dir),
                admission_value=admission_value,
                discharge_value=discharge_value,
            )
            return finalize(
                create_text_diaries(
                    status_files=status_files,
                    diary_files=[fallback],
                    output_dir=persistent_output_dir,
                    patient_name=patient_name,
                    admission_value=admission_value,
                    gender_source_name=gender_source_name,
                    discharge_value=discharge_value,
                    repeat_statuses=repeat_statuses,
                    force_final_diary=force_final_diary,
                    write_report=write_report,
                    doctor_name=doctor_name,
                    department_head_name=department_head_name,
                )
            )
