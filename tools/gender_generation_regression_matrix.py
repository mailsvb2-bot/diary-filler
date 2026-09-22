"""Differential regression matrix for gender-sensitive DOCX generation.

This is a test-only safety net for performance work. It compares the optimized
production adapter against the historical exhaustive algorithm, first on a
synthetic run-format corpus and then across every canonical medical document
kind for both male and female patients.
"""
from __future__ import annotations

import copy
import re
import tempfile
from pathlib import Path

from docx import Document

import medical_gender
from generation_performance_profile import _make_fixture
from golden_docx_regression import docx_fingerprint
from medical_constants import DOCUMENT_ORDER
from medical_docx_editor import iter_all_paragraphs, replace_paragraph_regex_preserving_runs
from medical_models import PatientData
from medical_text_utils import normalize_match
from shared_gender import GENDER_WORD_PAIRS


def _preserve_case(source: str, target: str) -> str:
    if source.isupper():
        return target.upper()
    if source[:1].isupper():
        return target[:1].upper() + target[1:]
    return target


def _reference_adapt_document_to_patient_gender(doc, data: PatientData) -> None:
    """Historical pre-optimization algorithm kept only inside this regression test."""
    gender = medical_gender.patient_gender(data)
    if gender not in {"male", "female"}:
        return

    pairs = sorted(GENDER_WORD_PAIRS, key=lambda pair: max(len(pair[0]), len(pair[1])), reverse=True)
    for paragraph in list(iter_all_paragraphs(doc)):
        original = paragraph.text or ""
        if not original.strip():
            continue
        if "диагноз" in normalize_match(original):
            continue
        for male, female in pairs:
            source, target = (female, male) if gender == "male" else (male, female)
            pattern = re.compile(
                rf"(?<![A-Za-zА-Яа-яЁё]){re.escape(source)}(?![A-Za-zА-Яа-яЁё])",
                re.IGNORECASE,
            )
            replace_paragraph_regex_preserving_runs(
                paragraph,
                pattern,
                lambda match, target=target: _preserve_case(match.group(0), target),
            )


def _source_target(pair: tuple[str, str], gender: str) -> tuple[str, str]:
    male, female = pair
    return (female, male) if gender == "male" else (male, female)


def _add_split_case(doc, label: str, source: str, must_change: list[tuple[int, str]]) -> None:
    paragraph = doc.add_paragraph()
    prefix = f"{label} START "
    suffix = " END"
    paragraph.add_run(prefix)
    split = max(1, len(source) // 2)
    first = paragraph.add_run(source[:split])
    first.bold = True
    second = paragraph.add_run(source[split:])
    second.italic = True
    tail = paragraph.add_run(suffix)
    tail.underline = True
    must_change.append((len(doc.paragraphs) - 1, prefix + source + suffix))


def _build_synthetic_document(gender: str):
    doc = Document()
    must_change: list[tuple[int, str]] = []
    must_stay: list[tuple[int, str]] = []
    ordered_pairs = sorted(GENDER_WORD_PAIRS, key=lambda pair: max(len(pair[0]), len(pair[1])), reverse=True)
    seen_sources: set[str] = set()

    # Every historically reachable source form must survive a cross-run
    # replacement. Duplicate source forms intentionally keep only one synthetic
    # case because the historical ordered algorithm itself determines which of
    # multiple е/ё targets wins. Exact parity is checked against that algorithm.
    for index, pair in enumerate(ordered_pairs):
        source, _target = _source_target(pair, gender)
        source_key = source.casefold()
        if source_key in seen_sources:
            continue
        seen_sources.add(source_key)
        _add_split_case(doc, f"PAIR-{index}", source, must_change)

        # Representative boundary-negative cases prove the cheap containment
        # prefilter cannot turn a non-match into a match.
        if index % 11 == 0:
            text = f"BOUND-{index} x{source}y"
            doc.add_paragraph(text)
            must_stay.append((len(doc.paragraphs) - 1, text))

        # Representative case-preservation cases exercise title/uppercase forms.
        if index % 13 == 0:
            for variant_name, variant in (("TITLE", source[:1].upper() + source[1:]), ("UPPER", source.upper())):
                _add_split_case(doc, f"{variant_name}-{index}", variant, must_change)

    # One dense paragraph includes every configured source, including duplicates,
    # and therefore exercises exact sequential/cascading order across the full rule set.
    dense_sources = [_source_target(pair, gender)[0] for pair in ordered_pairs]
    dense_text = "DENSE " + " | ".join(dense_sources)
    doc.add_paragraph(dense_text)
    must_change.append((len(doc.paragraphs) - 1, dense_text))

    # Diagnosis paragraphs are intentionally protected from gender rewriting.
    diagnosis_text = "Диагноз: " + " | ".join(dense_sources)
    doc.add_paragraph(diagnosis_text)
    must_stay.append((len(doc.paragraphs) - 1, diagnosis_text))
    return doc, must_change, must_stay


def _verify_synthetic_contract(doc, must_change: list[tuple[int, str]], must_stay: list[tuple[int, str]], gender: str) -> None:
    for index, original in must_change:
        actual = doc.paragraphs[index].text
        if actual == original:
            raise SystemExit(
                f"GENDER MATRIX FAILED: expected a historical gender replacement for {gender} paragraph {index}"
            )
    for index, expected in must_stay:
        actual = doc.paragraphs[index].text
        if actual != expected:
            raise SystemExit(
                "GENDER MATRIX FAILED: protected/boundary text changed for "
                f"{gender} paragraph {index}: expected {expected!r}, actual {actual!r}"
            )


def _verify_synthetic_parity(root: Path) -> int:
    checked = 0
    for gender, fio in (("female", "Маркер Женская Тестовая"), ("male", "Маркер Мужской Дополнительный")):
        optimized, must_change, must_stay = _build_synthetic_document(gender)
        reference, _, _ = _build_synthetic_document(gender)
        data = PatientData(fio=fio, output_fio=fio)

        medical_gender.adapt_document_to_patient_gender(optimized, data)
        _reference_adapt_document_to_patient_gender(reference, data)
        _verify_synthetic_contract(optimized, must_change, must_stay, gender)

        optimized_path = root / f"synthetic-{gender}-optimized.docx"
        reference_path = root / f"synthetic-{gender}-reference.docx"
        optimized.save(optimized_path)
        reference.save(reference_path)
        if docx_fingerprint(optimized_path) != docx_fingerprint(reference_path):
            raise SystemExit(f"GENDER MATRIX FAILED: optimized/reference run structure differs for {gender}")
        checked += len(must_change) + len(must_stay)
    return checked


def _document_text(path: Path) -> str:
    doc = Document(path)
    chunks = [paragraph.text for paragraph in iter_all_paragraphs(doc)]
    return "\n".join(chunks)


def _has_gender_target(text: str, gender: str) -> bool:
    for pair in GENDER_WORD_PAIRS:
        _source, target = _source_target(pair, gender)
        if re.search(
            rf"(?<![A-Za-zА-Яа-яЁё]){re.escape(target)}(?![A-Za-zА-Яа-яЁё])",
            text,
            re.IGNORECASE,
        ):
            return True
    return False


def _generate_with_adapter(service, nav: Path, data: PatientData, output_dir: Path, adapter):
    original = medical_gender.adapt_document_to_patient_gender
    medical_gender.adapt_document_to_patient_gender = adapter
    try:
        created, _ = service.create_documents(
            navigation_path=nav,
            output_dir=output_dir,
            selected_docs=DOCUMENT_ORDER,
            override_data=copy.deepcopy(data),
        )
        return created
    finally:
        medical_gender.adapt_document_to_patient_gender = original


def _verify_full_generation_parity(root: Path) -> int:
    nav, service, baseline = _make_fixture(root)
    compared = 0
    for gender, fio in (("female", "Маркер Женская Тестовая"), ("male", "Маркер Мужской Дополнительный")):
        data = copy.deepcopy(baseline)
        data.fio = fio
        data.output_fio = fio

        optimized_dir = root / f"generated-{gender}-optimized"
        reference_dir = root / f"generated-{gender}-reference"
        optimized = _generate_with_adapter(
            service, nav, data, optimized_dir, medical_gender.adapt_document_to_patient_gender
        )
        reference = _generate_with_adapter(
            service, nav, data, reference_dir, _reference_adapt_document_to_patient_gender
        )

        if len(optimized) != len(DOCUMENT_ORDER) or len(reference) != len(DOCUMENT_ORDER):
            raise SystemExit(
                f"GENDER MATRIX FAILED: expected {len(DOCUMENT_ORDER)} documents for {gender}, "
                f"got optimized={len(optimized)}, reference={len(reference)}"
            )

        combined_text: list[str] = []
        for kind, optimized_path, reference_path in zip(DOCUMENT_ORDER, optimized, reference):
            if optimized_path.name != reference_path.name:
                raise SystemExit(
                    f"GENDER MATRIX FAILED: output filename drift for {gender}/{kind}: "
                    f"{optimized_path.name!r} != {reference_path.name!r}"
                )
            if docx_fingerprint(optimized_path) != docx_fingerprint(reference_path):
                raise SystemExit(f"GENDER MATRIX FAILED: generated DOCX differs for {gender}/{kind}")
            combined_text.append(_document_text(optimized_path))
            compared += 1

        if not _has_gender_target("\n".join(combined_text), gender):
            raise SystemExit(f"GENDER MATRIX FAILED: generated {gender} corpus contains no gender-adapted target")
    return compared


def main() -> None:
    with tempfile.TemporaryDirectory(prefix="medical-autofill-gender-matrix-") as tmp:
        root = Path(tmp)
        synthetic_cases = _verify_synthetic_parity(root)
        document_pairs = _verify_full_generation_parity(root)
    print(
        "GENDER GENERATION REGRESSION MATRIX OK: "
        f"pairs={len(GENDER_WORD_PAIRS)}; genders=2; synthetic_cases={synthetic_cases}; "
        f"generated_document_pairs={document_pairs}; document_kinds={len(DOCUMENT_ORDER)}"
    )


if __name__ == "__main__":
    main()
