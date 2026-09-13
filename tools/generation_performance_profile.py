from __future__ import annotations

import argparse
import copy
import json
import statistics
import tempfile
import time
from collections import defaultdict
from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocxDocument

import medical_renderer_commission
import medical_renderer_labs
import medical_renderer_primary
import medical_renderer_special
import medical_service
from medical_constants import DOCUMENT_ORDER
from medical_docx_editor_core import DocxEditorCoreMixin
from medical_docx_editor_markers import DocxEditorMarkersMixin
from medical_service import MedicalDocumentService


RENDERER_MODULES = (
    medical_renderer_primary,
    medical_renderer_commission,
    medical_renderer_special,
    medical_renderer_labs,
)


def _ms(seconds: float) -> float:
    return round(seconds * 1000.0, 3)


def _make_fixture(root: Path):
    nav = root / "Направление_профиль.docx"
    nav_doc = Document()
    nav_doc.add_paragraph("10.06.2026 Первичный осмотр")
    nav_doc.add_paragraph("История болезни № 123")
    nav_doc.add_paragraph("Ф.И.О.: Иванова Ирина Ивановна")
    nav_doc.add_paragraph("Год рождения: 1980")
    nav_doc.add_paragraph("Зарегистрирован: Н. Новгород, тестовый район")
    nav_doc.add_paragraph("На учёте у психиатров: не состоит")
    nav_doc.add_paragraph("Направление от РВК: нет")
    nav_doc.add_paragraph("Работает в организации: не работает")
    nav_doc.add_paragraph("Должность: не работает")
    nav_doc.add_paragraph("В 3 отделение КДП поступает добровольно")
    nav_doc.add_paragraph("Жалобы на момент осмотра: тревога, нарушение сна")
    nav_doc.add_paragraph("Анамнез жизни: Со слов пациентки, росла и развивалась без особенностей.")
    nav_doc.add_paragraph("Анамнез заболевания: Ухудшение состояния в течение месяца.")
    nav_doc.add_paragraph("Психический статус: В сознании, ориентирована, контакт доступен.")
    nav_doc.add_paragraph("Соматический статус: Нормального питания, без грубой соматической патологии.")
    nav_doc.add_paragraph("План обследования: стандартный план")
    nav_doc.add_paragraph("План лечения: терапия по назначению врача")
    nav_doc.add_paragraph(
        "На основании данных анамнеза жизни и заболевания, психического статуса, "
        "данных клинических исследований был выставлен диагноз: F тестовый диагноз"
    )
    nav_doc.save(nav)

    service = MedicalDocumentService()
    data = service.parse_navigation(nav)
    data.discharge_date = "11.06.2026"
    data.diagnosis = "F99.9 Тестовый диагноз из UI"
    data.admission_occurrence = "повторно"
    data.rvk_act_number = "77-А"
    data.rvk_military_commissariat = "Ленинского"
    data.rvk_referral_present = "да"
    data.rvk_referral_commissariat = "Ленинского"
    data.rvk_work_position = "ООО РВК, программист"
    data.commission_date = "18.06.2026"
    data.commission_number = "9"
    data.vk_date = "16.06.2026"
    data.vk_protocol_number = "42"
    data.vk_protocol_date = "16.06.2026"
    data.vk_mse_work_org = "ГБУЗ НО Тест"
    data.vk_mse_position = "санитар"
    data.sick_leave_vk_date = "18.06.2026"
    data.sick_leave_vk_protocol_number = "55"
    data.sick_leave_vk_protocol_date = "18.06.2026"
    data.sick_leave_vk_commission_date = "18.06.2026"
    data.sick_leave_vk_work_org = "ООО Тест"
    data.sick_leave_vk_position = "инженер"
    data.sick_leave_vk_work_position = ""
    data.expert_work_status = "да"
    data.expert_work_org = "ООО Завод"
    data.expert_position = "инженер"
    data.expert_sick_leave_needed = "да"
    data.expert_sick_leave_from = "15.06.2026"
    data.disability_needed = "нет"
    data.disability = "не нужно"
    data.work_org = data.expert_work_org
    data.position = data.expert_position
    data.sick_leave = "нужен с 15.06.2026"
    data.epi_present = "нет"
    data.epi_text = ""
    return nav, service, data


class Probe:
    def __init__(self, service: MedicalDocumentService):
        self.service = service
        self.current_kind = ""
        self.render_seconds: dict[str, float] = defaultdict(float)
        self.render_calls: dict[str, int] = defaultdict(int)
        self.load_seconds: dict[str, float] = defaultdict(float)
        self.load_calls: dict[str, int] = defaultdict(int)
        self.save_seconds: dict[str, float] = defaultdict(float)
        self.save_calls: dict[str, int] = defaultdict(int)
        self.paragraph_seconds: dict[str, float] = defaultdict(float)
        self.paragraph_calls: dict[str, int] = defaultdict(int)
        self.find_seconds: dict[str, float] = defaultdict(float)
        self.find_calls: dict[str, int] = defaultdict(int)
        self.next_find_seconds: dict[str, float] = defaultdict(float)
        self.next_find_calls: dict[str, int] = defaultdict(int)
        self.replace_seconds = 0.0
        self.replace_calls = 0
        self._restorers: list[Any] = []

    def __enter__(self):
        original_render = self.service.renderer.render

        def timed_render(kind, template_path, output_path, data):
            previous = self.current_kind
            self.current_kind = kind
            started = time.perf_counter()
            try:
                return original_render(kind, template_path, output_path, data)
            finally:
                elapsed = time.perf_counter() - started
                self.render_seconds[kind] += elapsed
                self.render_calls[kind] += 1
                self.current_kind = previous

        self.service.renderer.render = timed_render
        self._restorers.append(lambda: setattr(self.service.renderer, "render", original_render))

        for module in RENDERER_MODULES:
            if not hasattr(module, "Document"):
                continue
            original_factory = module.Document

            def timed_factory(*args, __original=original_factory, **kwargs):
                started = time.perf_counter()
                try:
                    return __original(*args, **kwargs)
                finally:
                    kind = self.current_kind or "unknown"
                    self.load_seconds[kind] += time.perf_counter() - started
                    self.load_calls[kind] += 1

            module.Document = timed_factory
            self._restorers.append(lambda m=module, f=original_factory: setattr(m, "Document", f))

        original_save = DocxDocument.save

        def timed_save(document, *args, **kwargs):
            started = time.perf_counter()
            try:
                return original_save(document, *args, **kwargs)
            finally:
                kind = self.current_kind or "unknown"
                self.save_seconds[kind] += time.perf_counter() - started
                self.save_calls[kind] += 1

        DocxDocument.save = timed_save
        self._restorers.append(lambda: setattr(DocxDocument, "save", original_save))

        original_paragraphs = DocxEditorCoreMixin.paragraphs

        def timed_paragraphs(editor):
            started = time.perf_counter()
            try:
                return original_paragraphs.fget(editor)
            finally:
                kind = self.current_kind or "unknown"
                self.paragraph_seconds[kind] += time.perf_counter() - started
                self.paragraph_calls[kind] += 1

        DocxEditorCoreMixin.paragraphs = property(timed_paragraphs)
        self._restorers.append(lambda: setattr(DocxEditorCoreMixin, "paragraphs", original_paragraphs))

        original_find = DocxEditorMarkersMixin.find_paragraph_index

        def timed_find(editor, *args, **kwargs):
            started = time.perf_counter()
            try:
                return original_find(editor, *args, **kwargs)
            finally:
                kind = self.current_kind or "unknown"
                self.find_seconds[kind] += time.perf_counter() - started
                self.find_calls[kind] += 1

        DocxEditorMarkersMixin.find_paragraph_index = timed_find
        self._restorers.append(lambda: setattr(DocxEditorMarkersMixin, "find_paragraph_index", original_find))

        original_next = DocxEditorMarkersMixin.find_next_marker_index

        def timed_next(editor, *args, **kwargs):
            started = time.perf_counter()
            try:
                return original_next(editor, *args, **kwargs)
            finally:
                kind = self.current_kind or "unknown"
                self.next_find_seconds[kind] += time.perf_counter() - started
                self.next_find_calls[kind] += 1

        DocxEditorMarkersMixin.find_next_marker_index = timed_next
        self._restorers.append(lambda: setattr(DocxEditorMarkersMixin, "find_next_marker_index", original_next))

        original_replace = medical_service.os.replace

        def timed_replace(*args, **kwargs):
            started = time.perf_counter()
            try:
                return original_replace(*args, **kwargs)
            finally:
                self.replace_seconds += time.perf_counter() - started
                self.replace_calls += 1

        medical_service.os.replace = timed_replace
        self._restorers.append(lambda: setattr(medical_service.os, "replace", original_replace))
        return self

    def __exit__(self, exc_type, exc, tb):
        for restore in reversed(self._restorers):
            restore()

    def snapshot(self, total_seconds: float) -> dict[str, Any]:
        kinds = list(DOCUMENT_ORDER)
        per_kind = {}
        for kind in kinds:
            render = self.render_seconds.get(kind, 0.0)
            load = self.load_seconds.get(kind, 0.0)
            save = self.save_seconds.get(kind, 0.0)
            paragraphs = self.paragraph_seconds.get(kind, 0.0)
            finds = self.find_seconds.get(kind, 0.0) + self.next_find_seconds.get(kind, 0.0)
            per_kind[kind] = {
                "render_ms": _ms(render),
                "docx_load_ms": _ms(load),
                "docx_save_ms": _ms(save),
                "paragraph_list_ms": _ms(paragraphs),
                "marker_search_ms": _ms(finds),
                "paragraph_list_calls": self.paragraph_calls.get(kind, 0),
                "find_paragraph_calls": self.find_calls.get(kind, 0),
                "find_next_marker_calls": self.next_find_calls.get(kind, 0),
            }
        render_total = sum(self.render_seconds.values())
        load_total = sum(self.load_seconds.values())
        save_total = sum(self.save_seconds.values())
        paragraph_total = sum(self.paragraph_seconds.values())
        marker_total = sum(self.find_seconds.values()) + sum(self.next_find_seconds.values())
        return {
            "total_ms": _ms(total_seconds),
            "render_total_ms": _ms(render_total),
            "outside_render_ms": _ms(max(0.0, total_seconds - render_total)),
            "docx_load_total_ms": _ms(load_total),
            "docx_save_total_ms": _ms(save_total),
            "paragraph_list_total_ms": _ms(paragraph_total),
            "marker_search_total_ms": _ms(marker_total),
            "paragraph_list_calls": sum(self.paragraph_calls.values()),
            "find_paragraph_calls": sum(self.find_calls.values()),
            "find_next_marker_calls": sum(self.next_find_calls.values()),
            "os_replace_ms": _ms(self.replace_seconds),
            "os_replace_calls": self.replace_calls,
            "per_kind": per_kind,
        }


def _median(values):
    return round(statistics.median(values), 3)


def _aggregate(runs: list[dict[str, Any]]) -> dict[str, Any]:
    scalar_keys = (
        "total_ms",
        "render_total_ms",
        "outside_render_ms",
        "docx_load_total_ms",
        "docx_save_total_ms",
        "paragraph_list_total_ms",
        "marker_search_total_ms",
        "os_replace_ms",
    )
    result: dict[str, Any] = {f"median_{key}": _median([run[key] for run in runs]) for key in scalar_keys}
    for count_key in ("paragraph_list_calls", "find_paragraph_calls", "find_next_marker_calls", "os_replace_calls"):
        result[f"median_{count_key}"] = _median([run[count_key] for run in runs])
    per_kind = {}
    for kind in DOCUMENT_ORDER:
        per_kind[kind] = {
            "median_render_ms": _median([run["per_kind"][kind]["render_ms"] for run in runs]),
            "median_docx_load_ms": _median([run["per_kind"][kind]["docx_load_ms"] for run in runs]),
            "median_docx_save_ms": _median([run["per_kind"][kind]["docx_save_ms"] for run in runs]),
            "median_paragraph_list_ms": _median([run["per_kind"][kind]["paragraph_list_ms"] for run in runs]),
            "median_marker_search_ms": _median([run["per_kind"][kind]["marker_search_ms"] for run in runs]),
            "median_paragraph_list_calls": _median([run["per_kind"][kind]["paragraph_list_calls"] for run in runs]),
        }
    result["per_kind"] = per_kind
    return result


def run_profile(measured_runs: int) -> dict[str, Any]:
    with tempfile.TemporaryDirectory(prefix="medical-autofill-perf-") as tmp:
        root = Path(tmp)
        nav, service, baseline_data = _make_fixture(root)

        warmup_dir = root / "warmup"
        service.create_documents(
            navigation_path=nav,
            output_dir=warmup_dir,
            selected_docs=DOCUMENT_ORDER,
            override_data=copy.deepcopy(baseline_data),
        )

        runs = []
        for run_index in range(measured_runs):
            output_dir = root / f"run-{run_index + 1}"
            with Probe(service) as probe:
                started = time.perf_counter()
                created, _ = service.create_documents(
                    navigation_path=nav,
                    output_dir=output_dir,
                    selected_docs=DOCUMENT_ORDER,
                    override_data=copy.deepcopy(baseline_data),
                )
                elapsed = time.perf_counter() - started
            if len(created) != len(DOCUMENT_ORDER):
                raise RuntimeError(f"Expected {len(DOCUMENT_ORDER)} documents, got {len(created)}")
            runs.append(probe.snapshot(elapsed))

    return {
        "schema_version": 1,
        "measured_runs": measured_runs,
        "documents_per_run": len(DOCUMENT_ORDER),
        "document_order": list(DOCUMENT_ORDER),
        "runs": runs,
        "median": _aggregate(runs),
    }


def _print_summary(report: dict[str, Any]) -> None:
    median = report["median"]
    print("GENERATION PERFORMANCE PROFILE")
    print(f"measured runs: {report['measured_runs']}")
    print(f"documents/run: {report['documents_per_run']}")
    print(f"median total: {median['median_total_ms']:.3f} ms")
    print(f"median render total: {median['median_render_total_ms']:.3f} ms")
    print(f"median outside renderer: {median['median_outside_render_ms']:.3f} ms")
    print(f"median DOCX load: {median['median_docx_load_total_ms']:.3f} ms")
    print(f"median DOCX save: {median['median_docx_save_total_ms']:.3f} ms")
    print(f"median paragraph-list materialization: {median['median_paragraph_list_total_ms']:.3f} ms")
    print(f"median marker search: {median['median_marker_search_total_ms']:.3f} ms")
    print(f"median os.replace: {median['median_os_replace_ms']:.3f} ms")
    print(f"median paragraph-list calls: {median['median_paragraph_list_calls']:.0f}")
    print(f"median find_paragraph calls: {median['median_find_paragraph_calls']:.0f}")
    print(f"median find_next_marker calls: {median['median_find_next_marker_calls']:.0f}")
    print("per document:")
    for kind in report["document_order"]:
        row = median["per_kind"][kind]
        print(
            f"  {kind}: render={row['median_render_ms']:.3f} ms; "
            f"load={row['median_docx_load_ms']:.3f}; save={row['median_docx_save_ms']:.3f}; "
            f"paragraph-lists={row['median_paragraph_list_ms']:.3f}; "
            f"marker-search={row['median_marker_search_ms']:.3f}; "
            f"paragraph-list-calls={row['median_paragraph_list_calls']:.0f}"
        )
    print("PERF_PROFILE_JSON=" + json.dumps(report, ensure_ascii=False, sort_keys=True))


def main() -> int:
    parser = argparse.ArgumentParser(description="Profile canonical medical DOCX generation without changing it.")
    parser.add_argument("--runs", type=int, default=3)
    parser.add_argument("--json-out", type=Path)
    args = parser.parse_args()
    if args.runs < 1:
        parser.error("--runs must be >= 1")
    report = run_profile(args.runs)
    _print_summary(report)
    if args.json_out:
        args.json_out.write_text(json.dumps(report, ensure_ascii=False, indent=2, sort_keys=True), encoding="utf-8")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
