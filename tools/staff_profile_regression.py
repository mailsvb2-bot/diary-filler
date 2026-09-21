from __future__ import annotations

import json
import tempfile
from pathlib import Path
from types import SimpleNamespace
from tkinter import messagebox, simpledialog

from docx import Document

from actions_creation_orchestrator import ActionsCreationOrchestratorMixin
import actions_creation_orchestrator
from diary_service import DiaryService
from medical_constants import DOCUMENT_ORDER
from medical_docx_reader import extract_docx_text
from medical_formatting import format_staff_instrumental_short_name, format_staff_short_name
from medical_gender import normalize_staff_references_in_document
from medical_models import PatientData
from settings_mixin import SettingsMixin
from tools.generation_performance_profile import _make_fixture


class _SettingsHarness(SettingsMixin):
    pass


class _UnconfiguredGenerationHarness(ActionsCreationOrchestratorMixin):
    def __init__(self) -> None:
        self.root = object()
        self.status = ""
        self.prompt_calls = 0
        self.logs: list[str] = []

    def selected_medical_docs(self):
        return []

    def diaries_selected(self):
        return True

    def _staff_profile_is_configured(self):
        return False

    def _prompt_staff_profile(self, *, first_run=False):
        self.prompt_calls += 1
        return False

    def _set_status(self, text):
        self.status = str(text)

    def _log(self, text):
        self.logs.append(str(text))


def _assert_unconfigured_generation_fails_closed() -> None:
    app = _UnconfiguredGenerationHarness()
    warnings: list[tuple[str, str]] = []
    original_warning = actions_creation_orchestrator.messagebox.showwarning
    try:
        actions_creation_orchestrator.messagebox.showwarning = (
            lambda title, message, **_kwargs: warnings.append((str(title), str(message)))
        )
        app.create_selected_outputs(print_after=False)
    finally:
        actions_creation_orchestrator.messagebox.showwarning = original_warning
    assert app.prompt_calls == 1, app.prompt_calls
    assert app.status == "Создание отменено: укажите сотрудников", app.status
    assert app.logs == [], app.logs
    assert warnings and warnings[-1][0] == "Сотрудники не настроены", warnings


def _assert_fresh_staff_prompt_has_no_foreign_defaults(root: Path) -> None:
    app = _SettingsHarness()
    app.root = object()
    app._settings_path = root / "fresh-settings.json"
    app._settings = {}
    captured_defaults: list[str] = []
    answers = iter([
        "Орлов Олег Олегович",
        "Соколова Светлана Сергеевна",
        "Кузнецова Кира Константиновна",
    ])
    original_askstring = simpledialog.askstring
    original_info = messagebox.showinfo
    original_warning = messagebox.showwarning
    try:
        def fake_askstring(_title, _prompt, *, initialvalue="", **_kwargs):
            captured_defaults.append(str(initialvalue or ""))
            return next(answers)

        simpledialog.askstring = fake_askstring
        messagebox.showinfo = lambda *_args, **_kwargs: None
        messagebox.showwarning = lambda *_args, **_kwargs: None
        assert app._prompt_staff_profile(first_run=True)
    finally:
        simpledialog.askstring = original_askstring
        messagebox.showinfo = original_info
        messagebox.showwarning = original_warning

    assert captured_defaults == ["", "", ""], captured_defaults
    assert app._staff_profile_is_configured()
    assert app._effective_staff_profile()["doctor"] == "Орлов Олег Олегович"


def _assert_persistence_failure_fails_closed(root: Path) -> None:
    blocker = root / "blocked-settings-parent"
    blocker.write_text("not-a-directory", encoding="utf-8")

    app = _SettingsHarness()
    app._settings_path = blocker / "settings.json"
    app._settings = {}

    assert not app._set_staff_profile(
        doctor="Орлов Олег Олегович",
        department_head="Соколова Светлана Сергеевна",
        deputy_chief="Кузнецова Кира Константиновна",
    )
    assert not app._staff_profile_is_configured(), app._settings

    before = app._patient_folder_naming_settings()
    assert not app._set_patient_folder_naming_settings(
        parts=["full_fio", "admission_date"],
        date_format="full",
    )
    assert app._patient_folder_naming_settings() == before, app._settings

    assert app._set_desktop_intake_preference(True) is False


def _assert_settings(root: Path) -> None:
    settings = root / "settings.json"
    app = _SettingsHarness()
    app._settings_path = settings
    app._settings = {}
    assert app._set_staff_profile(
        doctor="Орлов Олег Олегович",
        department_head="Соколова Светлана Сергеевна",
        deputy_chief="Кузнецова Кира Константиновна",
    )
    app._set_desktop_intake_preference(False)
    payload = json.loads(settings.read_text(encoding="utf-8"))
    assert payload["staff_profile"]["doctor"] == "Орлов Олег Олегович"
    assert payload["staff_profile"]["department_head"] == "Соколова Светлана Сергеевна"
    assert payload["staff_profile"]["deputy_chief"] == "Кузнецова Кира Константиновна"
    assert payload["desktop_intake_enabled"] is False

    restored = _SettingsHarness()
    restored._settings_path = settings
    restored._settings = restored._load_settings()
    assert restored._staff_profile_is_configured()
    assert restored._effective_staff_profile() == {
        "doctor": "Орлов Олег Олегович",
        "department_head": "Соколова Светлана Сергеевна",
        "deputy_chief": "Кузнецова Кира Константиновна",
    }
    assert restored._desktop_intake_preference() is False


def _assert_name_formatting() -> None:
    assert format_staff_short_name("Иванов Иван Иванович") == "Иванов И.И."
    assert format_staff_short_name("Можарова Е.А.") == "Можарова Е.А."
    assert format_staff_instrumental_short_name("Иванов Иван Иванович") == "Ивановым И.И."
    assert format_staff_instrumental_short_name("Петрова Анна Сергеевна") == "Петровой А.С."


def _assert_medical_documents(root: Path) -> None:
    nav, service, data = _make_fixture(root)
    data.doctor = "Орлов Олег Олегович"
    data.head = "Соколова Светлана Сергеевна"
    data.deputy_chief = "Кузнецова Кира Константиновна"
    output = root / "medical"
    created, _ = service.create_documents(
        navigation_path=nav,
        output_dir=output,
        discharge_date=data.discharge_date,
        selected_docs=DOCUMENT_ORDER,
        override_data=data,
    )
    assert len(created) == len(DOCUMENT_ORDER)
    texts = {path.name: extract_docx_text(path) for path in created}
    joined = "\n".join(texts.values())
    for stale in ("Балаганин", "Можарова", "Можаровой", "Зуйкова", "Зуйковой"):
        assert stale not in joined, f"legacy staff name leaked: {stale}"

    primary = next(text for name, text in texts.items() if "Первичный осмотр" in name)
    discharge = next(text for name, text in texts.items() if "Выписной эпикриз" in name)
    commission = next(text for name, text in texts.items() if "Совместный осмотр" in name)
    admission = next(text for name, text in texts.items() if "приёмного покоя" in name)
    assert "Первичный осмотр с зав. отд. Соколовой С.С." in primary
    assert "Врач психиатр Орлов О.О." in primary
    assert "Зав. отделением Соколова С.С." in primary
    assert "Зав. отд. Соколова С.С." in discharge and "Врач-психиатр" in discharge and "Орлов О.О." in discharge
    assert "Совместный осмотр с зам глав врача Кузнецовой К.К." in commission
    assert "Врач психиатр Орлов О.О." in admission


def _assert_staff_replacement_scope() -> None:
    doc = Document()
    doc.add_paragraph("Ф.И.О.: Можарова Е.А.")
    doc.add_paragraph("Пациентка Можарова Е.А. сообщает о тревоге; Зуйкова А.А. указана в анамнезе семьи.")
    doc.add_paragraph("Зав. отделением Можарова Е.А.")
    doc.add_paragraph("10.06.2026 Первичный осмотр с зав. отд. Можаровой Е.А.")
    doc.add_paragraph("Председатель ВК Зуйкова А.А.")
    doc.add_paragraph("Врач-психиатр Балаганин С.В.")

    data = PatientData(
        doctor="Орлов Олег Олегович",
        head="Соколова Светлана Сергеевна",
        deputy_chief="Кузнецова Кира Константиновна",
    )
    normalize_staff_references_in_document(doc, data)
    lines = [paragraph.text for paragraph in doc.paragraphs]

    assert lines[0] == "Ф.И.О.: Можарова Е.А.", lines
    assert "Пациентка Можарова Е.А." in lines[1], lines
    assert "Зуйкова А.А. указана в анамнезе семьи" in lines[1], lines
    assert lines[2] == "Зав. отделением Соколова С.С.", lines
    assert "Первичный осмотр с зав. отд. Соколовой С.С." in lines[3], lines
    assert lines[4] == "Председатель ВК Кузнецова К.К.", lines
    assert lines[5] == "Врач-психиатр Орлов О.О.", lines


def _assert_diaries(root: Path) -> None:
    statuses = root / "statuses.docx"
    status_doc = Document()
    status_doc.add_paragraph("Состояние стабильное, контакт доступен, лечение переносит удовлетворительно.")
    status_doc.save(statuses)

    dates = root / "dates.docx"
    dates_doc = Document()
    table = dates_doc.add_table(rows=1, cols=4)
    for index, title in enumerate(("День госпитализации", "Число", "Месяц/Год", "Дневник наблюдения")):
        table.rows[0].cells[index].text = title
    for hospital_day in (1, 2, 3, 4):
        row = table.add_row()
        row.cells[0].text = str(hospital_day)
    dates_doc.save(dates)

    result = DiaryService().create_text_diaries(
        status_files=[statuses],
        diary_files=[dates],
        output_dir=root / "diaries",
        patient_name="Иванова Ирина Ивановна",
        gender_source_name="Иванова Ирина Ивановна",
        admission_value="10.06.2026",
        discharge_value="13.06.2026",
        doctor_name="Орлов Олег Олегович",
        department_head_name="Соколова Светлана Сергеевна",
    )
    assert len(result.created_files) == 1
    text = extract_docx_text(result.created_files[0])
    assert "Лечащий врач Орлов О.О." in text
    assert "Зав.отделением Соколова С.С." in text
    assert "Балаганин" not in text and "Можарова" not in text


def main() -> None:
    with tempfile.TemporaryDirectory(prefix="staff-profile-regression-") as tmp:
        root = Path(tmp)
        _assert_name_formatting()
        _assert_settings(root)
        _assert_persistence_failure_fails_closed(root)
        _assert_fresh_staff_prompt_has_no_foreign_defaults(root)
        _assert_unconfigured_generation_fails_closed()
        _assert_staff_replacement_scope()
        _assert_medical_documents(root)
        _assert_diaries(root)
    print("STAFF PROFILE REGRESSION OK: fail-closed onboarding + persistence failure rollback + blank fresh defaults + settings + 7 medical documents + production diary")


if __name__ == "__main__":
    main()
