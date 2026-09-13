"""Semantic/visual golden fingerprints for DOCX emitted by the canonical smoke corpus.

This module never generates a document. It fingerprints DOCX files already
created by the repository's existing smoke_test.py user-flow corpus.
"""
from __future__ import annotations

import hashlib
import json
from pathlib import Path
from typing import Any

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
MANIFEST = ROOT / "tests" / "golden_docx_manifest.json"

GOLDEN_RELATIVE_PATHS = (
    "medical_with_epi/Иванова Ирина Ивановна Акт для РВК.docx",
    "medical_with_epi/Иванова Ирина Ивановна ВК больничный.docx",
    "medical_with_epi/Иванова Ирина Ивановна ВК на МСЭ.docx",
    "medical_with_epi/Иванова Ирина Ивановна Выписной эпикриз.docx",
    "medical_with_epi/Иванова Ирина Ивановна Осмотр врача приёмного покоя.docx",
    "medical_with_epi/Иванова Ирина Ивановна Первичный осмотр.docx",
    "medical_with_epi/Иванова Ирина Ивановна Совместный осмотр.docx",
    "snapshot_diary_output/Снимок Пациента дневники.docx",
)


def _run_payload(run) -> dict[str, Any]:
    color = run.font.color.rgb
    return {
        "text": run.text,
        "bold": run.bold,
        "italic": run.italic,
        "underline": bool(run.underline) if run.underline is not None else None,
        "size_pt": float(run.font.size.pt) if run.font.size is not None else None,
        "color": str(color) if color is not None else None,
        "name": run.font.name,
    }


def _paragraph_payload(paragraph) -> dict[str, Any]:
    return {
        "text": paragraph.text,
        "style": paragraph.style.name if paragraph.style is not None else None,
        "alignment": int(paragraph.alignment) if paragraph.alignment is not None else None,
        "runs": [_run_payload(run) for run in paragraph.runs],
    }


def _cell_payload(cell) -> dict[str, Any]:
    return {
        "paragraphs": [_paragraph_payload(p) for p in cell.paragraphs],
        "tables": [_table_payload(t) for t in cell.tables],
    }


def _table_payload(table) -> dict[str, Any]:
    return {
        "style": table.style.name if table.style is not None else None,
        "rows": [[_cell_payload(cell) for cell in row.cells] for row in table.rows],
    }


def docx_fingerprint(path: Path) -> str:
    doc = Document(path)
    sections = []
    for section in doc.sections:
        sections.append({
            "top": int(section.top_margin or 0),
            "bottom": int(section.bottom_margin or 0),
            "left": int(section.left_margin or 0),
            "right": int(section.right_margin or 0),
            "orientation": int(section.orientation),
            "page_width": int(section.page_width or 0),
            "page_height": int(section.page_height or 0),
            "header": [_paragraph_payload(p) for p in section.header.paragraphs],
            "footer": [_paragraph_payload(p) for p in section.footer.paragraphs],
        })
    payload = {
        "paragraphs": [_paragraph_payload(p) for p in doc.paragraphs],
        "tables": [_table_payload(t) for t in doc.tables],
        "sections": sections,
    }
    raw = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    return hashlib.sha256(raw).hexdigest()


def current_manifest(smoke_root: Path) -> dict[str, str]:
    result: dict[str, str] = {}
    for relative in GOLDEN_RELATIVE_PATHS:
        path = smoke_root / relative
        if not path.is_file():
            raise SystemExit(f"GOLDEN DOCX FAILED: missing canonical smoke output: {relative}")
        result[relative] = docx_fingerprint(path)
    return result


def verify(smoke_root: Path) -> None:
    if not MANIFEST.is_file():
        raise SystemExit("GOLDEN DOCX FAILED: manifest is missing")
    expected = json.loads(MANIFEST.read_text(encoding="utf-8"))
    actual = current_manifest(smoke_root)
    if expected != actual:
        lines = ["GOLDEN DOCX FAILED: canonical output fingerprint changed"]
        for key in sorted(set(expected) | set(actual)):
            if expected.get(key) != actual.get(key):
                lines.append(f"- {key}: expected {expected.get(key)}, actual {actual.get(key)}")
        raise SystemExit("\n".join(lines))
    print(f"GOLDEN DOCX OK: {len(actual)} canonical documents match")


if __name__ == "__main__":
    verify(ROOT / "test_run_combined")
