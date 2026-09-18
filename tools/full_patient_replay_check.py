"""Replay the repository's existing canonical smoke user flow and lock its outputs.

No alternate generator exists here: this script runs smoke_test.py unchanged,
then inspects the artifacts that the existing flow produced.
"""
from __future__ import annotations

from pathlib import Path
import shutil
import subprocess
import sys

from docx import Document

from docx_block_boundary_regression import verify as verify_docx_block_boundaries
from golden_docx_regression import verify as verify_golden

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "test_run_combined"


def _doc_text(path: Path) -> str:
    doc = Document(path)
    chunks = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                chunks.extend(p.text for p in cell.paragraphs)
    return "\n".join(chunks)


def _require(path: Path) -> Path:
    if not path.is_file():
        raise SystemExit(f"FULL PATIENT REPLAY FAILED: missing {path.relative_to(ROOT)}")
    return path


def verify_user_flow() -> None:
    created = OUT / "user_contract_selection_popup_docx" / "created"
    primary = _require(created / "Петров Пётр Петрович Первичный осмотр.docx")
    discharge = _require(created / "Петров Пётр Петрович Выписной эпикриз.docx")
    primary_text = _doc_text(primary)
    discharge_text = _doc_text(discharge)
    required_primary = (
        "История болезни № К-900",
        "План лечения: терапия из пользовательского popup",
    )
    required_discharge = (
        "Выписной эпикриз № К-900",
        "по 11.06.2026",
        "В 3 отделение КДП поступает первично добровольно",
        "Лечение: терапия из пользовательского popup",
        "Рекомендовано: наблюдение у районного психиатра, приём препаратов",
    )
    for marker in required_primary:
        if marker not in primary_text:
            raise SystemExit(f"FULL PATIENT REPLAY FAILED: primary lost marker {marker!r}")
    for marker in required_discharge:
        if marker not in discharge_text:
            raise SystemExit(f"FULL PATIENT REPLAY FAILED: discharge lost marker {marker!r}")
    if "академического отпуска" in discharge_text.lower():
        raise SystemExit("FULL PATIENT REPLAY FAILED: stale academic-leave recommendation leaked into discharge DOCX")

    commission = _require(
        OUT / "user_contract_commission_popup_date" / "created" / "Петров Пётр Петрович Совместный осмотр.docx"
    )
    commission_doc = Document(commission)
    if not commission_doc.paragraphs or not commission_doc.paragraphs[0].text.startswith("21.06.2026 г. 10:00"):
        raise SystemExit("FULL PATIENT REPLAY FAILED: popup commission date did not reach DOCX header")
    commission_text = _doc_text(commission)
    if "предъявляет жалобы на" in commission_text.lower():
        raise SystemExit("FULL PATIENT REPLAY FAILED: legacy standalone complaint tail leaked into commission DOCX")

    partial = OUT / "user_contract_atomic_rollback" / "created"
    partial_docs = list(partial.glob("*.docx")) if partial.exists() else []
    if len(partial_docs) != 2:
        raise SystemExit(
            "FULL PATIENT REPLAY FAILED: diary-only failure must preserve the two successfully generated medical documents"
        )
    partial_names = {path.name for path in partial_docs}
    if not any("Первичный осмотр" in name for name in partial_names):
        raise SystemExit("FULL PATIENT REPLAY FAILED: partial set lost Primary Inspection")
    if not any("Выписной эпикриз" in name for name in partial_names):
        raise SystemExit("FULL PATIENT REPLAY FAILED: partial set lost Discharge Epicrisis")

    diary = _require(OUT / "snapshot_diary_output" / "Снимок Пациента дневники.docx")
    diary_text = _doc_text(diary)
    for marker in ("Лечащий врач Балаганин С.В.", "Зав.отделением Можарова Е.А."):
        if marker not in diary_text:
            raise SystemExit(f"FULL PATIENT REPLAY FAILED: diary lost canonical signature {marker!r}")

    verify_golden(OUT)
    print("FULL PATIENT REPLAY OK")


def main() -> None:
    verify_docx_block_boundaries()
    if OUT.exists():
        shutil.rmtree(OUT, ignore_errors=True)
    try:
        subprocess.run([sys.executable, "smoke_test.py"], cwd=ROOT, check=True, timeout=240)
        verify_user_flow()
    finally:
        shutil.rmtree(OUT, ignore_errors=True)


if __name__ == "__main__":
    main()
