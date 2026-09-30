"""Live Windows 10/11 GUI E2E for the production Tk application.

This test is intentionally different from the normal headless regression suite:
- it creates the real CombinedMedicalDiaryApp and a real Tk/TkDND root;
- it uses Windows cursor/mouse input, not Tk event_generate, to press visible tiles;
- it exercises the canonical production generation path and inspects real DOCX;
- it verifies the two recently repaired user paths: missing-diagnosis preflight
  and direct "ВК больничный" intent without a separate sick-leave prerequisite.

The workflow must run this only after windows11_live_e2e_preflight.ps1 proves
that the runner is a signed-in, unlocked Windows 11 workstation session.
"""
from __future__ import annotations

import ctypes
from ctypes import wintypes
import json
import os
from pathlib import Path
import threading
import time
from tempfile import TemporaryDirectory

if os.name != "nt":
    raise SystemExit("WINDOWS LIVE GUI E2E FAILED: Windows is required")

from docx import Document

from app_config import DIARY_KIND, DIARY_LABEL
from medical_constants import DOCUMENT_LABELS, DOCUMENT_ORDER


USER32 = ctypes.WinDLL("user32", use_last_error=True)
WM_CLOSE = 0x0010
VK_ESCAPE = 0x1B
KEYEVENTF_KEYUP = 0x0002
MOUSEEVENTF_LEFTDOWN = 0x0002
MOUSEEVENTF_LEFTUP = 0x0004


def _write_trace(artifact_dir: Path, payload: dict) -> None:
    artifact_dir.mkdir(parents=True, exist_ok=True)
    (artifact_dir / "live-gui-e2e.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


def _make_primary(path: Path) -> None:
    doc = Document()
    for line in (
        "10.06.2026 Первичный осмотр",
        "История болезни № LIVE-E2E-001",
        "Ф.И.О.: Маркер Мужской Лайв",
        "Год рождения: 1977",
        "Зарегистрирован: Н. Новгород, ул. Проверочная, д. 1",
        "Работает в организации: не работает",
        "В 3 отделение КДП поступает первично добровольно",
        "Жалобы на момент осмотра: тревога, сниженное настроение.",
        "Анамнез жизни: Со слов пациента, рос и развивался без особенностей.",
        "Анамнез заболевания: Ухудшение состояния около двух недель.",
        "Психический статус: Контактен, ориентирован, эмоционально напряжён.",
        "Соматический статус: Без грубой соматической патологии.",
        "Назначенное лечение: терапия по live E2E схеме.",
        (
            "На основании данных анамнеза жизни и заболевания, психического статуса, "
            "данных клинических исследований был выставлен диагноз: "
            "F41.2 Смешанное тревожное и депрессивное расстройство"
        ),
    ):
        doc.add_paragraph(line)
    doc.save(path)


def _make_diary_inputs(root: Path) -> tuple[Path, Path]:
    texts = root / "дневники тревожно депрессивное расстройство.docx"
    text_doc = Document()
    text_doc.add_paragraph("Пациент спокоен, жалоб активно не предъявляет, в беседе доступен.")
    text_doc.add_paragraph("Пациент сообщает об улучшении сна, поведение упорядоченное.")
    text_doc.add_paragraph("Пациент доступен продуктивному контакту, фон настроения ровный.")
    text_doc.save(texts)

    dates = root / "10.docx"
    dates_doc = Document()
    table = dates_doc.add_table(rows=1, cols=4)
    for index, header in enumerate(("День госпитализации", "Число", "Месяц/Год", "Дневник наблюдения")):
        table.rows[0].cells[index].text = header
    for hospital_day in (1, 2, 3, 7):
        row = table.add_row()
        row.cells[0].text = str(hospital_day)
        row.cells[3].text = "Лечащий врач Балаганин С.В.\nЗав.отделением Можарова Е.А."
    dates_doc.save(dates)
    return texts, dates


def _walk_widgets(widget):
    yield widget
    for child in widget.winfo_children():
        yield from _walk_widgets(child)


def _canvas_texts(canvas) -> list[str]:
    values: list[str] = []
    try:
        for item in canvas.find_all():
            if canvas.type(item) == "text":
                value = str(canvas.itemcget(item, "text") or "")
                if value:
                    values.append(value)
    except Exception:
        return []
    return values


def _find_canvas(root, needle: str):
    import tkinter as tk

    candidates = []
    for widget in _walk_widgets(root):
        if not isinstance(widget, tk.Canvas):
            continue
        texts = _canvas_texts(widget)
        if any(needle in value for value in texts):
            candidates.append((widget, texts))
    mapped = [pair for pair in candidates if pair[0].winfo_ismapped()]
    if len(mapped) != 1:
        raise AssertionError(f"visible Canvas for {needle!r}: expected 1, got {len(mapped)}; {mapped!r}")
    return mapped[0][0]


def _foreground(root) -> None:
    root.deiconify()
    root.lift()
    root.focus_force()
    root.update_idletasks()
    hwnd = int(root.winfo_id())
    USER32.SetForegroundWindow(wintypes.HWND(hwnd))
    time.sleep(0.08)


def _physical_click(root, widget) -> None:
    _foreground(root)
    root.update_idletasks()
    x = int(widget.winfo_rootx() + max(4, widget.winfo_width() // 2))
    y = int(widget.winfo_rooty() + max(4, widget.winfo_height() // 2))
    if not USER32.SetCursorPos(x, y):
        raise ctypes.WinError(ctypes.get_last_error())
    time.sleep(0.05)
    USER32.mouse_event(MOUSEEVENTF_LEFTDOWN, 0, 0, 0, 0)
    time.sleep(0.04)
    USER32.mouse_event(MOUSEEVENTF_LEFTUP, 0, 0, 0, 0)
    time.sleep(0.08)


def _top_windows_for_pid(pid: int) -> list[tuple[int, str]]:
    found: list[tuple[int, str]] = []
    enum_proc = ctypes.WINFUNCTYPE(wintypes.BOOL, wintypes.HWND, wintypes.LPARAM)

    @enum_proc
    def callback(hwnd, _lparam):
        owner_pid = wintypes.DWORD()
        USER32.GetWindowThreadProcessId(hwnd, ctypes.byref(owner_pid))
        if owner_pid.value != pid or not USER32.IsWindowVisible(hwnd):
            return True
        length = USER32.GetWindowTextLengthW(hwnd)
        if length:
            buf = ctypes.create_unicode_buffer(length + 1)
            USER32.GetWindowTextW(hwnd, buf, length + 1)
            found.append((int(hwnd), buf.value))
        return True

    USER32.EnumWindows(callback, 0)
    return found


class _PopupEscape:
    def __init__(self, *title_fragments: str):
        self.fragments = tuple(title_fragments)
        self.detected_title = ""
        self._stop = threading.Event()
        self._thread = threading.Thread(target=self._run, name="live-e2e-popup-escape", daemon=True)

    def __enter__(self):
        self._thread.start()
        return self

    def __exit__(self, exc_type, exc, tb):
        self._stop.set()
        self._thread.join(timeout=2)

    def _run(self) -> None:
        deadline = time.monotonic() + 15
        while not self._stop.is_set() and time.monotonic() < deadline:
            for hwnd, title in _top_windows_for_pid(os.getpid()):
                if any(fragment in title for fragment in self.fragments):
                    self.detected_title = title
                    USER32.SetForegroundWindow(wintypes.HWND(hwnd))
                    time.sleep(0.08)
                    USER32.keybd_event(VK_ESCAPE, 0, 0, 0)
                    USER32.keybd_event(VK_ESCAPE, 0, KEYEVENTF_KEYUP, 0)
                    return
            time.sleep(0.05)


def _pump(root, *, seconds: float = 0.4) -> None:
    deadline = time.monotonic() + seconds
    while time.monotonic() < deadline:
        root.update()
        time.sleep(0.02)


def _configure_complete_state(app, *, output_dir: Path, diary_text: Path, diary_dates: Path) -> None:
    app._manual_output_dir = True
    app.output_dir_var.set(str(output_dir))
    app.open_result_folder_var.set(False)
    app.patient_name_var.set("Маркер Мужской Лайв")
    app.admission_date_var.set("10.06.2026")
    app.discharge_date_var.set("13.06.2026")
    app.diagnosis_var.set("F41.2 Смешанное тревожное и депрессивное расстройство")
    app.case_number_var.set("LIVE-E2E-001")
    app.admission_occurrence_var.set("первично")
    app.assigned_treatment_var.set("терапия по live E2E схеме.")
    app.epi_present_var.set("нет")
    app.epi_path_var.set("")
    app.expert_work_status_var.set("нет")
    app.expert_work_org_var.set("")
    app.expert_position_var.set("")
    app.expert_sick_leave_needed_var.set("нет")
    app.expert_sick_leave_from_var.set("")
    app.expert_sick_leave_number_var.set("")
    app.disability_needed_var.set("нет")
    app.psych_account_status_var.set("не состоит")
    app.rvk_referral_present_var.set("нет")
    app.rvk_referral_commissariat_var.set("")
    app.commission_date_var.set("12.06.2026")
    app.commission_number_var.set("LIVE-COM-1")
    app.rvk_act_number_var.set("LIVE-RVK-1")
    app.rvk_military_commissariat_var.set("Ленинский")
    app.vk_date_var.set("12.06.2026")
    app.vk_protocol_number_var.set("LIVE-MSE-1")
    app.vk_protocol_date_var.set("12.06.2026")
    app.vk_mse_work_org_var.set("")
    app.vk_mse_position_var.set("")
    app.sick_leave_vk_date_var.set("12.06.2026")
    app.sick_leave_vk_protocol_number_var.set("LIVE-BL-1")
    app.sick_leave_vk_protocol_date_var.set("12.06.2026")
    app.sick_leave_vk_commission_date_var.set("12.06.2026")
    app.sick_leave_vk_work_org_var.set("")
    app.sick_leave_vk_position_var.set("")
    if hasattr(app, "sick_leave_vk_work_position_var"):
        app.sick_leave_vk_work_position_var.set("")
    app.status_files = [str(diary_text)]
    app.diary_files = [str(diary_dates)]
    app._loaded_diary_text_source_signatures = None
    app._loaded_diary_date_source_signatures = None
    if getattr(app, "data", None) is not None:
        app.data.diagnosis = app.diagnosis_var.get()
        app.data.case_number = app.case_number_var.get()
        app.data.admission_date = app.admission_date_var.get()
        app.data.discharge_date = app.discharge_date_var.get()
        app.data.treatment_plan = app.assigned_treatment_var.get()
        app.data.admission_occurrence = "первично"


def _set_all_outputs(app, value: bool) -> None:
    for var in app.output_vars.values():
        var.set(value)
    app._redraw_selection_controls()
    app._generation_action_cooldown_until = 0.0
    app._generation_action_last_fingerprint = None


def _output_canvas(root, kind: str):
    label = DIARY_LABEL if kind == DIARY_KIND else DOCUMENT_LABELS[kind]
    return _find_canvas(root, label)


def _assert_generated_bundle(output_dir: Path) -> list[str]:
    files = sorted(path.name for path in output_dir.glob("*.docx"))
    if len(files) != 8:
        raise AssertionError(f"expected 8 generated DOCX files, got {len(files)}: {files!r}")
    expected_markers = (
        "Первичный осмотр",
        "Выписной эпикриз",
        "Совместный осмотр",
        "ВК на МСЭ",
        "ВК больничный",
        "Акт",
        "приёмного покоя",
        "дневник",
    )
    lowered = [name.casefold() for name in files]
    for marker in expected_markers:
        if not any(marker.casefold() in name for name in lowered):
            raise AssertionError(f"generated bundle misses {marker!r}: {files!r}")
    return files


def main() -> None:
    artifact_dir = Path(os.environ.get("LIVE_E2E_ARTIFACT_DIR", "live-e2e-artifacts")).resolve()
    trace: dict = {"schema_version": 1, "scenarios": [], "status": "started"}
    root = None
    try:
        with TemporaryDirectory(prefix="diary-filler-live-e2e-") as raw:
            work = Path(raw)
            os.environ["APPDATA"] = str(work / "appdata")
            os.environ["MEDICAL_AUTOFILL_DISABLE_DESKTOP_INTAKE"] = "1"
            os.environ["CI"] = "1"

            settings_dir = Path(os.environ["APPDATA"]) / "MedicalDiaryAutofill"
            settings_dir.mkdir(parents=True, exist_ok=True)
            (settings_dir / "settings.json").write_text(
                json.dumps(
                    {
                        "staff_profile": {
                            "configured": True,
                            "doctor": "Балаганин С.В",
                            "department_head": "Можарова Е.А.",
                            "deputy_chief": "Зуйкова А.А.",
                        }
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )

            primary = work / "primary-live.docx"
            _make_primary(primary)
            diary_text, diary_dates = _make_diary_inputs(work)

            # Import only after the isolated APPDATA/test-user environment exists.
            from app import CombinedMedicalDiaryApp
            from startup import _create_root

            root = _create_root(require_dnd=True)
            app = CombinedMedicalDiaryApp(root)
            _foreground(root)
            _pump(root)

            applied = app._apply_primary_document_path(str(primary), prompt_for_referral=False)
            if applied is False:
                raise AssertionError("production primary-document route rejected the synthetic live fixture")
            _pump(root)

            all_output_dir = work / "all-output"
            _configure_complete_state(
                app,
                output_dir=all_output_dir,
                diary_text=diary_text,
                diary_dates=diary_dates,
            )

            save_canvas = _find_canvas(root, "Создать и сохранить")
            all_kinds = (*DOCUMENT_ORDER, DIARY_KIND)

            # 1) The real visible save button must reject an empty selection.
            _set_all_outputs(app, False)
            with _PopupEscape("Ничего не выбрано") as popup:
                _physical_click(root, save_canvas)
                _pump(root, seconds=0.8)
            assert popup.detected_title == "Ничего не выбрано", popup.detected_title
            assert not list(all_output_dir.glob("*.docx")) if all_output_dir.exists() else True
            trace["scenarios"].append("empty-selection-warning-via-physical-click")

            # 2) Every output tile must be physically clickable in both directions.
            for kind in all_kinds:
                tile = _output_canvas(root, kind)
                assert not bool(app.output_vars[kind].get()), kind
                _physical_click(root, tile)
                _pump(root, seconds=0.12)
                assert bool(app.output_vars[kind].get()), f"{kind}: physical select failed"
                _physical_click(root, tile)
                _pump(root, seconds=0.12)
                assert not bool(app.output_vars[kind].get()), f"{kind}: physical deselect failed"
            trace["scenarios"].append("all-eight-output-tiles-physical-toggle")

            # 3) Missing diagnosis must open a preflight popup instead of aborting invisibly.
            _set_all_outputs(app, False)
            app.diagnosis_var.set("")
            app._popup_diagnosis_override = ""
            app.data.diagnosis = ""
            primary_tile = _output_canvas(root, "primary")
            _physical_click(root, primary_tile)
            _pump(root, seconds=0.1)
            with _PopupEscape("Недостающие данные пациента", "Данные для выбранных документов") as popup:
                _physical_click(root, save_canvas)
                _pump(root, seconds=1.0)
            assert popup.detected_title, "missing diagnosis did not produce the visible preflight popup"
            assert not list(all_output_dir.glob("*.docx")) if all_output_dir.exists() else True
            trace["scenarios"].append("missing-diagnosis-visible-preflight-popup")

            # Restore complete clinical state after the cancelled popup.
            _set_all_outputs(app, False)
            _configure_complete_state(
                app,
                output_dir=work / "sick-vk-output",
                diary_text=diary_text,
                diary_dates=diary_dates,
            )

            # 4) Direct VK sick-leave selection is itself sufficient intent.
            # The global expert sick-leave choice deliberately stays "нет".
            assert app.expert_sick_leave_needed_var.get() == "нет"
            sick_tile = _output_canvas(root, "sick_leave_vk")
            _physical_click(root, sick_tile)
            _pump(root, seconds=0.1)
            _physical_click(root, save_canvas)
            _pump(root, seconds=0.8)
            sick_files = list((work / "sick-vk-output").glob("*.docx"))
            assert len(sick_files) == 1 and "ВК больничный" in sick_files[0].name, [p.name for p in sick_files]
            trace["scenarios"].append("direct-sick-leave-vk-without-separate-toggle")

            # 5) Select every output through physical Windows input and generate
            # the complete eight-document bundle through the production path.
            _set_all_outputs(app, False)
            bundle_dir = work / "bundle-output"
            _configure_complete_state(
                app,
                output_dir=bundle_dir,
                diary_text=diary_text,
                diary_dates=diary_dates,
            )
            for kind in all_kinds:
                _physical_click(root, _output_canvas(root, kind))
                _pump(root, seconds=0.08)
                assert bool(app.output_vars[kind].get()), kind
            _physical_click(root, save_canvas)
            _pump(root, seconds=1.0)
            generated = _assert_generated_bundle(bundle_dir)

            # Inspect real generated content, not just filenames.
            sick_path = next(path for path in bundle_dir.glob("*.docx") if "ВК больничный" in path.name)
            sick_doc = Document(sick_path)
            sick_text = "\n".join(p.text for p in sick_doc.paragraphs)
            for table in sick_doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        sick_text += "\n" + "\n".join(p.text for p in cell.paragraphs)
            assert "F41.2" in sick_text, "generated VK sick-leave document lost diagnosis"
            assert "продлить лечение по листу нетрудоспособности" in sick_text.casefold(), sick_text
            trace["scenarios"].append("all-eight-real-production-generation")
            trace["generated_files"] = generated

            trace["status"] = "passed"
            _write_trace(artifact_dir, trace)
            print(
                "WINDOWS11 LIVE GUI E2E OK: physical Windows input drove all 8 output tiles; "
                "empty selection + missing diagnosis + direct sick-leave VK + complete 8-output production bundle passed"
            )
    except Exception as exc:
        trace["status"] = "failed"
        trace["error_type"] = type(exc).__name__
        trace["error"] = str(exc)[:2000]
        _write_trace(artifact_dir, trace)
        raise
    finally:
        if root is not None:
            try:
                root.destroy()
            except Exception:
                pass


if __name__ == "__main__":
    main()
