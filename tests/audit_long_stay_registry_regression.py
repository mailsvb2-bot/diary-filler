"""Regression checks for long hospital stays and uncertain discharge drafts."""
from __future__ import annotations

from datetime import date
from pathlib import Path
from tempfile import TemporaryDirectory
import sys

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from diary_service import dynamic_epicrisis_dates, apply_sick_leave_dynamic_epicrises
from patient_registry import _discharge_hides_patient, _patient_folders


def main() -> None:
    with TemporaryDirectory() as tmp:
        root = Path(tmp)
        diary = root / "long_stay.docx"
        Document().save(diary)
        count = apply_sick_leave_dynamic_epicrises(
            diary, admission_value="01.01.2026", discharge_value="01.08.2026",
            sick_leave_from="01.01.2026", clinical_state_by_date={},
        )
        assert count > 12, count
        assert len([p for p in Document(diary).paragraphs if "Динамический эпикриз" in p.text]) == count

        folder = root / "patient"
        folder.mkdir()
        discharge = folder / "Выписной эпикриз.docx"
        Document().save(discharge)
        class Parsed: discharge_date = ""
        assert not _discharge_hides_patient(folder, date(2026, 10, 18), lambda path: Parsed())
        class Past: discharge_date = "01.10.2026"
        assert _discharge_hides_patient(folder, date(2026, 10, 18), lambda path: Past())
        nested = root
        for index in range(5):
            nested = nested / f"level{index}"
            nested.mkdir()
        Document().save(nested / "Первичный осмотр.docx")
        issues = []
        assert nested in _patient_folders(root, issues)

if __name__ == "__main__":
    main()
    print("AUDIT LONG STAY AND REGISTRY REGRESSION OK")
