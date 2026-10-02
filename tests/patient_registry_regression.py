from __future__ import annotations

from datetime import date, timedelta
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
import sys

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from medical_expert import build_expert_anamnesis
from medical_formatting import treatment_period_text
from medical_models import PatientData, parse_sick_leave_value
from medical_parser import MedicalTextParser
from medical_service import MedicalDocumentService
from patient_registry import (
    PatientRegistryEntry,
    first_sick_leave_vk_date,
    hospitalization_days_on,
    is_discharge_patient_filename,
    is_primary_patient_filename,
    next_sick_leave_vk_date,
    scan_patient_registry,
    sick_leave_days_on,
)


def _assert_sick_leave_parser() -> None:
    assert parse_sick_leave_value("да, с 01.09.2026") == ("да", "01.09.2026")
    assert parse_sick_leave_value("Да с 01092026") == ("да", "01092026")
    assert parse_sick_leave_value("с 01.09.2026") == ("да", "01.09.2026")
    assert parse_sick_leave_value("с 01.09.2026 числа") == ("да", "01.09.2026")
    assert parse_sick_leave_value("от 01092026") == ("да", "01092026")
    assert parse_sick_leave_value("нет") == ("нет", "")
    assert parse_sick_leave_value("не нужен") == ("нет", "")


def _assert_primary_text_sick_leave_label() -> None:
    data = MedicalTextParser().parse_text(
        "01.09.2026 Первичный осмотр\n"
        "Ф.И.О: Маркер Иванов Тестовый\n"
        "Год рождения: 1980\n"
        "Нужен больничный лист: да, с 20.08.2026\n"
        "Диагноз: F20.0"
    )
    assert data.sick_leave == "да, с 20.08.2026"
    assert parse_sick_leave_value(data.sick_leave) == ("да", "20.08.2026")


def _assert_primary_filename_contract() -> None:
    assert is_primary_patient_filename("Иванов первичный.docx")
    assert is_primary_patient_filename("Петров первичка.doc")
    assert is_primary_patient_filename("Сидоров (первичный).docm")
    assert not is_primary_patient_filename("Иванов выписной.docx")
    assert not is_primary_patient_filename("первичный.txt")
    assert is_discharge_patient_filename("Иванов Выписной.docx")
    assert is_discharge_patient_filename("Петров (выписной).doc")
    assert not is_discharge_patient_filename("Иванов первичный.docx")
    assert not is_discharge_patient_filename("Выписной.txt")


def _assert_real_docx_registry_ingestion() -> None:
    """Exercise the exact Word -> parser -> registry path used by «Мои пациенты»."""
    with TemporaryDirectory(prefix="patient-registry-real-docx-") as temp:
        root = Path(temp)

        paragraph_folder = root / "Иванов"
        paragraph_folder.mkdir()
        paragraph_primary = paragraph_folder / "Иванов первичный.docx"
        paragraph_doc = Document()
        paragraph_doc.add_paragraph("01.10.2026")
        paragraph_doc.add_paragraph("Ф.И.О.: Маркер Иванов Тестовый")
        paragraph_doc.add_paragraph("Год рождения: 1980")
        paragraph_doc.add_paragraph("Больничный лист с 25.09.2026 числа")
        paragraph_doc.add_paragraph("Диагноз: F20.0")
        paragraph_doc.save(paragraph_primary)

        table_folder = root / "Петров"
        table_folder.mkdir()
        table_primary = table_folder / "Петров первичка.docx"
        table_doc = Document()
        table_doc.add_paragraph("02.10.2026")
        table_doc.add_paragraph("Ф.И.О.: Маркер Петров Тестовый")
        table_doc.add_paragraph("Год рождения: 1975")
        table = table_doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Больничный лист"
        table.cell(0, 1).text = "с 20.09.2026 числа"
        table_doc.add_paragraph("Диагноз: F20.0")
        table_doc.save(table_primary)

        snapshot = scan_patient_registry(root, as_of=date(2026, 10, 2))
        assert len(snapshot.patients) == 2, snapshot.issues
        assert len(snapshot.sick_leave_patients) == 2, snapshot.issues

        by_fio = {patient.fio: patient for patient in snapshot.patients}
        ivanov = by_fio["Маркер Иванов Тестовый"]
        assert ivanov.admission_date == date(2026, 10, 1)
        assert ivanov.sick_leave_needed is True
        assert ivanov.sick_leave_from == date(2026, 9, 25)
        assert not ivanov.warning

        petrov = by_fio["Маркер Петров Тестовый"]
        assert petrov.admission_date == date(2026, 10, 2)
        assert petrov.sick_leave_needed is True
        assert petrov.sick_leave_from == date(2026, 9, 20)
        assert not petrov.warning


def _assert_registry_scan_and_independent_timelines() -> None:
    with TemporaryDirectory(prefix="patient-registry-") as temp:
        root = Path(temp)
        fixtures = {
            "Иванов": ("Иванов первичный.docx", "Маркер Иванов Тестовый", "01.09.2026", "да, с 20.08.2026"),
            "Петров": ("Петров первичка.docx", "Маркер Петров Тестовый", "15.09.2026", "нет"),
            "Сидоров": ("Сидоров первичный.docx", "Маркер Сидоров Тестовый", "20.09.2026", ""),
            "Будущий": ("Будущий первичный.docx", "Маркер Будущий Тестовый", "10.10.2026", "да, с 01.10.2026"),
            "Выписан": ("Выписан первичный.docx", "Маркер Выписан Тестовый", "05.09.2026", "да, с 05.09.2026"),
        }
        by_path = {}
        for folder_name, (filename, fio, admission, sick_leave) in fixtures.items():
            folder = root / folder_name
            folder.mkdir()
            path = folder / filename
            path.touch()
            by_path[path] = SimpleNamespace(
                fio=fio,
                admission_date=admission,
                sick_leave=sick_leave,
                expert_sick_leave_needed="",
                expert_sick_leave_from="",
            )
            if folder_name == "Выписан":
                (folder / "Выписан Выписной.docx").touch()

        def parser(path: Path):
            return by_path[path]

        snapshot = scan_patient_registry(root, as_of=date(2026, 10, 1), parser=parser)
        assert len(snapshot.patients) == 3
        assert len(snapshot.sick_leave_patients) == 1

        ivanov = next(item for item in snapshot.patients if "Иванов" in item.fio)
        assert ivanov.admission_date == date(2026, 9, 1)
        assert ivanov.sick_leave_from == date(2026, 8, 20)
        assert sick_leave_days_on(ivanov, date(2026, 10, 1)) == 43
        assert hospitalization_days_on(ivanov, date(2026, 10, 1)) == 31


def _assert_dated_sick_leave_census_contract() -> None:
    undated = PatientRegistryEntry(
        fio="Маркер Без Даты", folder=Path("."), primary_path=Path("Маркер первичный.docx"),
        admission_date=date(2026, 9, 1), sick_leave_needed=True, sick_leave_from=None,
    )
    future = PatientRegistryEntry(
        fio="Маркер Будущий ЛН", folder=Path("."), primary_path=Path("Будущий первичный.docx"),
        admission_date=date(2026, 9, 1), sick_leave_needed=True, sick_leave_from=date(2026, 10, 10),
    )
    active = PatientRegistryEntry(
        fio="Маркер Активный ЛН", folder=Path("."), primary_path=Path("Активный первичный.docx"),
        admission_date=date(2026, 9, 1), sick_leave_needed=True, sick_leave_from=date(2026, 9, 20),
    )
    as_of = date(2026, 10, 1)
    assert not undated.is_on_sick_leave_on(as_of)
    assert not future.is_on_sick_leave_on(as_of)
    assert active.is_on_sick_leave_on(as_of)

def _assert_candidate_fallback_and_missing_primary_diagnostics() -> None:
    with TemporaryDirectory(prefix="patient-registry-candidate-fallback-") as temp:
        root = Path(temp)
        missing = root / "БезПервичного"
        missing.mkdir()
        multi = root / "Несколько"
        multi.mkdir()
        first = multi / "Несколько первичный.docx"
        second = multi / "Несколько первичка.docm"
        first.touch()
        second.touch()
        by_path = {
            first: SimpleNamespace(fio="Маркер Первый Без Даты", admission_date="", sick_leave="нет", expert_sick_leave_needed="", expert_sick_leave_from=""),
            second: SimpleNamespace(fio="Маркер Валидный Второй", admission_date="01.09.2026", sick_leave="нет", expert_sick_leave_needed="", expert_sick_leave_from=""),
        }
        snapshot = scan_patient_registry(root, as_of=date(2026, 10, 1), parser=lambda path: by_path[path])
        assert len(snapshot.patients) == 1, snapshot.issues
        assert snapshot.patients[0].primary_path == second
        assert snapshot.patients[0].fio == "Маркер Валидный Второй"
        missing_issue = next(issue for issue in snapshot.issues if issue.folder == missing)
        assert "первичный/первичка" in missing_issue.message

def _base_sick_leave_vk_data() -> PatientData:
    return PatientData(
        fio="Маркер Иванов Тестовый",
        birth="01.01.1980",
        admission_date="01.09.2026",
        case_number="1",
        diagnosis="F20.0",
        treatment_plan="Терапия",
        expert_work_status="нет",
        psych_account_status="нет",
        epi_present="нет",
        sick_leave_vk_date="09.09.2026",
        sick_leave_vk_protocol_number="7",
        sick_leave_vk_protocol_date="09.09.2026",
        sick_leave_vk_commission_date="09.09.2026",
    )


def _assert_sick_leave_vk_service_boundary() -> None:
    service = MedicalDocumentService()

    missing_start = _base_sick_leave_vk_data()
    try:
        service._validate_and_normalize_selected_data(missing_start, ["sick_leave_vk"])
    except ValueError as exc:
        assert "Дата начала больничного" in str(exc)
    else:
        raise AssertionError("sick_leave_vk accepted a missing opening date")

    roundtrip = _base_sick_leave_vk_data()
    roundtrip.sick_leave = "нужен с 25.08.2026"
    service._validate_and_normalize_selected_data(roundtrip, ["sick_leave_vk"])
    assert roundtrip.expert_sick_leave_from == "25.08.2026"
    assert roundtrip.sick_leave == "нужен с 25.08.2026"

    before_opening = _base_sick_leave_vk_data()
    before_opening.expert_sick_leave_from = "10.09.2026"
    before_opening.sick_leave = "нужен с 10.09.2026"
    try:
        service._validate_and_normalize_selected_data(before_opening, ["sick_leave_vk"])
    except ValueError as exc:
        assert "раньше даты начала больничного" in str(exc)
    else:
        raise AssertionError("sick_leave_vk accepted a commission before LN opening")

def _assert_vk_wednesday_schedule() -> None:
    # Day one is exactly the date from «Нужен больничный лист с ...».
    # For every possible weekday, the first VK must be a Wednesday inside the
    # inclusive 7..15-day window, choosing the latest Wednesday that does not
    # exceed day 15.
    anchor = date(2026, 9, 1)
    for offset in range(7):
        sick_from = anchor + timedelta(days=offset)
        first = first_sick_leave_vk_date(sick_from)
        inclusive_day = (first - sick_from).days + 1
        assert first.weekday() == 2
        assert 7 <= inclusive_day <= 15
        assert first + timedelta(days=7) > sick_from + timedelta(days=14)

    sick_from = date(2026, 9, 1)
    first = first_sick_leave_vk_date(sick_from)
    assert first == date(2026, 9, 9)
    assert (first - sick_from).days + 1 == 9

    assert next_sick_leave_vk_date(sick_from, date(2026, 9, 1)) == date(2026, 9, 9)
    assert next_sick_leave_vk_date(sick_from, date(2026, 9, 9)) == date(2026, 9, 9)
    second = next_sick_leave_vk_date(sick_from, date(2026, 9, 10))
    assert second == date(2026, 9, 23)
    assert second.weekday() == 2
    assert (second - first).days + 1 == 15

    third = next_sick_leave_vk_date(sick_from, date(2026, 10, 1))
    assert third == date(2026, 10, 7)
    assert third.weekday() == 2

    # The VK document itself describes hospitalization days, not sick-leave days.
    assert treatment_period_text("01.09.2026", "09.09.2026") == (
        "Находится на лечении с 01.09.2026 (9 дней)"
    )

def _assert_discharge_stays_admission_based() -> None:
    data = PatientData(
        admission_date="10.09.2026",
        discharge_date="20.09.2026",
        expert_work_status="нет",
        expert_sick_leave_needed="да",
        expert_sick_leave_from="01.09.2026",
        expert_sick_leave_number="123456",
    )
    rendered = build_expert_anamnesis(data)
    assert "Срок лечения с 10.09.2026 по 20.09.2026, 11 дней." in rendered
    assert "Срок лечения с 01.09.2026" not in rendered


def _assert_desktop_wiring_contract() -> None:
    main_source = Path("main.py").read_text(encoding="utf-8")
    mixin_source = Path("patient_registry_mixin.py").read_text(encoding="utf-8")
    window_source = Path("window_mixin.py").read_text(encoding="utf-8")
    settings_source = Path("settings_mixin.py").read_text(encoding="utf-8")

    for snippet in (
        'PATIENT_SUMMARY_ARGUMENT = "--patient-summary"',
        'PATIENT_SUMMARY_TRAY_ARGUMENT = "--patient-summary-tray"',
        "app._ensure_patient_registry_folder(first_run=True)",
        "_run_patient_summary_mode(start_in_tray=True)",
        "remove_patient_summary_autostart()",
    ):
        assert snippet in main_source
    assert "configure_patient_registry=not bool(intake_primary)" not in main_source
    for snippet in (
        'title="Из какой папки анализировать пациентов?"',
        "install_patient_summary_autostart()",
        "show_my_patients(self, *, startup_mode: bool = False, start_in_tray: bool = False)",
        "Подготовить ВК по больничному на",
        "Открыть папку пациентов",
        "Сменить путь",
        'text="× Закрыть"',
        'text="Свернуть в трей"',
        "minimize_to_tray",
        "PatientSummaryTray",
        "launch_patient_summary_tray_process",
        "open_patient_folder",
    ):
        assert snippet in mixin_source
    assert 'text="Мои пациенты", command=self.show_my_patients' in window_source
    assert 'text="Папка пациентов", command=self.show_patient_registry_folder_settings' in window_source
    assert 'patient_root_button.grid(row=0, column=2' in window_source
    registry_source = Path("patient_registry.py").read_text(encoding="utf-8")
    for snippet in (
        "class PatientSummaryTray",
        "PATIENT_SUMMARY_TRAY_ARGUMENT",
        "launch_patient_summary_tray_process",
        "_desktop_hidden_popen",
        "Shell_NotifyIcon",
        "NIM_ADD",
        "NIM_DELETE",
        "consume_restore_request",
        "consume_close_request",
    ):
        assert snippet in registry_source
    assert "DIR_PATIENT_REGISTRY" in settings_source
    assert "not Path(current_root).expanduser().is_dir()" in mixin_source


def _function_source(path: str, function_name: str) -> str:
    source = Path(path).read_text(encoding="utf-8")
    marker_text = f"    def {function_name}("
    start = source.index(marker_text)
    next_def = source.find("\n    def ", start + len(marker_text))
    return source[start:] if next_def < 0 else source[start:next_def]


def _assert_sick_leave_popup_date_contract() -> None:
    popup_functions = [
        ("dialog_expert.py", "_prompt_sick_leave_start_date_if_needed"),
        ("dialog_expert.py", "_prompt_shared_clinical_options_if_needed"),
        ("dialog_expert.py", "_prompt_expert_anamnesis_details"),
        ("dialog_expert.py", "_prompt_discharge_sick_leave_number"),
        ("dialog_expert.py", "_prompt_discharge_output_requirements"),
        ("dialog_document_details.py", "_prompt_sick_leave_vk_details"),
    ]
    for path, function_name in popup_functions:
        block = _function_source(path, function_name)
        assert "С какого числа больничный лист" in block, (path, function_name)

    expert_source = Path("dialog_expert.py").read_text(encoding="utf-8")
    assert "def _store_sick_leave_start_date_value" in expert_source
    assert 'self.data.sick_leave = f"нужен с {normalized}"' in expert_source

    orchestrator_source = Path("actions_creation_orchestrator.py").read_text(encoding="utf-8")
    assert "self.expert_sick_leave_from_var.get().strip()" in orchestrator_source


def _assert_pre_admission_sick_leave_is_not_rejected() -> None:
    data = PatientData(
        fio="Маркер Иванов Тестовый",
        birth="01.01.1980",
        admission_date="10.09.2026",
        discharge_date="20.09.2026",
        case_number="1",
        diagnosis="F20.0",
        treatment_plan="Терапия",
        expert_work_status="нет",
        expert_sick_leave_needed="да",
        expert_sick_leave_from="01.09.2026",
        psych_account_status="нет",
        epi_present="нет",
        admission_occurrence="первично",
    )
    MedicalDocumentService()._validate_and_normalize_selected_data(data, ["discharge"])
    assert data.expert_sick_leave_from == "01.09.2026"
    assert data.sick_leave == "нужен с 01.09.2026"

    dialog_source = Path("dialog_expert.py").read_text(encoding="utf-8")
    assert "Дата начала больничного не может быть раньше даты госпитализации." not in dialog_source


def main() -> None:
    _assert_sick_leave_parser()
    _assert_primary_text_sick_leave_label()
    _assert_primary_filename_contract()
    _assert_real_docx_registry_ingestion()
    _assert_registry_scan_and_independent_timelines()
    _assert_dated_sick_leave_census_contract()
    _assert_candidate_fallback_and_missing_primary_diagnostics()
    _assert_sick_leave_vk_service_boundary()
    _assert_vk_wednesday_schedule()
    _assert_discharge_stays_admission_based()
    _assert_desktop_wiring_contract()
    _assert_sick_leave_popup_date_contract()
    _assert_pre_admission_sick_leave_is_not_rejected()
    print(
        "PATIENT REGISTRY REGRESSION OK: folder scan, discharged-folder exclusion, "
        "real DOCX ingestion, dated sick-leave census, candidate fallback diagnostics, "
        "service-boundary VK chronology, all sick-leave popup opening dates, persistent tray, "
        "7..15-day Wednesday VK schedule and admission-based discharge duration are locked"
    )


if __name__ == "__main__":
    main()
