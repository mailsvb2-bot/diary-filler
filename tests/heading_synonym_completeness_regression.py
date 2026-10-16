"""Clinical section heading synonyms must preserve all parsed blocks."""
from pathlib import Path
from tempfile import TemporaryDirectory
from docx import Document
from medical_parser import MedicalTextParser

HEAD = ["25.12.2026 10:00 Первичный осмотр", "Ф.И.О.: Маркер Синонимный Тестовый", "Дата рождения: 03.04.1991", "В 3 отделение КДП поступает первично добровольно"]
BODY = {
    "complaints": ("Жалобы", "На сниженное настроение, маркер альфа."),
    "life_anamnesis": ("Анамнез жизни", "Родилась в семье служащих, маркер гамма."),
    "disease_anamnesis": ("Анамнез заболевания", "Ухудшение после развода, маркер эпсилон."),
    "mental_status": ("Психический статус", "Сознание ясное, маркер дзета."),
    "somatic_status": ("Соматический статус", "Кожные покровы обычной окраски, маркер тета."),
    "treatment_plan": ("План лечения", "Сертралин 50 мг утром, маркер мю."),
    "diagnosis": ("Диагноз", "F32.1 Депрессивный эпизод"),
}
SYNONYMS = {
    "life_anamnesis": ["Anamnesis vitae"],
    "disease_anamnesis": ["Анамнез болезни", "Anamnesis morbi"],
    "mental_status": ["Психическое состояние", "Status psychicus"],
    "somatic_status": ["Соматическое состояние", "Status somaticus"],
}

def main():
    parser = MedicalTextParser()
    with TemporaryDirectory() as raw:
        for field, aliases in SYNONYMS.items():
            for alias in aliases:
                doc = Document()
                for line in HEAD: doc.add_paragraph(line)
                for key, (label, value) in BODY.items():
                    doc.add_paragraph(f"{alias if key == field else label}: {value}")
                path = Path(raw) / f"{field}_{abs(hash(alias))}.docx"
                doc.save(path)
                parsed = parser.parse_docx(path)
                for key, (_, expected) in BODY.items():
                    actual = getattr(parsed, key)
                    assert actual == expected, (alias, key, expected, actual)
    print("OK: clinical synonym completeness")

if __name__ == "__main__": main()
