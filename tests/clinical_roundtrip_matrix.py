from __future__ import annotations

from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document
from docx.oxml import OxmlElement

from medical_constants import DOCUMENT_ORDER
from medical_docx_reader import extract_docx_text
from medical_parser import MedicalTextParser
from tools.generation_performance_profile import _make_fixture


def _long_lines(prefix: str, count: int = 40) -> list[str]:
    lines = [
        f"{prefix}_LINE_{index:03d}: полный клинический текст строки {index}; не сокращать."
        for index in range(1, count + 1)
    ]
    lines[-1] = f"{prefix}_FINAL_SENTINEL_НЕ_ОБРЕЗАТЬ"
    return lines


def _append_textbox_lines(host_paragraph, lines: list[str]) -> None:
    container = OxmlElement("w:txbxContent")
    for value in lines:
        paragraph = OxmlElement("w:p")
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = value
        run.append(text)
        paragraph.append(run)
        container.append(paragraph)
    host_paragraph._p.append(container)


def assert_word_story_sources_are_parsed() -> None:
    with TemporaryDirectory(prefix="medical-autofill-word-stories-") as temp_dir:
        source = Path(temp_dir) / "word-stories.docx"
        doc = Document()
        doc.add_paragraph("10.06.2026 Первичный осмотр")
        doc.add_paragraph("История болезни № STORY-001")
        doc.add_paragraph("Ф.И.О.: Маркер Историй Тестовый")
        doc.add_paragraph("Дата рождения: 01.01.1980")

        host = doc.add_paragraph()
        _append_textbox_lines(
            host,
            [
                "Жалобы:",
                "TEXTBOX_COMPLAINT_01 тревога и нарушение сна.",
                "TEXTBOX_COMPLAINT_FINAL не терять хвост.",
            ],
        )
        doc.add_paragraph("Анамнез жизни:")
        doc.add_paragraph("Обычный основной текст.")
        doc.add_paragraph("Диагноз:")
        doc.add_paragraph("F99.9 Тестовый диагноз")
        doc.add_paragraph("План лечения:")
        doc.add_paragraph("Тестовое лечение")

        header = doc.sections[0].header
        header.paragraphs[0].text = "Работает в организации: HEADER_WORK_ORG"
        footer = doc.sections[0].footer
        footer.paragraphs[0].text = "Должность: FOOTER_POSITION"

        doc.save(source)
        parsed = MedicalTextParser().parse_docx(source)

        assert "TEXTBOX_COMPLAINT_01" in parsed.complaints, parsed.complaints
        assert "TEXTBOX_COMPLAINT_FINAL" in parsed.complaints, parsed.complaints
        assert "__DOCX_STORY_BOUNDARY__" not in parsed.complaints, parsed.complaints
        assert parsed.work_org == "HEADER_WORK_ORG", parsed.work_org
        assert parsed.position == "FOOTER_POSITION", parsed.position


def assert_complete_field_matrix_roundtrip() -> None:
    with TemporaryDirectory(prefix="medical-autofill-complete-field-matrix-") as temp_dir:
        root = Path(temp_dir)
        navigation, service, data = _make_fixture(root)

        blocks = {
            "complaints": _long_lines("COMPLAINTS"),
            "life_anamnesis": _long_lines("LIFE"),
            "disease_anamnesis": _long_lines("DISEASE"),
            "mental_status": _long_lines("MENTAL"),
            "somatic_status": _long_lines("SOMATIC"),
            "investigation_results": _long_lines("INVESTIGATIONS"),
            "treatment_plan": _long_lines("TREATMENT"),
        }
        for field_name, lines in blocks.items():
            setattr(data, field_name, "\n".join(lines))

        recommendation_lines = _long_lines("DISCHARGE_RECOMMENDATIONS", 24)
        data.discharge_recommendations = "\n".join(recommendation_lines)

        created, _ = service.create_documents(
            navigation_path=navigation,
            output_dir=root / "output",
            discharge_date=data.discharge_date,
            selected_docs=DOCUMENT_ORDER,
            override_data=data,
        )
        assert len(created) == len(DOCUMENT_ORDER), created

        for kind, path in zip(DOCUMENT_ORDER, created):
            text = extract_docx_text(path)
            for field_name, expected_lines in blocks.items():
                positions = []
                for line in expected_lines:
                    count = text.count(line)
                    assert count == 1, (
                        f"{kind}/{path.name}: {field_name} line lost or duplicated: "
                        f"{line!r}; count={count}"
                    )
                    positions.append(text.find(line))
                assert positions == sorted(positions), (
                    f"{kind}/{path.name}: {field_name} source order changed"
                )

            recommendation_hits = [
                text.count(line) for line in recommendation_lines
            ]
            if kind == "discharge":
                assert all(count == 1 for count in recommendation_hits), (
                    f"{kind}/{path.name}: discharge recommendations were truncated "
                    f"or duplicated: {recommendation_hits}"
                )
            else:
                assert not any(recommendation_hits), (
                    f"{kind}/{path.name}: discharge-only recommendations leaked "
                    f"into another form: {recommendation_hits}"
                )


def main() -> None:
    assert_word_story_sources_are_parsed()
    assert_complete_field_matrix_roundtrip()
    print(
        "CLINICAL ROUNDTRIP MATRIX OK: text boxes/header/footer are parsed; "
        "every line of complaints/life/disease/mental/somatic/investigations/"
        "treatment survives all 7 forms; discharge recommendations stay discharge-only"
    )


if __name__ == "__main__":
    main()
