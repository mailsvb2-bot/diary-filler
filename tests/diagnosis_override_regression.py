"""Regression for UI diagnosis authority and diary-text source selection.

The doctor-visible Diagnosis field is authoritative for the complete generation
run. Automatic diary text selection follows that frozen diagnosis; an explicit
Word-file choice pins a manual override. Both modern OOXML and legacy .doc text
sources are part of the supported UI contract.
"""
from __future__ import annotations

import copy
import hashlib
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace

from docx import Document

import diary_batch
import diary_text_parser
import files_mixin
import dnd_mixin
import actions_creation_orchestrator
import printer_support
import window_mixin
from actions_medical_flow import ActionsMedicalFlowMixin
from app_initialization import AppInitializationMixin
from actions_diary_flow import ActionsDiaryFlowMixin
from actions_creation_orchestrator import ActionsCreationOrchestratorMixin
from diary_text_selection import (
    find_diary_text_file_for_diagnosis,
    iter_diary_text_docx_files,
    normalize_diary_diagnosis_name,
)
from files_mixin import FilesMixin
from dnd_mixin import DragDropMixin
from window_mixin import WindowMixin
from medical_constants import DOCUMENT_ORDER
from medical_docx_reader import extract_docx_text
from medical_models import PatientData
from medical_parser import MedicalTextParser
from medical_docx_editor_epi import remove_epi_mentions_from_document
from tools.generation_performance_profile import _make_fixture


class _Var:
    def __init__(self, value=""):
        self.value = value

    def get(self):
        return self.value

    def set(self, value):
        self.value = value


class _SnapshotHarness(ActionsMedicalFlowMixin):
    def _apply_staff_profile_to_patient_data(self, data):
        return data


class _CloseRoot:
    def __init__(self):
        self.destroy_calls = 0
        self.protocol_calls: list[tuple[str, object]] = []

    def destroy(self):
        self.destroy_calls += 1

    def protocol(self, name, callback):
        self.protocol_calls.append((str(name), callback))


class _CloseSafetyHarness(WindowMixin):
    def __init__(self, pending=None):
        self.root = _CloseRoot()
        self._pending_print_retry_files = list(pending or [])


class _MissingSourceEarlyFailHarness(ActionsCreationOrchestratorMixin):
    def __init__(self, missing_path: Path):
        self.navigation_path_var = _Var(str(missing_path))
        self._pending_print_retry_files = []
        self.status = ""
        self.staff_checks = 0
        self.popup_checks = 0

    def selected_medical_docs(self):
        return ["primary"]

    def diaries_selected(self):
        return False

    def _ensure_staff_profile_for_generation(self):
        self.staff_checks += 1
        raise AssertionError("staff prompt must not run when source is already missing")

    def _prompt_missing_patient_identity_if_needed(self):
        self.popup_checks += 1
        raise AssertionError("medical popups must not run when source is already missing")

    def _set_status(self, text):
        self.status = str(text)


class _OutputPathEarlyFailHarness(_MissingSourceEarlyFailHarness):
    def __init__(self, source_path: Path, output_target: Path):
        super().__init__(source_path)
        self.output_target = output_target

    def _result_output_dir(self):
        return self.output_target


class _ChangedSourceEarlyFailHarness(_MissingSourceEarlyFailHarness):
    def __init__(self, source_path: Path):
        super().__init__(source_path)
        self._loaded_primary_source_signature = ("loaded", 1, 2, 3, "a" * 64)

    def _primary_document_source_signature(self, _path):
        return ("current", 4, 5, 6, "b" * 64)


class _ChangedDiarySourceEarlyFailHarness(_ChangedSourceEarlyFailHarness):
    def selected_medical_docs(self):
        return []

    def diaries_selected(self):
        return True


class _ManualDiaryWithoutPrimaryHarness(ActionsCreationOrchestratorMixin):
    def __init__(self):
        self.navigation_path_var = _Var("")
        self.status = ""

    def _set_status(self, text):
        self.status = str(text)


class _DiaryInputRevisionHarness(ActionsCreationOrchestratorMixin, FilesMixin):
    def __init__(self):
        self.status_files: list[str] = []
        self.diary_files: list[str] = []
        self.status = ""

    def _set_status(self, text):
        self.status = str(text)


class _DroppedEpiIdentityHarness(ActionsCreationOrchestratorMixin, DragDropMixin, FilesMixin):
    def __init__(self):
        self.root = None
        self.epi_path_var = _Var("")
        self.epi_present_var = _Var("")
        self.output_dir_var = _Var("already-selected")
        self.status = ""
        self.logs: list[str] = []

    def _classify_dropped_file(self, _path):
        return "epi"

    def _remember_dialog_directory(self, *_args, **_kwargs):
        pass

    def reparse_navigation(self, **_kwargs):
        pass

    def _log(self, text):
        self.logs.append(str(text))

    def _set_status(self, text):
        self.status = str(text)


class _ChangedEpiEarlyFailHarness(_MissingSourceEarlyFailHarness):
    def __init__(self, primary_path: Path, epi_path: Path):
        super().__init__(primary_path)
        self.epi_path_var = _Var(str(epi_path))
        self.epi_present_var = _Var("да")
        self._loaded_primary_source_signature = ("primary", 1, 2, 3, "p" * 64)
        self._loaded_epi_source_signature = ("epi", 1, 2, 3, "e" * 64)

    def _primary_document_source_signature(self, path):
        if Path(path) == Path(self.navigation_path_var.get()):
            return self._loaded_primary_source_signature
        return ("epi", 1, 2, 3, "x" * 64)


class _GenerationGateHarness(ActionsCreationOrchestratorMixin):
    def __init__(self):
        self.impl_calls = 0
        self.status = ""

    def _set_status(self, text):
        self.status = str(text)

    def _create_selected_outputs_impl(self, *, print_after: bool = False):
        self.impl_calls += 1
        if self.impl_calls == 1:
            # Simulate a nested/re-entrant click while the first action is active.
            self.create_selected_outputs(print_after=print_after)


class _ParseCacheService:
    def __init__(self):
        self.calls = 0

    def parse_primary_document(self, _path):
        self.calls += 1
        return PatientData(fio="Свежий Пациент Тестович")


class _ParseCacheHarness(AppInitializationMixin):
    def __init__(self):
        self.service = _ParseCacheService()
        self._primary_parse_cache = {}


class _TextHarness(FilesMixin):
    def __init__(self, folder: Path):
        self.folder = folder
        self.diagnosis_var = _Var("")
        self.output_dir_var = _Var("")
        self.status_files: list[str] = []
        self.diary_texts_dir = str(folder)
        self._diary_text_files_auto_selected = False
        self.data = PatientData()
        self.events: list[str] = []

    def _candidate_diary_text_dirs(self):
        return [self.folder]

    def _update_diary_text_label(self, **_kwargs):
        pass

    def _redraw_selection_controls(self):
        pass

    def _remember_dialog_directory(self, *_args, **_kwargs):
        pass

    def _dialog_initial_dir(self, *_args, **_kwargs):
        return str(self.folder)

    def _set_output_dir_auto(self, value):
        self.output_dir_var.set(str(value))

    def _log(self, text):
        self.events.append(text)


class _PatientSwitchHarness(FilesMixin):
    def __init__(self):
        for name, default in (
            files_mixin.PATIENT_SESSION_ALWAYS_VAR_DEFAULTS
            + files_mixin.PATIENT_SESSION_TRACKED_UI_VAR_DEFAULTS
            + files_mixin.PATIENT_SESSION_SWITCH_ONLY_VAR_DEFAULTS
        ):
            setattr(self, name, _Var(copy.deepcopy(default)))
        for name, default in (
            files_mixin.PATIENT_SESSION_ALWAYS_ATTR_DEFAULTS
            + files_mixin.PATIENT_SESSION_SWITCH_ONLY_ATTR_DEFAULTS
        ):
            setattr(self, name, copy.deepcopy(default))
        for name in files_mixin.PATIENT_SESSION_SWITCH_ONLY_LIST_ATTRS:
            setattr(self, name, [])
        self.diary_texts_dir = ""
        self.diary_template_dir = ""
        self.output_vars = {
            "primary": _Var(False),
            "discharge": _Var(False),
            "diaries": _Var(False),
        }
        self.data = PatientData()

    def _set_ui_var(self, var, value):
        var.set(value)

    def _update_expert_sick_leave_display(self):
        pass

    def _update_diary_text_label(self, **_kwargs):
        pass

    def _update_diary_template_label(self, **_kwargs):
        pass

    def _set_primary_drop_empty(self):
        pass

    def _redraw_selection_controls(self):
        pass


class _InvalidPrimaryHarness(FilesMixin):
    def __init__(self, current_path: str):
        self.navigation_path_var = _Var(current_path)
        self.errors: list[tuple[str, str]] = []
        self.reset_calls = 0

    def _parse_primary_document(self, _path):
        raise ValueError("broken-docx")

    def _show_error(self, title, exc):
        self.errors.append((str(title), type(exc).__name__))

    def _reset_primary_document_runtime_state(self, **_kwargs):
        self.reset_calls += 1


class _DropBatchHarness(DragDropMixin):
    def __init__(self):
        self.root = None
        self.applied: list[str] = []
        self.logs: list[str] = []
        self.status = ""

    def _classify_dropped_file(self, _path):
        return "primary"

    def _apply_primary_document_path(self, path, **_kwargs):
        self.applied.append(str(path))
        return True

    def _log(self, text):
        self.logs.append(str(text))

    def _set_status(self, text):
        self.status = str(text)


class _AtomicDropHarness(DragDropMixin):
    def __init__(self, classifications: dict[str, str], *, primary_result: bool):
        self.root = None
        self.classifications = dict(classifications)
        self.primary_result = bool(primary_result)
        self.applied: list[str] = []
        self.logs: list[str] = []
        self.status = ""

    def _classify_dropped_file(self, path):
        return self.classifications.get(Path(path).name, "unknown")

    def _apply_primary_document_path(self, path, **_kwargs):
        self.applied.append(str(path))
        return self.primary_result

    def _log(self, text):
        self.logs.append(str(text))

    def _set_status(self, text):
        self.status = str(text)


class _DiaryFallbackHarness(ActionsDiaryFlowMixin):
    def __init__(self):
        self.navigation_path_var = _Var("")
        self.diagnosis_var = _Var("F42.2 Смешанные навязчивые мысли и действия")
        self.diary_files = ["dates.docx"]
        self.status_files = []
        self._diary_files_auto_selected = False
        self._diary_text_files_auto_selected = False
        self.diary_texts_dir = "C:/Тексты"
        self.manual_offer = None

    def _auto_select_diary_text_by_diagnosis(self, **_kwargs):
        return False

    def _offer_manual_diary_text_file(self, *, diagnosis="", initial_dir=None):
        self.manual_offer = (diagnosis, str(initial_dir or ""))
        return False


class _PrintRetryHarness(ActionsCreationOrchestratorMixin):
    def __init__(self, output_dir: Path):
        self.output_dir = output_dir
        self.expert_sick_leave_needed_var = _Var("нет")
        self.printer_var = _Var("Test Printer")
        self.open_result_folder_var = _Var(False)
        self.output_vars = {"primary": _Var(True)}
        self.root = SimpleNamespace(update_idletasks=lambda: None)
        self.logs: list[str] = []
        self.status = ""
        self.generation_calls = 0
        self.print_calls: list[list[Path]] = []
        self._pending_print_retry_files: list[Path] = []
        self.redraw_count = 0

    def selected_medical_docs(self):
        return ["primary"] if self.output_vars["primary"].get() else []

    def diaries_selected(self):
        return False

    def _selected_output_names(self, selected_medical, selected_diaries):
        return list(selected_medical)

    def _prompt_missing_patient_identity_if_needed(self):
        return True

    def _prompt_shared_clinical_options_if_needed(self, _selected):
        return True

    def _prompt_common_output_requirements(self, **_kwargs):
        return True

    def _selected_docs_need_expert_anamnesis(self, _selected):
        return False

    def _prompt_assigned_treatment_if_needed(self, **_kwargs):
        return True

    def _start_progress(self):
        pass

    def _stop_progress(self):
        pass

    def _result_output_dir(self):
        return self.output_dir

    def _capture_generation_patient_data(self, **_kwargs):
        return PatientData(
            fio="Печатнов Тест Тестович",
            output_fio="Печатнов Тест Тестович",
            admission_date="01.09.2026",
            diagnosis="F42.2 Тест",
        )

    def _create_medical_documents_impl(self, _selected, *, output_dir_override, **_kwargs):
        self.generation_calls += 1
        path = Path(output_dir_override) / "Печатнов Тест Тестович Первичный осмотр.docx"
        path.write_bytes(b"single-generated-document")
        return [path]

    def _run_print_files_safely(self, paths):
        paths = [Path(path) for path in paths]
        self.print_calls.append(paths)
        if len(self.print_calls) == 1:
            return printer_support.PrintResult([], [f"{paths[0].name}: printer offline"])
        return printer_support.PrintResult(paths, [])

    def _write_creation_report(self, **_kwargs):
        return None

    def _open_output_folder_after_creation(self, **_kwargs):
        return False

    def _log(self, text):
        self.logs.append(str(text))

    def _set_status(self, text):
        self.status = str(text)

    def _redraw_selection_controls(self):
        self.redraw_count += 1


class _PartialSetHarness(ActionsCreationOrchestratorMixin):
    def __init__(self, output_dir: Path):
        self.output_dir = output_dir
        self.expert_sick_leave_needed_var = _Var("нет")
        self.printer_var = _Var("Test Printer")
        self.open_result_folder_var = _Var(False)
        self.output_vars = {
            "primary": _Var(True),
            "diaries": _Var(True),
        }
        self.root = SimpleNamespace(update_idletasks=lambda: None)
        self.redraw_count = 0
        self.logs: list[str] = []
        self.status = ""
        self.reports: list[dict] = []
        self.medical_generation_calls = 0
        self.diary_generation_calls = 0
        self.diary_should_fail = True
        self._pending_print_retry_files: list[Path] = []

    def selected_medical_docs(self):
        return ["primary"] if self.output_vars["primary"].get() else []

    def diaries_selected(self):
        return bool(self.output_vars["diaries"].get())

    def _selected_output_names(self, selected_medical, selected_diaries):
        return [*selected_medical, "diaries"] if selected_diaries else list(selected_medical)

    def _prompt_missing_patient_identity_if_needed(self):
        return True

    def _prompt_shared_clinical_options_if_needed(self, _selected):
        return True

    def _prompt_common_output_requirements(self, **_kwargs):
        return True

    def _selected_docs_need_expert_anamnesis(self, _selected):
        return False

    def _prompt_assigned_treatment_if_needed(self, **_kwargs):
        return True

    def _start_progress(self):
        pass

    def _stop_progress(self):
        pass

    def _result_output_dir(self):
        return self.output_dir

    def _capture_generation_patient_data(self, **_kwargs):
        return PatientData(
            fio="Тестов Тест Тестович",
            output_fio="Тестов Тест Тестович",
            admission_date="01.09.2026",
            discharge_date="05.09.2026",
            diagnosis="F42.2 Смешанные навязчивые мысли и действия",
        )

    def _create_medical_documents_impl(self, _selected, *, output_dir_override, **_kwargs):
        self.medical_generation_calls += 1
        path = Path(output_dir_override) / "Первичный осмотр.docx"
        path.write_bytes(b"medical-document-ok")
        return [path]

    def _create_diaries_impl(self, *, output_dir_override, **_kwargs):
        self.diary_generation_calls += 1
        if self.diary_should_fail:
            raise ValueError("Тексты дневников не выбраны")
        path = Path(output_dir_override) / "Дневники.docx"
        path.write_bytes(b"diary-document-ok")
        return SimpleNamespace(
            created_files=[path],
            report_path=None,
            processed_files=1,
            filled_rows=1,
            month_cells_filled=1,
            final_rows_filled=1,
            removed_after_discharge_rows=0,
        )

    def _write_creation_report(self, **kwargs):
        self.reports.append(kwargs)
        return None

    def _open_output_folder_after_creation(self, **_kwargs):
        return False

    def _log(self, text):
        self.logs.append(text)

    def _set_status(self, text):
        self.status = text

    def _redraw_selection_controls(self):
        self.redraw_count += 1


def _assert_pending_print_close_safety(root: Path) -> None:
    pending = root / "pending-print.docx"
    pending.write_bytes(b"saved-document")

    answers = iter([False, True])
    prompts: list[tuple[str, str]] = []
    original_askyesno = window_mixin.messagebox.askyesno
    try:
        window_mixin.messagebox.askyesno = (
            lambda title, message, **_kwargs: (
                prompts.append((str(title), str(message))),
                next(answers),
            )[1]
        )

        app = _CloseSafetyHarness([pending])
        app._install_close_handler()
        assert len(app.root.protocol_calls) == 1, app.root.protocol_calls
        name, callback = app.root.protocol_calls[0]
        assert name == "WM_DELETE_WINDOW" and callable(callback), app.root.protocol_calls

        # Custom close / Alt+F4 must both respect an unfinished print retry.
        app._request_close()
        assert app.root.destroy_calls == 0, app.root.destroy_calls
        assert prompts and prompts[-1][0] == "Печать не завершена", prompts
        assert "DOCX останутся на диске" in prompts[-1][1], prompts[-1][1]

        callback()
        assert app.root.destroy_calls == 1, app.root.destroy_calls

        # Missing/stale paths are not a real pending queue and must not nag.
        quiet = _CloseSafetyHarness([root / "already-removed.docx"])
        prompt_count = len(prompts)
        quiet._request_close()
        assert quiet.root.destroy_calls == 1, quiet.root.destroy_calls
        assert len(prompts) == prompt_count, prompts
    finally:
        window_mixin.messagebox.askyesno = original_askyesno


def _assert_compact_address_never_consumes_clinical_residence_narrative() -> None:
    text = (
        "30.09.2025 10:00 Осмотр врача приёмного покоя.\n"
        "Иванова Ирина Ивановна, 24.07.1997, по адресу: "
        "Н. Новгород, Ленинский район, ул. Паскаля 1а.\n"
        "Работает в организации: Тестовая организация\n"
        "Должность: специалист\n"
        "Жалобы на момент осмотра: тревога.\n"
        "Анамнез жизни: Родилась в полной семье.\n"
        "Анамнез заболевания: Считает себя больной несколько лет. "
        "В 2018 году переехала в г. Нижний Новгород. "
        "В настоящее время проживает с сестрой. "
        "С 2018 года по 2023 год длительно лечилась амбулаторно. "
        "Принимала терапию, состояние менялось постепенно. "
        "КОНЕЦ_АНАМНЕЗА_НЕ_АДРЕС.\n"
        "Психический статус: Контакту доступна.\n"
        "Соматический статус: Без особенностей.\n"
        "План обследования: ОАК, ОАМ, ЭКГ, ФЛГ, ЭПИ, ЭЭГ.\n"
        "На основании данных осмотра был выставлен диагноз: F20.0 Тестовый диагноз.\n"
        "Эпидемиологический анамнез: контактов с инфекционными больными не было."
    )
    data = MedicalTextParser().parse_text(text)
    assert data.registered == "Н. Новгород, Ленинский район, ул. Паскаля 1а.", data.registered
    assert "проживает с сестрой" in data.disease_anamnesis.lower(), data.disease_anamnesis
    assert "КОНЕЦ_АНАМНЕЗА_НЕ_АДРЕС" in data.disease_anamnesis, data.disease_anamnesis
    assert "лечилась амбулаторно" not in data.registered.lower(), data.registered
    assert data.examination_plan.endswith("ЭПИ, ЭЭГ."), data.examination_plan


def _assert_epidemiology_stops_before_post_section_admission_prose() -> None:
    text = (
        "Эпидемиологический анамнез: со слов пациентки, за пределы области не выезжала, "
        "в контакте с инфекционными больными не была.\n"
        "Венерические заболевания, туберкулёз, вирусные гепатиты отрицает.\n"
        "Пациентка предъявляет жалобы на апатию и плохой сон.\n"
        "Целесообразна госпитализация пациентки в профильное отделение.\n"
        "В связи с психическим состоянием, направляется на лечение в стационар."
    )
    data = MedicalTextParser().parse_text(text)
    assert "за пределы области не выезжала" in data.epidemiology, data.epidemiology
    assert "Венерические заболевания" in data.epidemiology, data.epidemiology
    assert "предъявляет жалобы" not in data.epidemiology.lower(), data.epidemiology
    assert "целесообразна госпитализация" not in data.epidemiology.lower(), data.epidemiology
    assert "направляется на лечение" not in data.epidemiology.lower(), data.epidemiology


def _assert_epi_cleanup_preserves_examination_plan_item() -> None:
    doc = Document()
    plan = doc.add_paragraph("План обследования: ОАК, ОАМ, ФЛГ, ЭПИ, ЭЭГ.")
    doc.add_paragraph("ЭПИ - служебный шаблонный блок")
    remove_epi_mentions_from_document(doc)
    lines = [paragraph.text for paragraph in doc.paragraphs]
    assert plan.text == "План обследования: ОАК, ОАМ, ФЛГ, ЭПИ, ЭЭГ.", lines
    assert all(not line.startswith("ЭПИ") for line in lines), lines


def _assert_missing_source_fails_before_any_medical_popup(root: Path) -> None:
    app = _MissingSourceEarlyFailHarness(root / "already-moved-or-deleted.docx")
    errors: list[tuple[str, str]] = []
    original_error = actions_creation_orchestrator.messagebox.showerror
    try:
        actions_creation_orchestrator.messagebox.showerror = (
            lambda title, message, **_kwargs: errors.append((str(title), str(message)))
        )
        app._create_selected_outputs_impl(print_after=False)
    finally:
        actions_creation_orchestrator.messagebox.showerror = original_error

    assert app.staff_checks == 0, app.staff_checks
    assert app.popup_checks == 0, app.popup_checks
    assert app.status == "Создание отменено: источник пациента недоступен", app.status
    assert errors and errors[-1][0] == "Источник пациента недоступен", errors
    assert "Выберите исходный Word-документ заново" in errors[-1][1], errors[-1][1]


def _assert_output_path_file_fails_before_any_medical_popup(root: Path) -> None:
    source = root / "valid-primary-for-output-preflight.docx"
    source.write_bytes(b"primary-source")
    output_target = root / "not-a-folder.docx"
    output_target.write_bytes(b"already-a-file")

    app = _OutputPathEarlyFailHarness(source, output_target)
    errors: list[tuple[str, str]] = []
    original_error = actions_creation_orchestrator.messagebox.showerror
    try:
        actions_creation_orchestrator.messagebox.showerror = (
            lambda title, message, **_kwargs: errors.append((str(title), str(message)))
        )
        app._create_selected_outputs_impl(print_after=False)
    finally:
        actions_creation_orchestrator.messagebox.showerror = original_error

    assert app.staff_checks == 0, app.staff_checks
    assert app.popup_checks == 0, app.popup_checks
    assert app.status == "Создание отменено: выберите папку результата", app.status
    assert errors and errors[-1][0] == "Папка результата недоступна", errors
    assert "указывает на файл" in errors[-1][1], errors[-1][1]


def _assert_changed_source_fails_before_any_medical_popup(root: Path) -> None:
    source = root / "same-name-primary.docx"
    source.write_bytes(b"patient-b-replaced-content")
    app = _ChangedSourceEarlyFailHarness(source)
    errors: list[tuple[str, str]] = []
    original_error = actions_creation_orchestrator.messagebox.showerror
    try:
        actions_creation_orchestrator.messagebox.showerror = (
            lambda title, message, **_kwargs: errors.append((str(title), str(message)))
        )
        app._create_selected_outputs_impl(print_after=False)
    finally:
        actions_creation_orchestrator.messagebox.showerror = original_error

    assert app.staff_checks == 0, app.staff_checks
    assert app.popup_checks == 0, app.popup_checks
    assert app.status == "Создание отменено: источник пациента изменился", app.status
    assert errors and errors[-1][0] == "Источник пациента изменился", errors
    assert "Чтобы не смешать данные разных пациентов" in errors[-1][1], errors[-1][1]


def _assert_changed_source_blocks_diary_only_but_manual_diary_mode_survives(root: Path) -> None:
    source = root / "same-name-diary-source.docx"
    source.write_bytes(b"patient-b-replaced-content")
    app = _ChangedDiarySourceEarlyFailHarness(source)
    errors: list[tuple[str, str]] = []
    original_error = actions_creation_orchestrator.messagebox.showerror
    try:
        actions_creation_orchestrator.messagebox.showerror = (
            lambda title, message, **_kwargs: errors.append((str(title), str(message)))
        )
        assert not app._ensure_primary_source_available_for_generation([], True)
    finally:
        actions_creation_orchestrator.messagebox.showerror = original_error

    assert app.status == "Создание отменено: источник пациента изменился", app.status
    assert errors and errors[-1][0] == "Источник пациента изменился", errors

    manual = _ManualDiaryWithoutPrimaryHarness()
    assert manual._ensure_primary_source_available_for_generation([], True)
    assert manual.status == "", manual.status


def _assert_changed_diary_input_files_fail_closed(root: Path) -> None:
    text_file = root / "selected-texts.docx"
    dates_file = root / "selected-dates.docx"
    text_file.write_bytes(b"texts-v1")
    dates_file.write_bytes(b"dates-v1")

    app = _DiaryInputRevisionHarness()
    # Programmatic/legacy paths without a UI selection baseline stay compatible.
    app.status_files = [str(text_file)]
    app.diary_files = [str(dates_file)]
    assert app._ensure_diary_input_sources_available_for_generation(True)

    app._pin_diary_text_source_signatures()
    app._pin_diary_date_source_signatures()
    assert app._loaded_diary_text_source_signatures
    assert app._loaded_diary_date_source_signatures

    # A stale auto-match path that disappeared before pinning must never become
    # a valid stable baseline merely because the same missing sentinel repeats.
    vanished = root / "vanished-before-pin.docx"
    app.status_files = [str(vanished)]
    app._pin_diary_text_source_signatures()
    errors: list[tuple[str, str]] = []
    original_error = actions_creation_orchestrator.messagebox.showerror
    try:
        actions_creation_orchestrator.messagebox.showerror = (
            lambda title, message, **_kwargs: errors.append((str(title), str(message)))
        )
        assert not app._ensure_diary_input_sources_available_for_generation(True)
    finally:
        actions_creation_orchestrator.messagebox.showerror = original_error
    assert errors and "«Тексты»" in errors[-1][1], errors

    app.status_files = [str(text_file)]
    app.diary_files = [str(dates_file)]
    app._pin_diary_text_source_signatures()
    app._pin_diary_date_source_signatures()

    errors = []
    original_error = actions_creation_orchestrator.messagebox.showerror
    try:
        actions_creation_orchestrator.messagebox.showerror = (
            lambda title, message, **_kwargs: errors.append((str(title), str(message)))
        )

        text_file.write_bytes(b"texts-v2-replaced")
        assert not app._ensure_diary_input_sources_available_for_generation(True)
        assert app.status == "Создание отменено: источник «Тексты» изменился", app.status
        assert errors and errors[-1][0] == "Источник дневников изменился", errors
        assert "«Тексты»" in errors[-1][1], errors[-1][1]

        app._pin_diary_text_source_signatures()
        app._pin_diary_date_source_signatures()
        errors.clear()
        dates_file.write_bytes(b"dates-v2-replaced")
        assert not app._ensure_diary_input_sources_available_for_generation(True)
        assert app.status == "Создание отменено: источник «Даты» изменился", app.status
        assert errors and "«Даты»" in errors[-1][1], errors
    finally:
        actions_creation_orchestrator.messagebox.showerror = original_error


def _assert_dropped_epi_is_bound_to_drop_time_revision(root: Path) -> None:
    epi = root / "dropped-epi.docx"
    epi.write_bytes(b"epi-at-drop-time")
    app = _DroppedEpiIdentityHarness()

    app._handle_dropped_files([str(epi)])
    expected = app._primary_document_source_signature(epi)
    assert app.epi_path_var.get() == str(epi), app.epi_path_var.get()
    assert app.epi_present_var.get() == "да", app.epi_present_var.get()
    assert app._loaded_epi_source_signature == expected, app._loaded_epi_source_signature

    epi.write_bytes(b"epi-replaced-after-drop")
    errors: list[tuple[str, str]] = []
    original_error = actions_creation_orchestrator.messagebox.showerror
    try:
        actions_creation_orchestrator.messagebox.showerror = (
            lambda title, message, **_kwargs: errors.append((str(title), str(message)))
        )
        assert not app._ensure_epi_source_available_for_generation(["primary"])
    finally:
        actions_creation_orchestrator.messagebox.showerror = original_error

    assert app.status == "Создание отменено: файл ЭПИ изменился", app.status
    assert errors and errors[-1][0] == "Файл ЭПИ изменился", errors


def _assert_changed_epi_fails_before_any_medical_popup(root: Path) -> None:
    primary = root / "stable-primary.docx"
    epi = root / "selected-epi.docx"
    primary.write_bytes(b"stable-primary")
    epi.write_bytes(b"replaced-epi")
    app = _ChangedEpiEarlyFailHarness(primary, epi)
    errors: list[tuple[str, str]] = []
    original_error = actions_creation_orchestrator.messagebox.showerror
    try:
        actions_creation_orchestrator.messagebox.showerror = (
            lambda title, message, **_kwargs: errors.append((str(title), str(message)))
        )
        app._create_selected_outputs_impl(print_after=False)
    finally:
        actions_creation_orchestrator.messagebox.showerror = original_error

    assert app.staff_checks == 0, app.staff_checks
    assert app.popup_checks == 0, app.popup_checks
    assert app.status == "Создание отменено: файл ЭПИ изменился", app.status
    assert errors and errors[-1][0] == "Файл ЭПИ изменился", errors
    assert "не смешать данные разных пациентов" in errors[-1][1], errors[-1][1]


def _assert_generation_action_gate_blocks_reentry_and_queued_double_click() -> None:
    app = _GenerationGateHarness()

    app.create_selected_outputs(print_after=False)
    assert app.impl_calls == 1, app.impl_calls
    assert app.status == "Создание уже выполняется", app.status

    # Represents the second ButtonRelease that Tk queued for a double click.
    app.create_selected_outputs(print_after=False)
    assert app.impl_calls == 1, app.impl_calls
    assert app.status == "Предыдущее создание уже завершено", app.status

    # A deliberate later action remains available.
    app._generation_action_cooldown_until = 0.0
    app.create_selected_outputs(print_after=True)
    assert app.impl_calls == 2, app.impl_calls


def _assert_ui_diagnosis_wins_snapshot() -> None:
    app = _SnapshotHarness()
    app.data = PatientData(fio="Иванов Иван Иванович", diagnosis="F99.9 Старый диагноз из первичного")
    app.navigation_path_var = _Var("")
    app.patient_name_var = _Var("Иванов Иван Иванович")
    app.admission_date_var = _Var("10.06.2026")
    app.discharge_date_var = _Var("11.06.2026")
    app.diagnosis_var = _Var("F20.0 Параноидная шизофрения")
    app._popup_discharge_date_override = ""
    app._popup_diagnosis_override = "F06.8 Старый диагноз из popup"

    snapshot = app._capture_generation_patient_data(require_primary=False)
    assert snapshot.diagnosis == "F20.0 Параноидная шизофрения", snapshot.diagnosis

    app.diagnosis_var.set("Органическое расстройство личности")
    snapshot = app._capture_generation_patient_data(require_primary=False)
    assert snapshot.diagnosis == "Органическое расстройство личности", snapshot.diagnosis


def _assert_all_medical_documents_receive_new_diagnosis(root: Path) -> None:
    nav, service, data = _make_fixture(root)
    old = data.diagnosis
    new = "F20.0 Параноидная шизофрения"
    data.diagnosis = new
    output = root / "medical-diagnosis-override"
    created, used = service.create_documents(
        navigation_path=nav,
        output_dir=output,
        discharge_date=data.discharge_date,
        selected_docs=DOCUMENT_ORDER,
        override_data=data,
    )
    assert used.diagnosis == new
    assert len(created) == len(DOCUMENT_ORDER)
    texts = {path.name: extract_docx_text(path) for path in created}
    joined = "\n".join(texts.values())
    assert old not in joined, f"old diagnosis leaked into generated set: {old!r}"
    for name, text in texts.items():
        assert "Параноидная шизофрения" in text, f"new UI diagnosis missing from {name}"


def _assert_verbal_matching_and_word_formats(root: Path) -> None:
    folder = root / "Тексты"
    folder.mkdir()
    (folder / "F20.0 Органическое расстройство.docx").touch()
    (folder / "шизофрения параноидная.docx").touch()
    (folder / "органическое расстройство личности.doc").touch()
    (folder / "депрессивный эпизод.docm").touch()

    visible = {p.name for p in iter_diary_text_docx_files(folder)}
    assert visible == {
        "F20.0 Органическое расстройство.docx",
        "шизофрения параноидная.docx",
        "органическое расстройство личности.doc",
        "депрессивный эпизод.docm",
    }, visible

    found = find_diary_text_file_for_diagnosis(folder, "F20.0 Параноидная шизофрения")
    assert found is not None and found.name == "шизофрения параноидная.docx", found

    found = find_diary_text_file_for_diagnosis(folder, "F06.8 Органическое расстройство личности")
    assert found is not None and found.name == "органическое расстройство личности.doc", found

    assert normalize_diary_diagnosis_name("F20.0") == ""
    assert find_diary_text_file_for_diagnosis(folder, "F20.0") is None


def _assert_auto_refresh_and_manual_pin(root: Path) -> None:
    folder = root / "auto-texts"
    folder.mkdir()
    old_file = folder / "органическое расстройство личности.docx"
    new_file = folder / "шизофрения параноидная.docx"
    manual_file = folder / "мой выбранный текст.doc"
    for path in (old_file, new_file, manual_file):
        path.touch()

    app = _TextHarness(folder)
    app.status_files = [str(old_file)]
    app._diary_text_files_auto_selected = True
    assert app._auto_select_diary_text_by_diagnosis(
        diagnosis_override="F20.0 Параноидная шизофрения",
        ask_folder=False,
    )
    assert app.status_files == [str(new_file)], app.status_files
    assert app._diary_text_files_auto_selected is True
    assert app._loaded_diary_text_source_signatures, app._loaded_diary_text_source_signatures

    assert not app._auto_select_diary_text_by_diagnosis(
        diagnosis_override="Кататоническое состояние редкого типа",
        ask_folder=False,
    )
    assert app.status_files == [], app.status_files
    assert app._loaded_diary_text_source_signatures is None

    app.status_files = [str(manual_file)]
    app._diary_text_files_auto_selected = False
    assert app._auto_select_diary_text_by_diagnosis(
        diagnosis_override="F20.0 Параноидная шизофрения",
        ask_folder=False,
    )
    assert app.status_files == [str(manual_file)]
    assert app._diary_text_files_auto_selected is False


def _assert_manual_text_is_patient_scoped(root: Path) -> None:
    manual_file = root / "patient-a-manual-text.docx"
    manual_file.touch()
    app = _PatientSwitchHarness()

    app.status_files = [str(manual_file)]
    app._diary_text_files_auto_selected = False
    app._reset_primary_document_runtime_state(clear_patient_inputs=False)
    assert app.status_files == [str(manual_file)], app.status_files
    assert app._diary_text_files_auto_selected is False

    app.status_files = [str(manual_file)]
    app._diary_text_files_auto_selected = False
    app._reset_primary_document_runtime_state(clear_patient_inputs=True)
    assert app.status_files == [], app.status_files
    assert app._diary_text_files_auto_selected is False


def _assert_full_patient_switch_reset_matrix() -> None:
    """Every high-risk patient-specific field must be owned by the reset contract."""
    required_always_vars = {
        "assigned_treatment_var",
        "case_number_var",
        "admission_occurrence_var",
        "expert_work_status_var",
        "expert_work_org_var",
        "expert_position_var",
        "expert_sick_leave_needed_var",
        "expert_sick_leave_from_var",
        "expert_sick_leave_number_var",
        "disability_needed_var",
        "psych_account_status_var",
        "psych_account_since_year_var",
        "rvk_referral_present_var",
        "rvk_referral_commissariat_var",
        "vk_mse_work_org_var",
        "vk_mse_position_var",
        "sick_leave_vk_work_org_var",
        "sick_leave_vk_position_var",
        "sick_leave_vk_work_position_var",
    }
    required_tracked_ui_vars = {
        "patient_name_var",
        "admission_date_var",
        "discharge_date_var",
        "diagnosis_var",
    }
    required_switch_vars = {
        "rvk_act_number_var",
        "rvk_military_commissariat_var",
        "rvk_work_position_var",
        "vk_date_var",
        "vk_protocol_number_var",
        "vk_protocol_date_var",
        "sick_leave_vk_date_var",
        "sick_leave_vk_protocol_number_var",
        "sick_leave_vk_protocol_date_var",
        "sick_leave_vk_commission_date_var",
        "commission_date_var",
        "commission_number_var",
        "epi_path_var",
        "epi_present_var",
    }
    required_always_attrs = {
        "_primary_work_org_default",
        "_primary_work_position_default",
        "_work_details_manually_edited",
        "_manual_patient_name",
        "_manual_admission_date",
        "_manual_discharge_date",
        "_manual_diagnosis",
        "_popup_diagnosis_override",
        "_popup_discharge_date_override",
        "_popup_fio_override",
        "_popup_birth_override",
    }
    required_switch_attrs = {
        "_last_committee_date",
        "_last_protocol_date",
        "_diary_text_files_auto_selected",
        "_diary_files_auto_selected",
        "_loaded_diary_text_source_signatures",
        "_loaded_diary_date_source_signatures",
    }
    required_switch_lists = {"status_files", "diary_files", "_pending_print_retry_files"}

    always_vars = dict(files_mixin.PATIENT_SESSION_ALWAYS_VAR_DEFAULTS)
    tracked_ui_vars = dict(files_mixin.PATIENT_SESSION_TRACKED_UI_VAR_DEFAULTS)
    switch_vars = dict(files_mixin.PATIENT_SESSION_SWITCH_ONLY_VAR_DEFAULTS)
    always_attrs = dict(files_mixin.PATIENT_SESSION_ALWAYS_ATTR_DEFAULTS)
    switch_attrs = dict(files_mixin.PATIENT_SESSION_SWITCH_ONLY_ATTR_DEFAULTS)
    switch_lists = set(files_mixin.PATIENT_SESSION_SWITCH_ONLY_LIST_ATTRS)

    assert required_always_vars <= set(always_vars), required_always_vars - set(always_vars)
    assert required_tracked_ui_vars <= set(tracked_ui_vars), required_tracked_ui_vars - set(tracked_ui_vars)
    assert required_switch_vars <= set(switch_vars), required_switch_vars - set(switch_vars)
    assert required_always_attrs <= set(always_attrs), required_always_attrs - set(always_attrs)
    assert required_switch_attrs <= set(switch_attrs), required_switch_attrs - set(switch_attrs)
    assert required_switch_lists <= switch_lists, required_switch_lists - switch_lists

    app = _PatientSwitchHarness()
    for name in required_always_vars | required_tracked_ui_vars | required_switch_vars:
        getattr(app, name).set("PATIENT_A_LEAK")
    for name in required_always_attrs | required_switch_attrs:
        default = always_attrs.get(name, switch_attrs.get(name))
        setattr(app, name, True if isinstance(default, bool) else "PATIENT_A_LEAK")
    app.status_files = ["patient-a-text.docx"]
    app.diary_files = ["patient-a-dates.docx"]
    app._pending_print_retry_files = [Path("patient-a-unprinted.docx")]
    for var in app.output_vars.values():
        var.set(True)
    app.data = PatientData(
        fio="Пациент А",
        admission_date="01.09.2026",
        discharge_date="05.09.2026",
        diagnosis="F99.9 Данные пациента А",
    )

    app._reset_primary_document_runtime_state(clear_patient_inputs=True)

    for name in required_always_vars:
        assert getattr(app, name).get() == always_vars[name], (name, getattr(app, name).get())
    for name in required_tracked_ui_vars:
        assert getattr(app, name).get() == tracked_ui_vars[name], (name, getattr(app, name).get())
    for name in required_switch_vars:
        assert getattr(app, name).get() == switch_vars[name], (name, getattr(app, name).get())
    for name in required_always_attrs:
        assert getattr(app, name) == always_attrs[name], (name, getattr(app, name))
    for name in required_switch_attrs:
        assert getattr(app, name) == switch_attrs[name], (name, getattr(app, name))
    assert app.status_files == []
    assert app.diary_files == []
    assert app._pending_print_retry_files == []
    assert all(not var.get() for var in app.output_vars.values()), app.output_vars
    assert app.data == PatientData(), app.data


def _assert_primary_cache_rejects_same_metadata_wrong_digest(root: Path) -> None:
    source = root / "digest-cache.docx"
    source.write_bytes(b"fresh-patient-bytes")
    stat = source.stat()
    key = str(source.resolve())
    app = _ParseCacheHarness()
    stale = PatientData(fio="Старый Пациент Ошибочный")
    app._primary_parse_cache[key] = (
        int(getattr(stat, "st_mtime_ns", int(stat.st_mtime * 1_000_000_000))),
        int(getattr(stat, "st_ctime_ns", int(stat.st_ctime * 1_000_000_000))),
        int(stat.st_size),
        "0" * 64,
        stale,
    )

    parsed = app._parse_primary_document(source)
    assert parsed.fio == "Свежий Пациент Тестович", parsed.fio
    assert app.service.calls == 1, app.service.calls

    # Unchanged content now hits the freshly written digest-bound cache.
    parsed_again = app._parse_primary_document(source)
    assert parsed_again.fio == "Свежий Пациент Тестович", parsed_again.fio
    assert app.service.calls == 1, app.service.calls


def _assert_same_path_replacement_is_patient_switch(root: Path) -> None:
    source = root / "Первичный осмотр.docx"
    source.write_bytes(b"patient-a")
    app = _PatientSwitchHarness()
    first_signature = app._primary_document_source_signature(source)
    assert len(first_signature) == 5, first_signature
    assert first_signature[-1] == hashlib.sha256(b"patient-a").hexdigest(), first_signature
    app._loaded_primary_source_signature = first_signature

    source.write_bytes(b"patient-b-replacement")
    second_signature = app._primary_document_source_signature(source)
    assert first_signature != second_signature, (first_signature, second_signature)
    assert app._is_primary_document_switch(str(source), str(source), second_signature) is True
    assert app._is_primary_document_switch("", str(source), second_signature) is False


def _assert_invalid_new_source_preserves_open_patient(root: Path) -> None:
    current = root / "current.docx"
    current.write_bytes(b"current")
    broken = root / "broken.docx"
    broken.write_bytes(b"not-a-valid-docx")
    app = _InvalidPrimaryHarness(str(current))
    assert app._apply_primary_document_path(str(broken), prompt_for_referral=True) is False
    assert app.navigation_path_var.get() == str(current), app.navigation_path_var.get()
    assert app.reset_calls == 0, app.reset_calls
    assert app.errors == [("Не удалось прочитать медицинский документ", "ValueError")], app.errors


def _assert_multi_primary_drop_fails_safe(root: Path) -> None:
    first = root / "patient-a.docx"
    second = root / "patient-b.docx"
    first.write_bytes(b"a")
    second.write_bytes(b"b")
    app = _DropBatchHarness()
    warnings: list[str] = []
    original_warning = dnd_mixin.messagebox.showwarning
    try:
        dnd_mixin.messagebox.showwarning = lambda _title, message, **_kwargs: warnings.append(str(message))
        app._handle_dropped_files([str(first), str(second)])
    finally:
        dnd_mixin.messagebox.showwarning = original_warning
    assert app.applied == [], app.applied
    assert len(warnings) == 1 and "несколько медицинских документов" in warnings[0].lower(), warnings
    assert "Выберите один" in app.status, app.status


def _assert_failed_primary_aborts_entire_drop_batch(root: Path) -> None:
    primary = root / "new-patient.docx"
    epi = root / "new-patient-epi.docx"
    texts = root / "new-patient-texts.docx"
    for path in (primary, epi, texts):
        path.write_bytes(path.name.encode("utf-8"))

    app = _AtomicDropHarness(
        {
            primary.name: "primary",
            epi.name: "epi",
            texts.name: "diary_status",
        },
        primary_result=False,
    )
    # If the implementation touches any auxiliary state after primary=False,
    # this deliberately minimal harness raises instead of hiding partial apply.
    app._handle_dropped_files([str(primary), str(epi), str(texts)])
    assert app.applied == [str(primary)], app.applied
    assert "остальные файлы не изменены" in app.status, app.status


def _assert_multi_epi_drop_fails_before_primary_apply(root: Path) -> None:
    primary = root / "primary-with-two-epi.docx"
    epi_a = root / "epi-a.docx"
    epi_b = root / "epi-b.docx"
    for path in (primary, epi_a, epi_b):
        path.write_bytes(path.name.encode("utf-8"))

    app = _AtomicDropHarness(
        {
            primary.name: "primary",
            epi_a.name: "epi",
            epi_b.name: "epi",
        },
        primary_result=True,
    )
    warnings: list[tuple[str, str]] = []
    original_warning = dnd_mixin.messagebox.showwarning
    try:
        dnd_mixin.messagebox.showwarning = (
            lambda title, message, **_kwargs: warnings.append((str(title), str(message)))
        )
        app._handle_dropped_files([str(primary), str(epi_a), str(epi_b)])
    finally:
        dnd_mixin.messagebox.showwarning = original_warning

    assert app.applied == [], app.applied
    assert warnings and warnings[-1][0] == "Несколько файлов ЭПИ", warnings
    assert "один файл ЭПИ" in warnings[-1][1], warnings
    assert app.status == "Выберите один файл ЭПИ", app.status


def _assert_multi_diary_source_folder_drop_fails_before_primary_apply(root: Path) -> None:
    primary = root / "primary-with-two-source-folders.docx"
    primary.write_bytes(b"primary")

    dates_a = root / "dates-a"
    dates_b = root / "dates-b"
    texts_a = root / "texts-a"
    texts_b = root / "texts-b"
    for folder in (dates_a, dates_b, texts_a, texts_b):
        folder.mkdir()

    cases = [
        (
            {
                primary.name: "primary",
                dates_a.name: "numbered_diary_template_dir",
                dates_b.name: "numbered_diary_template_dir",
            },
            [str(primary), str(dates_a), str(dates_b)],
            "Несколько папок «Даты»",
            "Выберите одну папку «Даты»",
        ),
        (
            {
                primary.name: "primary",
                texts_a.name: "diary_text_dir",
                texts_b.name: "diary_text_dir",
            },
            [str(primary), str(texts_a), str(texts_b)],
            "Несколько папок «Тексты»",
            "Выберите одну папку «Тексты»",
        ),
    ]

    original_warning = dnd_mixin.messagebox.showwarning
    try:
        for classifications, paths, expected_title, expected_status in cases:
            app = _AtomicDropHarness(classifications, primary_result=True)
            warnings: list[tuple[str, str]] = []
            dnd_mixin.messagebox.showwarning = (
                lambda title, message, **_kwargs: warnings.append((str(title), str(message)))
            )
            app._handle_dropped_files(paths)
            assert app.applied == [], app.applied
            assert warnings and warnings[-1][0] == expected_title, warnings
            assert app.status == expected_status, app.status
    finally:
        dnd_mixin.messagebox.showwarning = original_warning


def _assert_manual_picker_accepts_doc(root: Path) -> None:
    folder = root / "manual-picker"
    folder.mkdir()
    legacy = folder / "ручной текст.doc"
    legacy.write_bytes(b"legacy-placeholder")
    app = _TextHarness(folder)

    original = files_mixin.filedialog.askopenfilename
    captured = {}
    try:
        def fake_picker(*_args, **kwargs):
            captured.update(kwargs)
            return str(legacy)

        files_mixin.filedialog.askopenfilename = fake_picker
        app.choose_status_file()
    finally:
        files_mixin.filedialog.askopenfilename = original

    assert app.status_files == [str(legacy)]
    assert app.diary_texts_dir == str(folder)
    assert app._diary_text_files_auto_selected is False
    assert app._loaded_diary_text_source_signatures, app._loaded_diary_text_source_signatures
    filetypes_repr = repr(captured.get("filetypes"))
    for suffix in ("*.doc", "*.docx", "*.docm"):
        assert suffix in filetypes_repr, filetypes_repr


def _assert_legacy_doc_parser_route(root: Path) -> None:
    legacy = root / "legacy.doc"
    legacy.write_bytes(b"legacy-placeholder")
    original = diary_text_parser._convert_legacy_doc_to_docx
    try:
        def fake_convert(_source: Path, target: Path) -> None:
            doc = Document()
            doc.add_paragraph("Пациент спокоен, ориентирован, продуктивному контакту доступен.")
            doc.save(target)

        diary_text_parser._convert_legacy_doc_to_docx = fake_convert
        statuses = diary_text_parser.extract_statuses_from_docx(legacy)
    finally:
        diary_text_parser._convert_legacy_doc_to_docx = original
    assert statuses and "Пациент спокоен" in statuses[0]

    original_extract = diary_batch.extract_statuses_from_docx
    try:
        diary_batch.extract_statuses_from_docx = lambda *_args, **_kwargs: [
            "Пациент спокоен, ориентирован, продуктивному контакту доступен."
        ]
        statuses = diary_batch.read_statuses_from_files([legacy])
    finally:
        diary_batch.extract_statuses_from_docx = original_extract
    assert statuses
    try:
        diary_batch._existing_docx_files([legacy], "таблица дневников")
    except ValueError:
        pass
    else:
        raise AssertionError("legacy .doc must not be accepted as a Dates template")


def _assert_failed_auto_match_offers_manual_word_file(root: Path) -> None:
    folder = root / "manual-fallback-texts"
    folder.mkdir()
    manual = folder / "нужный врачом шаблон.docx"
    manual.touch()
    app = _TextHarness(folder)
    app.diagnosis_var.set("F42.2 Смешанные навязчивые мысли и действия")

    original_dir = files_mixin.filedialog.askdirectory
    original_file = files_mixin.filedialog.askopenfilename
    original_yesno = files_mixin.messagebox.askyesno
    captured: dict[str, object] = {"yesno": 0}
    try:
        files_mixin.filedialog.askdirectory = lambda **_kwargs: str(folder)

        def fake_yesno(title, message, **_kwargs):
            captured["yesno"] = int(captured["yesno"]) + 1
            captured["title"] = title
            captured["message"] = message
            return True

        def fake_file(*_args, **kwargs):
            captured["file_kwargs"] = kwargs
            return str(manual)

        files_mixin.messagebox.askyesno = fake_yesno
        files_mixin.filedialog.askopenfilename = fake_file
        app.choose_status_files()
    finally:
        files_mixin.filedialog.askdirectory = original_dir
        files_mixin.filedialog.askopenfilename = original_file
        files_mixin.messagebox.askyesno = original_yesno

    assert captured["yesno"] == 1, captured
    assert "Выбрать нужный файл вручную" in str(captured.get("message", "")), captured
    picker = captured.get("file_kwargs")
    assert isinstance(picker, dict), captured
    assert Path(str(picker.get("initialdir"))).resolve() == folder.resolve(), picker
    filetypes_repr = repr(picker.get("filetypes"))
    assert "*.doc" in filetypes_repr and "*.docx" in filetypes_repr and "*.docm" in filetypes_repr, picker
    assert app.status_files == [str(manual)], app.status_files
    assert app._diary_text_files_auto_selected is False


def _assert_diary_creation_path_offers_manual_fallback() -> None:
    app = _DiaryFallbackHarness()
    snapshot = PatientData(
        fio="Тестов Тест Тестович",
        admission_date="01.09.2026",
        discharge_date="05.09.2026",
        diagnosis="F42.2 Смешанные навязчивые мысли и действия",
    )
    try:
        app._create_diaries_impl(patient_data_snapshot=snapshot)
    except ValueError as exc:
        assert "Тексты дневников не выбраны" in str(exc), exc
    else:
        raise AssertionError("diary flow must stop only after manual fallback was declined")
    assert app.manual_offer is not None
    assert app.manual_offer[0] == snapshot.diagnosis, app.manual_offer


def _assert_commit_rollback_reports_residual_file(root: Path) -> None:
    final_dir = root / "commit-rollback-output"
    staging_dir = root / "commit-rollback-staging"
    final_dir.mkdir()
    staging_dir.mkdir()
    first = staging_dir / "first.docx"
    second = staging_dir / "second.docx"
    first.write_bytes(b"first")
    second.write_bytes(b"second")

    app = ActionsCreationOrchestratorMixin()
    original_replace = actions_creation_orchestrator.os.replace
    original_unlink = Path.unlink
    replace_calls = 0

    def fail_second_replace(src, dst):
        nonlocal replace_calls
        replace_calls += 1
        if replace_calls == 2:
            raise OSError("synthetic second commit failure")
        return original_replace(src, dst)

    def block_committed_unlink(path_obj, *args, **kwargs):
        if Path(path_obj) == final_dir / "first.docx":
            raise OSError("synthetic rollback lock")
        return original_unlink(path_obj, *args, **kwargs)

    caught = None
    try:
        actions_creation_orchestrator.os.replace = fail_second_replace  # type: ignore[assignment]
        Path.unlink = block_committed_unlink  # type: ignore[assignment]
        try:
            app._commit_staged_generation(
                final_output_dir=final_dir,
                staged_medical=[first, second],
                diary_result=None,
            )
        except Exception as exc:
            caught = exc
    finally:
        Path.unlink = original_unlink  # type: ignore[assignment]
        actions_creation_orchestrator.os.replace = original_replace  # type: ignore[assignment]

    assert caught is not None
    assert getattr(caught, "rollback_remaining", 0) == 1, caught
    assert (final_dir / "first.docx").is_file()
    note = app._commit_failure_rollback_note(caught)
    assert "Откат завершён не полностью" in note, note
    assert "осталось 1 файл(ов)" in note, note
    assert "Файлы, уже перенесённые этой попыткой, удалены." not in note, note


def _assert_print_failure_retries_same_saved_file_without_regeneration(root: Path) -> None:
    output = root / "print-retry-output"
    app = _PrintRetryHarness(output)
    warnings: list[tuple[str, str]] = []
    original_warning = actions_creation_orchestrator.messagebox.showwarning
    try:
        actions_creation_orchestrator.messagebox.showwarning = (
            lambda title, message, **_kwargs: warnings.append((str(title), str(message)))
        )
        app.create_selected_outputs(print_after=True)

        saved_after_first = sorted(output.glob("*.docx"))
        assert len(saved_after_first) == 1, saved_after_first
        assert app.generation_calls == 1, app.generation_calls
        assert len(app.print_calls) == 1, app.print_calls
        assert app.output_vars["primary"].get() is False
        assert app._pending_print_retry_files == saved_after_first, app._pending_print_retry_files
        assert "печать не завершена" in app.status.lower(), app.status
        assert warnings and warnings[-1][0] == "Создано, но печать с ошибками", warnings

        # With no outputs selected, the same print button is now a print-only
        # retry. No document generation or collision-suffix copy may happen.
        app.create_selected_outputs(print_after=True)
    finally:
        actions_creation_orchestrator.messagebox.showwarning = original_warning

    saved_after_retry = sorted(output.glob("*.docx"))
    assert saved_after_retry == saved_after_first, (saved_after_first, saved_after_retry)
    assert app.generation_calls == 1, app.generation_calls
    assert len(app.print_calls) == 2, app.print_calls
    assert app.print_calls[1] == saved_after_first, app.print_calls
    assert app._pending_print_retry_files == [], app._pending_print_retry_files
    assert app.status == "Готово: сохранённые документы отправлены на печать", app.status
    assert not list(output.glob("* (2).docx")), saved_after_retry


def _assert_diary_failure_keeps_medical_documents(root: Path) -> None:
    output = root / "partial-set-output"
    app = _PartialSetHarness(output)
    original_warning = actions_creation_orchestrator.messagebox.showwarning
    original_error = actions_creation_orchestrator.messagebox.showerror
    original_print_files = printer_support.print_files
    warnings: list[tuple[str, str]] = []
    errors: list[tuple[str, str]] = []
    print_calls: list[tuple[list[Path], str]] = []

    def fake_print_files(paths, printer):
        normalized = [Path(path) for path in paths]
        print_calls.append((normalized, printer))
        return printer_support.PrintResult(normalized, [])

    try:
        actions_creation_orchestrator.messagebox.showwarning = lambda title, message, **_kwargs: warnings.append((title, message))
        actions_creation_orchestrator.messagebox.showerror = lambda title, message, **_kwargs: errors.append((title, message))
        printer_support.print_files = fake_print_files  # type: ignore[assignment]
        app.create_selected_outputs(print_after=True)

        saved = list(output.glob("Первичный осмотр*.docx"))
        assert len(saved) == 1 and saved[0].read_bytes() == b"medical-document-ok", saved
        assert app._pending_print_retry_files == saved, app._pending_print_retry_files
        assert not errors, errors
        assert warnings and warnings[-1][0] == "Комплект создан частично", warnings
        assert "Дневники" in warnings[-1][1] and "Тексты дневников не выбраны" in warnings[-1][1], warnings
        assert "Автоматическая печать не запускалась" in warnings[-1][1], warnings
        assert "поставлены в очередь" in warnings[-1][1], warnings
        assert print_calls == [], print_calls
        assert app.output_vars["primary"].get() is False, app.output_vars["primary"].get()
        assert app.output_vars["diaries"].get() is True, app.output_vars["diaries"].get()
        assert app.medical_generation_calls == 1, app.medical_generation_calls
        assert app.diary_generation_calls == 1, app.diary_generation_calls
        assert app.redraw_count == 1, app.redraw_count
        assert app.reports and any("Дневники:" in item for item in (app.reports[-1].get("errors") or [])), app.reports
        assert app.status == "Готово частично: доступные документы сохранены", app.status

        # The doctor fixes the diary source and repeats the original print action.
        # Only diaries are generated now, but printing must include both the
        # previously saved medical file and the newly created diary.
        app.diary_should_fail = False
        app.create_selected_outputs(print_after=True)
    finally:
        printer_support.print_files = original_print_files  # type: ignore[assignment]
        actions_creation_orchestrator.messagebox.showwarning = original_warning
        actions_creation_orchestrator.messagebox.showerror = original_error

    diaries = list(output.glob("Дневники*.docx"))
    assert len(diaries) == 1 and diaries[0].read_bytes() == b"diary-document-ok", diaries
    assert app.medical_generation_calls == 1, app.medical_generation_calls
    assert app.diary_generation_calls == 2, app.diary_generation_calls
    assert len(print_calls) == 1, print_calls
    assert print_calls[0][1] == "Test Printer", print_calls
    assert print_calls[0][0] == [saved[0], diaries[0]], print_calls
    assert app._pending_print_retry_files == [], app._pending_print_retry_files
    assert len(list(output.glob("Первичный осмотр*.docx"))) == 1



def _assert_diary_source_buttons_route_to_expected_picker() -> None:
    source = (Path(__file__).resolve().parents[1] / "layout_sources.py").read_text(encoding="utf-8")
    start = source.index("    def _diary_compact_row")
    tail = source[start:]
    assert 'text="Тексты"' in tail and 'command=self.choose_status_file' in tail, tail[:5000]
    assert 'text="Папка"' in tail and 'command=self.choose_status_files' in tail, tail[:5000]
    assert 'text="Даты"' in tail and 'command=self.choose_diary_files' in tail, tail[:5000]


def main() -> None:
    _assert_generation_action_gate_blocks_reentry_and_queued_double_click()
    _assert_ui_diagnosis_wins_snapshot()
    _assert_full_patient_switch_reset_matrix()
    _assert_diary_source_buttons_route_to_expected_picker()
    with TemporaryDirectory(prefix="diagnosis-override-regression-") as temp_dir:
        root = Path(temp_dir)
        _assert_pending_print_close_safety(root)
        _assert_compact_address_never_consumes_clinical_residence_narrative()
        _assert_epidemiology_stops_before_post_section_admission_prose()
        _assert_epi_cleanup_preserves_examination_plan_item()
        _assert_missing_source_fails_before_any_medical_popup(root)
        _assert_output_path_file_fails_before_any_medical_popup(root)
        _assert_changed_source_fails_before_any_medical_popup(root)
        _assert_changed_source_blocks_diary_only_but_manual_diary_mode_survives(root)
        _assert_changed_diary_input_files_fail_closed(root)
        _assert_dropped_epi_is_bound_to_drop_time_revision(root)
        _assert_changed_epi_fails_before_any_medical_popup(root)
        _assert_primary_cache_rejects_same_metadata_wrong_digest(root)
        _assert_same_path_replacement_is_patient_switch(root)
        _assert_invalid_new_source_preserves_open_patient(root)
        _assert_multi_primary_drop_fails_safe(root)
        _assert_failed_primary_aborts_entire_drop_batch(root)
        _assert_multi_epi_drop_fails_before_primary_apply(root)
        _assert_multi_diary_source_folder_drop_fails_before_primary_apply(root)
        _assert_all_medical_documents_receive_new_diagnosis(root)
        _assert_verbal_matching_and_word_formats(root)
        _assert_auto_refresh_and_manual_pin(root)
        _assert_manual_text_is_patient_scoped(root)
        _assert_manual_picker_accepts_doc(root)
        _assert_failed_auto_match_offers_manual_word_file(root)
        _assert_commit_rollback_reports_residual_file(root)
        _assert_print_failure_retries_same_saved_file_without_regeneration(root)
        _assert_diary_failure_keeps_medical_documents(root)
        _assert_legacy_doc_parser_route(root)
    _assert_diary_creation_path_offers_manual_fallback()
    print(
        "DIAGNOSIS OVERRIDE REGRESSION OK: UI diagnosis + verbal matching + "
        "manual fallback + visible Word picker + compact address/anamnesis isolation + digest-bound source/cache + missing/changed primary/EPI/Texts/Dates + invalid output-path early fail + DnD EPI revision pin + pending-print close safety + generation double-click guard + partial-set survival + complete partial-print retry + print retry without regeneration + .doc/.docx source + same-path replacement isolation + transactional invalid-source handling + atomic primary DnD + multi-primary/multi-EPI/multi-folder fail-safe + complete patient-session reset matrix"
    )


if __name__ == "__main__":
    main()
