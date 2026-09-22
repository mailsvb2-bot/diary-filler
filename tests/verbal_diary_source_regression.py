"""Regression for the real physician diary-text folder contract.

The «Тексты» folder contains Word files named by diagnosis WORDS, often in
informal Russian, and normally contains no ICD codes at all. It is a different
source type from optional numeric «Даты» 01–31.
"""
from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document

import diary_service
from actions_diary_flow import ActionsDiaryFlowMixin
from diary_service import dynamic_epicrisis_base_date, dynamic_epicrisis_dates
from diary_template_selection import DiaryTemplateSelectionMixin
from diary_text_selection import (
    find_diary_text_file_for_diagnosis,
    normalize_diary_diagnosis_name,
)
from medical_models import PatientData


class _Var:
    def __init__(self, value=""):
        self.value = value

    def get(self):
        return self.value

    def set(self, value):
        self.value = value


class _Result:
    created_files: list[Path] = []
    report_path = None
    processed_files = 0
    filled_rows = 0
    month_cells_filled = 0
    final_rows_filled = 0
    removed_after_discharge_rows = 0


class _DiaryFlowHarness(ActionsDiaryFlowMixin):
    def __init__(self, text_file: Path, output_dir: Path):
        self.navigation_path_var = _Var("")
        self.patient_name_var = _Var("Маркер Мужской Тестовый")
        self.admission_date_var = _Var("18.05.2026")
        self.discharge_date_var = _Var("25.05.2026")
        self.diagnosis_var = _Var("F06.8 Органическое расстройство личности")
        self.status_files = [str(text_file)]
        self._diary_text_files_auto_selected = True
        self.diary_texts_dir = str(text_file.parent)
        self.diary_files: list[str] = []
        self._diary_files_auto_selected = False
        self.repeat_statuses_var = _Var(True)
        self.force_final_diary_var = _Var(True)
        self.output_dir = output_dir
        self.word_refresh_calls = 0
        self.numbered_lookup_calls = 0

    def _auto_select_diary_text_by_diagnosis(self, **_kwargs):
        self.word_refresh_calls += 1
        return True

    def _offer_manual_diary_text_file(self, **_kwargs):
        raise AssertionError("manual fallback must not open when word-matched text exists")

    def _auto_select_numbered_diary_template(self, **_kwargs):
        self.numbered_lookup_calls += 1
        raise AssertionError("numeric Dates lookup must not run for a valid Texts source")

    def _result_output_dir(self):
        return self.output_dir

    def _effective_staff_profile(self):
        return {"doctor": "Врач", "department_head": "Заведующий"}

    def _diagnostic_reports_enabled(self):
        return False

    def _log(self, _text):
        pass


class _DateLookupHarness(DiaryTemplateSelectionMixin):
    def __init__(self, text_dir: Path):
        self.diary_texts_dir = str(text_dir)
        self.lookup_calls = 0

    def _get_saved_directory(self, _key):
        return ""

    def _find_numbered_diary_template(self, _folder, _day):
        self.lookup_calls += 1
        raise AssertionError("Texts folder reached numeric 01-31 scanner")


def _assert_words_only_matching(root: Path) -> None:
    folder = root / "Тексты"
    folder.mkdir()
    organic = folder / "дневники на органичку.docx"
    schizophrenia = folder / "дневники на шизофреника.doc"
    unrelated = folder / "дневники на депрессию.docx"
    numeric = folder / "20.docx"
    for path in (organic, schizophrenia, unrelated, numeric):
        path.touch()

    # Codes and all numeric fragments disappear before comparison.
    normalized = normalize_diary_diagnosis_name("F06.8 Органическое расстройство личности 2026")
    assert "f06" not in normalized and "06" not in normalized and "2026" not in normalized, normalized
    assert "органичес" in normalized, normalized
    assert normalize_diary_diagnosis_name("F20.0") == ""

    found = find_diary_text_file_for_diagnosis(folder, "F06.8 Органическое расстройство личности")
    assert found == organic, found
    found = find_diary_text_file_for_diagnosis(folder, "F20.0 Параноидная шизофрения")
    assert found == schizophrenia, found
    assert find_diary_text_file_for_diagnosis(folder, "F20.0") is None


def _assert_selected_text_does_not_require_dates(root: Path) -> None:
    text_file = root / "дневники на органичку.docx"
    text_file.touch()
    output = root / "out"
    output.mkdir()
    app = _DiaryFlowHarness(text_file, output)
    snapshot = PatientData(
        fio="Маркер Мужской Тестовый",
        output_fio="Маркер Мужской Тестовый",
        birth="01.01.1980",
        admission_date="18.05.2026",
        discharge_date="25.05.2026",
        diagnosis="F06.8 Органическое расстройство личности",
        complaints="Жалоб не предъявляет",
        treatment_plan="Терапия по листу назначений",
        mental_status="Состояние стабильное",
        expert_sick_leave_needed="да",
        expert_sick_leave_from="20.05.2026",
        doctor="Врач",
        head="Заведующий",
    )

    captured: dict[str, object] = {}
    original = diary_service.DiaryService.create_text_diaries
    try:
        def fake_create(_service, **kwargs):
            captured.update(kwargs)
            return _Result()

        diary_service.DiaryService.create_text_diaries = fake_create
        app._create_diaries_impl(patient_data_snapshot=snapshot, log_created=False)
    finally:
        diary_service.DiaryService.create_text_diaries = original

    assert app.word_refresh_calls == 1, app.word_refresh_calls
    assert app.numbered_lookup_calls == 0, app.numbered_lookup_calls
    assert captured.get("status_files") == [str(text_file)], captured
    assert captured.get("diary_files") == [], captured
    assert captured.get("sick_leave_dynamic_epicrisis") is True, captured
    assert captured.get("sick_leave_from") == "20.05.2026", captured
    assert captured.get("birth_date") == "01.01.1980", captured
    assert captured.get("complaints") == "Жалоб не предъявляет", captured
    assert captured.get("treatment") == "Терапия по листу назначений", captured
    assert captured.get("profile_status") == "Состояние стабильное", captured


def _assert_fallback_output_survives_temporary_date_source(root: Path) -> None:
    text_dir = root / "persistent-text-source"
    text_dir.mkdir()
    text_file = text_dir / "дневники на органичку.docx"
    doc = Document()
    doc.add_paragraph("Состояние спокойное, поведение упорядоченное, продуктивной психопатологической симптоматики не выявляет.")
    doc.save(str(text_file))

    result = diary_service.DiaryService().create_text_diaries(
        status_files=[text_file],
        diary_files=[],
        output_dir=None,
        patient_name="Маркер Мужской Тестовый",
        admission_value="18.05.2026",
        discharge_value="25.05.2026",
    )

    assert result.created_files, result
    for created in result.created_files:
        assert created.exists(), f"fallback output disappeared with temporary Dates source: {created}"
        assert created.parent.resolve() == text_dir.resolve(), created
        assert ".diary-date-source-" not in str(created), created


def _assert_text_folder_never_enters_numeric_scanner(root: Path) -> None:
    text_dir = root / "Тексты"
    text_dir.mkdir(exist_ok=True)
    app = _DateLookupHarness(text_dir)
    result = app._try_find_template_in_dirs(
        [text_dir],
        [(18, "дате госпитализации", datetime(2026, 5, 18)), (19, "первому дню дневника", datetime(2026, 5, 19))],
    )
    assert result == (None, None, "", None), result
    assert app.lookup_calls == 0, app.lookup_calls


def _assert_dynamic_epicrisis_calendar_contract() -> None:
    admission = date(2026, 9, 1)
    assert dynamic_epicrisis_base_date(admission, "06.09.2026") == date(2026, 9, 6)
    assert dynamic_epicrisis_base_date(admission, "") == admission

    # 19.09.2026 is Saturday, therefore the first +10 day epicrisis moves to Monday.
    assert dynamic_epicrisis_dates(
        date(2026, 9, 9), discharge_date=date(2026, 10, 5)
    )[:2] == (date(2026, 9, 21), date(2026, 9, 29))
    # If the shifted working day reaches discharge, the entry is not created.
    assert dynamic_epicrisis_dates(
        date(2026, 9, 9), discharge_date=date(2026, 9, 21)
    ) == ()
    # Historical fixed holiday calendar: 01-09 January are non-working days.
    assert dynamic_epicrisis_dates(
        date(2026, 12, 22), discharge_date=date(2027, 1, 20)
    )[0] == date(2027, 1, 11)


def _assert_dynamic_epicrisis_is_additive(root: Path) -> None:
    text_dir = root / "dynamic-epicrisis"
    text_dir.mkdir()
    text_file = text_dir / "status.docx"
    doc = Document()
    doc.add_paragraph("Состояние спокойное. Контакт продуктивный.")
    doc.save(str(text_file))

    without_dir = text_dir / "without"
    with_dir = text_dir / "with"
    without_dir.mkdir()
    with_dir.mkdir()
    service = diary_service.DiaryService()

    without_result = service.create_text_diaries(
        status_files=[text_file],
        diary_files=[],
        output_dir=without_dir,
        patient_name="Маркер Мужской Тестовый",
        admission_value="01.09.2026",
        discharge_value="25.09.2026",
        doctor_name="Врач В.В.",
        department_head_name="Заведующий З.З.",
        sick_leave_dynamic_epicrisis=False,
    )
    without_text = [p.text for p in Document(str(without_result.created_files[0])).paragraphs]
    assert not any("Динамический эпикриз." in text for text in without_text), without_text
    assert not hasattr(without_result, "dynamic_epicrisis_count")

    with_result = service.create_text_diaries(
        status_files=[text_file],
        diary_files=[],
        output_dir=with_dir,
        patient_name="Маркер Мужской Тестовый",
        admission_value="01.09.2026",
        discharge_value="25.09.2026",
        doctor_name="Врач В.В.",
        department_head_name="Заведующий З.З.",
        sick_leave_dynamic_epicrisis=True,
        sick_leave_from="01.09.2026",
        birth_date="01.01.1980",
        complaints="Жалоб не предъявляет",
        treatment="Терапия по листу назначений",
        profile_status="Состояние стабильное",
    )
    paragraphs = [p.text for p in Document(str(with_result.created_files[0])).paragraphs]
    epicrisis_heads = [text for text in paragraphs if "Динамический эпикриз." in text]
    assert epicrisis_heads == [
        "11.09.26 Динамический эпикриз.",
        "21.09.26 Динамический эпикриз.",
    ], epicrisis_heads
    assert getattr(with_result, "dynamic_epicrisis_count", 0) == 2
    assert "Психический статус: Состояние стабильное." in paragraphs, paragraphs
    assert not any("Профильный статус:" in text for text in paragraphs), paragraphs

    # The ordinary clinical sequence remains present; an epicrisis on 11.09 is
    # inserted after that date's ordinary diary, never replacing it.
    ordinary_11 = [i for i, text in enumerate(paragraphs) if text.startswith("11.09.26 ") and "Динамический эпикриз." not in text]
    dynamic_11 = [i for i, text in enumerate(paragraphs) if text == "11.09.26 Динамический эпикриз."]
    assert ordinary_11 and dynamic_11 and ordinary_11[0] < dynamic_11[0], paragraphs
    assert any(text.startswith("25.09.26 ") for text in paragraphs), "final discharge diary disappeared"


def main() -> None:
    with TemporaryDirectory(prefix="verbal-diary-source-") as temp_dir:
        root = Path(temp_dir)
        _assert_words_only_matching(root)
        _assert_selected_text_does_not_require_dates(root)
        _assert_fallback_output_survives_temporary_date_source(root)
        _assert_text_folder_never_enters_numeric_scanner(root)
        _assert_dynamic_epicrisis_calendar_contract()
        _assert_dynamic_epicrisis_is_additive(root)
    print(
        "VERBAL DIARY SOURCE REGRESSION OK: words-only matching; Texts never scanned as numeric Dates; "
        "fallback output persists; sick-leave dynamic epicrises are additive and calendar-locked"
    )


if __name__ == "__main__":
    main()
