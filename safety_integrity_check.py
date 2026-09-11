"""Focused regressions for patient-fact and output-integrity safety."""
from __future__ import annotations

import json
import re
import shutil
from datetime import date, timedelta
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document

import diary_batch as diary_batch_module
from app_config import DIR_DIARY_TEXTS, DIR_PRIMARY_DOCUMENTS
from diary_batch import fill_diary_batch
from diary_service import DiaryService
from diary_gender import adapt_text_to_patient_gender, detect_gender_from_patient_name
from diary_models import FillResult
from medical_gender import adapt_document_to_patient_gender, normalize_facility_references_in_document
from medical_models import PatientData
from medical_parser import MedicalTextParser
from medical_service import MedicalDocumentService
from settings_mixin import SettingsMixin

ROOT = Path(__file__).resolve().parent


def _test_missing_patient_facts() -> None:
    data = MedicalTextParser().parse_text(
        "12.01.2026 Первичный осмотр\n"
        "Ф.И.О.: Иванов Иван Иванович\n"
        "Дата рождения: 01.01.1980\n"
        "Диагноз: F41.2 тест\n"
        "Лечение: тестовое лечение"
    )
    assert data.registered == "", data.registered
    assert data.epidemiology == "", data.epidemiology
    assert any("адрес регистрации" in warning for warning in data.warnings), data.warnings
    assert any("эпидемиологический анамнез" in warning for warning in data.warnings), data.warnings


def _test_conservative_gender() -> None:
    assert detect_gender_from_patient_name("Иванов И.И.") == "male"
    assert detect_gender_from_patient_name("Иванова И.И.") == "female"
    assert detect_gender_from_patient_name("Шевченко Алексей Сергеевич") == "male"
    assert detect_gender_from_patient_name("Шевченко Анна Сергеевна") == "female"
    assert detect_gender_from_patient_name("Шевченко А.А.") is None


def _make_diary_template(path: Path, days: tuple[int, ...]) -> None:
    doc = Document()
    table = doc.add_table(rows=1, cols=4)
    for index, header in enumerate(("День госпитализации", "Число", "Месяц/Год", "Дневник наблюдения")):
        table.rows[0].cells[index].text = header
    for day in days:
        row = table.add_row()
        row.cells[0].text = str(day)
        row.cells[3].text = "Лечащий врач Балаганин С.В."
    doc.save(path)


def _make_statuses(path: Path) -> None:
    doc = Document()
    doc.add_paragraph("Пациент спокоен, жалоб не предъявляет, контакт доступен, сон достаточный.")
    doc.save(path)



def _test_common_diary_gender_agreement() -> None:
    source = (
        "В отделение явился без сопровождения. На лечение пришёл самостоятельно. "
        "Остаётся вялым, подавленным, пассивным, несобранным. "
        "В отделении малозаметен, бездеятелен, замкнута. "
        "Легко становится негативистичным, эмоционально монотонен, однообразен, "
        "выхолощен, холоден."
    )
    adapted, changed = adapt_text_to_patient_gender(source, "female")
    assert changed >= 13, (changed, adapted)
    import re
    for wrong in (
        "явился", "пришёл", "вялым", "подавленным", "пассивным",
        "несобранным", "малозаметен", "бездеятелен", "негативистичным",
        "монотонен", "однообразен", "выхолощен", "холоден",
    ):
        assert re.search(rf"(?<![А-Яа-яЁё]){re.escape(wrong)}(?![А-Яа-яЁё])", adapted) is None, adapted
    for expected in (
        "явилась", "пришла", "вялой", "подавленной", "пассивной",
        "несобранной", "малозаметна", "бездеятельна", "негативистичной",
        "монотонна", "однообразна", "выхолощена", "холодна",
    ):
        assert expected in adapted, adapted


def _test_text_diary_signature_fallbacks(tmp: Path) -> None:
    doctor_only = tmp / "doctor-only-dates.docx"
    doc = Document()
    table = doc.add_table(rows=2, cols=4)
    for index, header in enumerate(("День госпитализации", "Число", "Месяц/Год", "Дневник наблюдения")):
        table.rows[0].cells[index].text = header
    table.rows[1].cells[0].text = "2"
    table.rows[1].cells[3].text = "Лечащий врач Балаганин С.В."
    doc.save(doctor_only)
    assert diary_batch_module._signature_lines_from_diary_sources([doctor_only]) == (
        "Лечащий врач Балаганин С.В.",
        "Зав. отделением ____________________",
    )

    head_only = tmp / "head-only-dates.docx"
    doc = Document()
    table = doc.add_table(rows=2, cols=4)
    for index, header in enumerate(("День госпитализации", "Число", "Месяц/Год", "Дневник наблюдения")):
        table.rows[0].cells[index].text = header
    table.rows[1].cells[0].text = "2"
    table.rows[1].cells[3].text = "Заведующий отделением Можарова Е.А."
    doc.save(head_only)
    assert diary_batch_module._signature_lines_from_diary_sources([head_only]) == (
        "Лечащий врач ____________________",
        "Заведующий отделением Можарова Е.А.",
    )


def _test_text_diary_user_route(tmp: Path) -> None:
    statuses = tmp / "text-route-texts.docx"
    template = tmp / "text-route-dates.docx"
    doc = Document()
    repeated = "Пациент спокоен, контактен, жалоб активно не предъявляет, назначения выполняет."
    doc.add_paragraph(repeated)
    doc.add_paragraph(repeated)
    doc.save(statuses)

    doc = Document()
    table = doc.add_table(rows=1, cols=4)
    for index, header in enumerate(("День госпитализации", "Число", "Месяц/Год", "Дневник наблюдения")):
        table.rows[0].cells[index].text = header
    for hospital_day in (2, 5, 9, 12):
        row = table.add_row()
        row.cells[0].text = str(hospital_day)
        row.cells[3].text = "Лечащий врач Балаганин С.В."
        if hospital_day == 5:
            row.cells[3].text += "\nЗав.отделением Можарова Е.А."
    doc.save(template)

    preserved = diary_batch_module.read_statuses_from_files([statuses], preserve_duplicates=True)
    assert len(preserved) == 2, preserved

    result = DiaryService().create_text_diaries(
        status_files=[statuses],
        diary_files=[template],
        output_dir=tmp / "text-route-out",
        patient_name="Иванова Анна Сергеевна",
        gender_source_name="Иванова Анна Сергеевна",
        admission_value="01.01.2026",
        discharge_value="09.01.2026",
        force_final_diary=True,
        repeat_statuses=True,
    )
    rendered = Document(result.created_files[0])
    assert rendered.tables == [], "user-facing diary output must be text-only"
    text = "\n".join(paragraph.text for paragraph in rendered.paragraphs)
    # Proven Dokkomplekt clinical cadence: D0+1, +2, +3, +7, then
    # alternating +3/+4 day steps; discharge is one separate final diary.
    expected_prefixes = ("02.01.26", "03.01.26", "04.01.26", "08.01.26", "09.01.26")
    diary_lines = [
        paragraph.text.strip()
        for paragraph in rendered.paragraphs
        if re.match(r"^\d{2}\.\d{2}\.\d{2}\s", paragraph.text.strip())
    ]
    assert tuple(line.split()[0] for line in diary_lines) == expected_prefixes, diary_lines
    assert "02.01.26 Пациентка спокойна" in text, text
    assert "03.01.26 Пациентка спокойна" in text, text
    assert "04.01.26 Пациентка спокойна" in text, text
    assert "08.01.26 Пациентка спокойна" in text, text
    assert "09.01.26 Состояние улучшилось" in text, text
    for omitted in ("05.01.26", "06.01.26", "07.01.26"):
        assert omitted not in text, (omitted, text)
    assert "12.01.26" not in text, text
    assert result.detected_rows == 4, result.detected_rows
    assert result.filled_rows == 5, result.filled_rows
    assert "Лечащий врач Балаганин С.В." in text, text
    assert "Зав.отделением Можарова Е.А." in text, text
    lines = [paragraph.text.strip() for paragraph in rendered.paragraphs if paragraph.text.strip()]
    nonempty_paragraphs = [paragraph for paragraph in rendered.paragraphs if paragraph.text.strip()]
    for prefix in expected_prefixes:
        index = next(i for i, line in enumerate(lines) if line.startswith(prefix))
        assert lines[index + 1:index + 3] == [
            "Лечащий врач Балаганин С.В.",
            "Зав.отделением Можарова Е.А.",
        ], lines[index:index + 3]
        entry_paragraph = next(paragraph for paragraph in nonempty_paragraphs if paragraph.text.startswith(prefix))
        doctor_paragraph = nonempty_paragraphs[nonempty_paragraphs.index(entry_paragraph) + 1]
        assert entry_paragraph.paragraph_format.keep_together is True
        assert entry_paragraph.paragraph_format.keep_with_next is True
        assert doctor_paragraph.paragraph_format.keep_with_next is True
    assert len(Document(template).tables) == 1, "doctor-owned Dates source must not be modified"
    assert result.final_rows_filled == 1
    action_source = (ROOT / "actions_diary_flow.py").read_text(encoding="utf-8")
    assert "DiaryService().create_text_diaries" in action_source, "GUI diary button must use the production DiaryService"
    assert "fill_diary_batch" not in action_source, "GUI must not call the legacy batch facade"
    assert "text_output" not in action_source, "GUI must not select diary architecture through a boolean flag"


def _test_holiday_default_is_safe(tmp: Path) -> None:
    statuses = tmp / "texts.docx"
    template = tmp / "01.docx"
    _make_statuses(statuses)
    _make_diary_template(template, (2, 3))
    result = fill_diary_batch(
        status_files=[statuses],
        diary_files=[template],
        output_dir=tmp / "holiday-out",
        patient_name="Иванов Иван Иванович",
        admission_value="01.01.2026",
        force_final_diary=False,
        open_result_folder=False,
    )
    out_doc = Document(result.created_files[0])
    assert len(out_doc.tables[0].rows) == 3, "holiday rows must not be deleted unless explicitly enabled"
    assert result.removed_holiday_rows == 0


def _test_run_formatting_preserved() -> None:
    doc = Document()
    paragraph = doc.add_paragraph()
    paragraph.add_run("Пациент ")
    bold = paragraph.add_run("ВАЖНЫЙ")
    bold.bold = True
    italic = paragraph.add_run(" находится в ГБУЗ НО ПБ №2")
    italic.italic = True
    data = PatientData(fio="Иванова Ирина Ивановна")
    adapt_document_to_patient_gender(doc, data)
    normalize_facility_references_in_document(doc)
    assert paragraph.runs[1].text == "ВАЖНЫЙ" and paragraph.runs[1].bold is True
    assert paragraph.runs[2].italic is True
    assert "Пациентка" in paragraph.text
    assert "ГБУЗ НО «НКЦПЗ» диспансер №2" in paragraph.text


class _FailSecondRenderer:
    def __init__(self) -> None:
        self.calls = 0

    def render(self, kind, template_path, output_path, data) -> None:
        self.calls += 1
        Path(output_path).write_bytes(b"staged")
        if self.calls == 2:
            raise RuntimeError("synthetic renderer failure")


def _test_medical_transaction(tmp: Path) -> None:
    nav = tmp / "primary.docx"
    Document().save(nav)
    service = MedicalDocumentService()
    service.renderer = _FailSecondRenderer()
    data = PatientData(
        case_number="123",
        fio="Иванов Иван Иванович",
        admission_date="10.06.2026",
        discharge_date="11.06.2026",
        diagnosis="F41.2 тест",
        treatment_plan="тестовое лечение",
    )
    out = tmp / "medical-tx"
    try:
        service.create_documents(
            navigation_path=nav,
            output_dir=out,
            discharge_date="11.06.2026",
            selected_docs=["primary", "discharge"],
            override_data=data,
        )
    except RuntimeError as exc:
        assert "synthetic renderer failure" in str(exc)
    else:
        raise AssertionError("synthetic renderer failure must propagate")
    assert not list(out.glob("*.docx")), list(out.glob("*.docx"))
    assert not list(out.glob(".medical-autofill-*"))


def _test_diary_transaction(tmp: Path) -> None:
    statuses = tmp / "tx-texts.docx"
    first = tmp / "tx-1.docx"
    second = tmp / "tx-2.docx"
    _make_statuses(statuses)
    _make_diary_template(first, (2,))
    shutil.copy2(first, second)
    original = diary_batch_module.fill_diary_file
    calls = {"count": 0}

    def fail_second(path, status_texts, **kwargs):
        calls["count"] += 1
        if calls["count"] == 2:
            raise RuntimeError("synthetic diary writer failure")
        return FillResult(1, 1, 1, 0, 1)

    diary_batch_module.fill_diary_file = fail_second
    out = tmp / "diary-tx"
    try:
        try:
            diary_batch_module.fill_diary_batch(
                status_files=[statuses],
                diary_files=[first, second],
                output_dir=out,
                patient_name="Иванов Иван Иванович",
                admission_value="15.04.2026",
                force_final_diary=False,
                open_result_folder=False,
            )
        except RuntimeError as exc:
            assert "synthetic diary writer failure" in str(exc)
        else:
            raise AssertionError("synthetic diary writer failure must propagate")
        assert not list(out.glob("*.docx")), list(out.glob("*.docx"))
        assert not list(out.glob(".diary-autofill-*"))
    finally:
        diary_batch_module.fill_diary_file = original


def _test_daily_diary_coverage(tmp: Path) -> None:
    statuses = tmp / "daily-texts.docx"
    template = tmp / "daily-sparse.docx"
    _make_statuses(statuses)
    _make_diary_template(template, (2, 3, 4, 7, 11))

    result = fill_diary_batch(
        status_files=[statuses],
        diary_files=[template],
        output_dir=tmp / "daily-out",
        patient_name="Иванов Иван Иванович",
        admission_value="01.01.2026",
        discharge_value="10.02.2026",
        force_final_diary=True,
        open_result_folder=False,
    )

    table = Document(result.created_files[0]).tables[0]
    data_rows = table.rows[1:]
    hospitalization_days = [int(row.cells[0].text.strip()) for row in data_rows]
    actual_dates = [f"{row.cells[1].text.strip()}.{row.cells[2].text.strip()}" for row in data_rows]

    assert hospitalization_days == list(range(2, 42)), hospitalization_days
    expected_dates = []
    current = date(2026, 1, 2)
    while current <= date(2026, 2, 10):
        expected_dates.append(current.strftime("%d.%m.%Y"))
        current += timedelta(days=1)
    assert actual_dates == expected_dates, actual_dates
    assert hospitalization_days[-1] > 31, hospitalization_days[-1]

    # The actual GUI uses the text route and the proven Dokkomplekt clinical
    # calendar, not a daily expansion.  This also locks cadence beyond day 31.
    assert diary_batch_module._clinical_diary_offsets(21) == (1, 2, 3, 7, 10, 14, 17, 21)
    text_result = DiaryService().create_text_diaries(
        status_files=[statuses],
        diary_files=[template],
        output_dir=tmp / "clinical-text-out",
        patient_name="Иванов Иван Иванович",
        admission_value="01.01.2026",
        discharge_value="10.02.2026",
        force_final_diary=True,
    )
    text_doc = Document(text_result.created_files[0])
    assert text_doc.tables == [], "GUI text diary must never restore the source table"
    text_dates = [
        paragraph.text.strip().split()[0]
        for paragraph in text_doc.paragraphs
        if re.match(r"^\d{2}\.\d{2}\.\d{2}\s", paragraph.text.strip())
    ]
    expected_text_dates = [
        "02.01.26", "03.01.26", "04.01.26", "08.01.26", "11.01.26",
        "15.01.26", "18.01.26", "22.01.26", "25.01.26", "29.01.26",
        "01.02.26", "05.02.26", "08.02.26", "10.02.26",
    ]
    assert text_dates == expected_text_dates, text_dates
    assert text_result.detected_rows == 13, text_result.detected_rows
    assert text_result.filled_rows == 14, text_result.filled_rows
    assert text_result.final_rows_filled == 1, text_result.final_rows_filled


class _SettingsHarness(SettingsMixin):
    pass


def _test_settings_privacy(tmp: Path) -> None:
    harness = _SettingsHarness()
    harness._settings_path = tmp / "settings.json"
    harness._settings = {}
    patient_dir = tmp / "Иванов Иван Иванович 12345"
    diary_dir = tmp / "Reusable diary texts"
    patient_dir.mkdir()
    diary_dir.mkdir()
    harness._remember_dialog_directory(DIR_PRIMARY_DOCUMENTS, str(patient_dir), selected_is_dir=True)
    assert harness._get_saved_directory(DIR_PRIMARY_DOCUMENTS) == str(patient_dir)
    harness._remember_dialog_directory(DIR_DIARY_TEXTS, str(diary_dir), selected_is_dir=True)
    on_disk = json.loads(harness._settings_path.read_text(encoding="utf-8"))
    folders = on_disk.get("folders", {})
    assert DIR_PRIMARY_DOCUMENTS not in folders, folders
    assert folders.get(DIR_DIARY_TEXTS) == str(diary_dir), folders

    harness._settings_path.write_text(
        json.dumps({"folders": {DIR_PRIMARY_DOCUMENTS: str(patient_dir), DIR_DIARY_TEXTS: str(diary_dir)}, "patient_name": "Иванов"}),
        encoding="utf-8",
    )
    loaded = harness._load_settings()
    assert DIR_PRIMARY_DOCUMENTS not in loaded.get("folders", {}), loaded
    assert "patient_name" not in loaded, loaded


def _test_release_pins() -> None:
    build_requirements = (ROOT / "requirements_build.txt").read_text(encoding="utf-8")
    assert "python-docx==1.2.0" in build_requirements
    assert "pyinstaller==6.21.0" in build_requirements
    workflow = (ROOT / ".github/workflows/windows-build.yml").read_text(encoding="utf-8")
    assert "actions/checkout@34e114876b0b11c390a56381ad16ebd13914f8d5" in workflow
    assert "actions/setup-python@a26af69be951a213d495a4c3e4e4022e16d87065" in workflow
    assert "actions/upload-artifact@ea165f8d65b6e75b540449e92b4886f43607fa02" in workflow


def main() -> None:
    _test_missing_patient_facts()
    _test_conservative_gender()
    _test_common_diary_gender_agreement()
    _test_run_formatting_preserved()
    _test_release_pins()
    with TemporaryDirectory(prefix="medical-autofill-safety-") as temp_dir:
        tmp = Path(temp_dir)
        _test_holiday_default_is_safe(tmp)
        _test_text_diary_signature_fallbacks(tmp)
        _test_text_diary_user_route(tmp)
        _test_daily_diary_coverage(tmp)
        _test_medical_transaction(tmp)
        _test_diary_transaction(tmp)
        _test_settings_privacy(tmp)
    print("SAFETY INTEGRITY CHECK OK")


if __name__ == "__main__":
    main()
